"""Host-AI section writing: portable inputs, declared evidence, and editable output.

Validation checks references and asset readability, not the truth of free prose.
"""

from __future__ import annotations

import json
import re
import shutil
import tempfile
from contextlib import contextmanager
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Annotated, Literal

from PIL import Image
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    StrictStr,
    field_validator,
    model_serializer,
    model_validator,
)

from cfdpaper.publication.elements import MathNode, SectionEquation, SectionTable, math_text
from cfdpaper.publication.style import FigureSizing, PublicationStyle, figure_placement


class _Record(BaseModel):
    model_config = ConfigDict(extra="forbid")

    @field_validator("*", mode="after")
    @classmethod
    def reject_blank_required_text(cls, value, info):
        if isinstance(value, str) and info.field_name not in {"context", "unit"}:
            if not value.strip():
                raise ValueError(f"{info.field_name} must not be blank")
        return value


class _Figure(_Record):
    id: str = Field(pattern=r"^[A-Za-z0-9_.-]+$")
    path: str = Field(min_length=1)
    caption: str = Field(min_length=1)
    description: str = Field(min_length=1)
    sizing: FigureSizing = Field(default_factory=FigureSizing)
    placement: Literal["default", "section-end"] = "default"

    @model_serializer(mode="wrap")
    def serialize_figure(self, handler):
        result = handler(self)
        if self.placement == "default":
            result.pop("placement", None)
        return result


class _ResultRef(_Record):
    calculation_id: StrictStr
    group: StrictStr
    field: Literal[
        "count",
        "sum",
        "mean",
        "cv",
        "area",
        "rate",
        "mean_flux",
        "regional_flux",
        "shares",
        "value",
        "difference",
        "relative_change",
        "relative_reduction",
        "ratio",
        "weighted_mean",
        "weighted_std",
        "weight_sum",
        "measure_above",
        "exceedance_fraction",
        "sample_count",
        "duration",
        "start_value",
        "end_value",
        "change",
        "minimum",
        "maximum",
        "peak_time",
        "time_mean",
        "integral",
        "first_crossing_time",
        "reference_value",
        "reference_lower",
        "reference_upper",
        "outside_distance",
        "net",
        "positive",
        "negative",
        "absolute",
        "measure",
        "reference_mean",
        "comparison_mean",
        "difference_mean",
        "residual",
        "input",
        "output",
        "source",
        "storage",
    ]
    source_record: int | None = Field(default=None, strict=True, ge=2)
    places: int = Field(default=3, strict=True, ge=0, le=12)
    percentage: bool = Field(default=False, strict=True)


class _JsonValueRef(BaseModel):
    model_config = ConfigDict(extra="forbid")
    source: StrictStr = Field(min_length=1)
    pointer: StrictStr
    unit: StrictStr
    places: int = Field(default=3, strict=True, ge=0, le=12)
    unit_pointer: StrictStr | None = None


class _Evidence(_Record):
    id: str = Field(pattern=r"^[A-Za-z0-9_.-]+$")
    text: str = Field(min_length=1)
    source: str = Field(min_length=1)
    kind: Literal["observation", "metric", "interpretation", "literature"]
    value: StrictStr | None = None
    unit: StrictStr = ""
    result_ref: _ResultRef | None = None
    json_value_ref: _JsonValueRef | None = None
    citation_record: dict | None = None
    evidence_role: (
        Literal[
            "algebraic",
            "same-solution-diagnostic",
            "condition-contrast",
            "model-contrast",
            "independent-validation",
        ]
        | None
    ) = None

    @field_validator("citation_record")
    @classmethod
    def citation_metadata(cls, value):
        if value is not None:
            from cfdpaper.publication.literature import validate_citation_record

            return validate_citation_record(value)
        return value

    @model_serializer(mode="wrap")
    def serialize_evidence(self, handler):
        result = handler(self)
        if self.json_value_ref is None:
            result.pop("json_value_ref", None)
        if self.citation_record is None:
            result.pop("citation_record", None)
        if self.evidence_role is None:
            result.pop("evidence_role", None)
        return result

    @model_validator(mode="after")
    def finite_metric(self):
        if self.citation_record is not None and self.kind != "literature":
            raise ValueError("citation_record is only for literature evidence")
        if self.json_value_ref is not None:
            if (
                self.value is not None
                or self.unit
                or self.kind != "metric"
                or self.result_ref is not None
            ):
                raise ValueError(
                    "json_value_ref needs metric evidence without manual value/unit or result_ref"
                )
        if self.result_ref is not None:
            if self.value is not None or self.unit or self.kind != "metric":
                raise ValueError("result_ref requires metric evidence without manual value or unit")
        if self.kind == "metric" and self.value is not None:
            try:
                valid = Decimal(self.value).is_finite()
            except InvalidOperation:
                valid = False
            if not valid:
                raise ValueError("Metric value must be a finite numeric string")
        return self


class _Duty(_Record):
    purpose: str = Field(min_length=1)
    evidence_ids: list[str] = Field(default_factory=list)
    figure_ids: list[str]


class _TableCalculation(_Record):
    id: str = Field(pattern=r"^[A-Za-z0-9_.-]+$")
    source: str
    operation: Literal[
        "population",
        "partition",
        "scalar_select",
        "paired_change",
        "weighted_population",
        "weighted_exceedance",
        "temporal",
        "temporal_difference",
        "reference_range",
        "paired_reference_change",
    ]
    columns: dict[str, StrictStr]
    units: dict[str, StrictStr]
    domain: str
    support: dict | None = None
    group_by: str | None = None
    pair_by: StrictStr | None = None
    reference: StrictStr | None = None
    comparison: StrictStr | None = None
    quantity_kind: Literal["ordinary", "absolute-temperature", "temperature-difference"] = (
        "ordinary"
    )
    temperature_reference: float | None = Field(default=None, allow_inf_nan=False)
    weight_kind: Literal["area", "volume"] | None = None
    region_fraction: StrictStr | None = None
    region_complement: bool = Field(default=False, strict=True)
    row_filters: dict[StrictStr, StrictStr] = Field(default_factory=dict)
    unit_column: StrictStr | None = None
    time_window: (
        tuple[
            Annotated[float, Field(allow_inf_nan=False)],
            Annotated[float, Field(allow_inf_nan=False)],
        ]
        | None
    ) = None
    temporal_value_kind: Literal["instantaneous", "cumulative"] = "instantaneous"
    threshold: float | None = Field(default=None, allow_inf_nan=False)
    exceedance_relation: Literal["above", "at-or-above"] = "above"
    crossing_direction: Literal["at-or-above", "at-or-below"] = "at-or-above"
    reference_kind: (
        Literal["reported-range", "figure-reading-envelope", "declared-bound"] | None
    ) = None

    @field_validator("threshold", mode="before")
    @classmethod
    def numeric_exceedance_threshold(cls, value, info):
        if (
            info.data.get("operation") == "weighted_exceedance"
            and value is not None
            and type(value) not in (int, float)
        ):
            raise ValueError("Weighted exceedance threshold must be a numeric value")
        return value

    @model_serializer(mode="wrap")
    def serialize_calculation(self, handler):
        result = handler(self)
        if self.support is None:
            result.pop("support", None)
        if not self.row_filters:
            result.pop("row_filters", None)
        if self.unit_column is None:
            result.pop("unit_column", None)
        if self.region_fraction is None and not self.region_complement:
            result.pop("region_fraction", None)
            result.pop("region_complement", None)
        # New paired options must not make an unchanged historical input look stale.
        if self.operation not in {"paired_change", "paired_reference_change"}:
            for key in (
                "pair_by",
                "reference",
                "comparison",
                "temperature_reference",
            ):
                result.pop(key, None)
        if self.operation not in {"temporal", "temporal_difference"}:
            for key in ("time_window", "temporal_value_kind", "threshold", "crossing_direction"):
                if key == "threshold" and self.operation == "weighted_exceedance":
                    continue
                result.pop(key, None)
        if self.operation != "weighted_exceedance":
            result.pop("exceedance_relation", None)
        if self.operation not in {"reference_range", "paired_reference_change"}:
            result.pop("reference_kind", None)
        if self.operation not in {
            "paired_change",
            "weighted_population",
            "weighted_exceedance",
            "temporal",
            "temporal_difference",
            "reference_range",
            "paired_reference_change",
        } and not (self.operation == "scalar_select" and self.quantity_kind != "ordinary"):
            result.pop("quantity_kind", None)
        if self.operation not in {"weighted_population", "weighted_exceedance"}:
            result.pop("weight_kind", None)
        return result

    @model_validator(mode="after")
    def explicit_definition(self):
        required = {
            "partition": {"area", "rate"},
            "weighted_population": {"value", "weight"},
            "weighted_exceedance": {"value", "weight"},
            "temporal": {"time", "value"},
            "temporal_difference": {"time", "reference", "comparison"},
        }.get(self.operation, {"value"})
        if self.operation in {"reference_range", "paired_reference_change"}:
            if set(self.columns) not in (
                {"value", "reference", "lower", "upper"},
                {"value", "reference", "half_width"},
            ):
                raise ValueError("Reference range requires value/reference and declared bounds")
            required = set(self.columns)
            if self.reference_kind is None or len(set(self.units.values())) != 1:
                raise ValueError("Reference range requires reference_kind and matching role units")
        elif self.reference_kind is not None:
            raise ValueError("reference_kind is only for reference range operations")
        if set(self.columns) != required or set(self.units) != required:
            raise ValueError("Calculation columns and units must match its operation")
        if any(not c.strip() for c in self.columns.values()):
            raise ValueError("Calculation columns must not be blank")
        if self.operation in {"temporal", "temporal_difference"}:
            if self.units["time"] != "s":
                raise ValueError("Temporal calculations require time in seconds")
            if self.time_window is not None and self.time_window[0] >= self.time_window[1]:
                raise ValueError("Time window requires increasing sampled endpoints")
            if self.operation == "temporal_difference":
                from cfdpaper.publication.temporal import validate_temporal_difference_options

                validate_temporal_difference_options(
                    self.units,
                    self.time_window,
                    self.temporal_value_kind,
                    self.threshold,
                    self.crossing_direction,
                )
        elif (
            self.time_window is not None
            or self.temporal_value_kind != "instantaneous"
            or (self.threshold is not None and self.operation != "weighted_exceedance")
            or self.crossing_direction != "at-or-above"
        ):
            raise ValueError("Time window and event options are only for temporal calculations")
        if self.operation == "weighted_exceedance":
            if self.threshold is None:
                raise ValueError("Weighted exceedance requires a finite threshold")
        elif self.exceedance_relation != "above":
            raise ValueError("exceedance_relation is only for weighted_exceedance")
        if any(not key.strip() for key in self.row_filters):
            raise ValueError("Row filter column names must not be blank")
        if self.unit_column is not None and "value" not in self.columns:
            raise ValueError("unit_column checks the declared value unit only")
        paired = (self.pair_by, self.reference, self.comparison)
        if self.region_fraction is not None:
            if (
                self.operation != "weighted_population"
                or self.region_fraction in self.columns.values()
            ):
                raise ValueError(
                    "Region fractions require weighted_population and a distinct column"
                )
        elif self.region_complement:
            raise ValueError("Region complement requires region_fraction")
        if self.operation in {"weighted_population", "weighted_exceedance"}:
            if self.weight_kind is None:
                raise ValueError("Weighted population requires area or volume weight_kind")
        elif self.weight_kind is not None:
            raise ValueError("weight_kind is only for weighted operations")
        if self.operation in {"paired_change", "paired_reference_change"}:
            if not all(paired) or self.reference == self.comparison:
                raise ValueError("Paired change requires distinct reference/comparison and pair_by")
        elif (
            any(x is not None for x in paired)
            or self.temperature_reference is not None
            or (
                self.quantity_kind != "ordinary"
                and self.operation
                not in {
                    "weighted_population",
                    "weighted_exceedance",
                    "scalar_select",
                    "temporal",
                    "temporal_difference",
                    "reference_range",
                }
            )
        ):
            raise ValueError("Pair selectors and temperature_reference are only for paired_change")
        return self


class _FieldCalculation(_Record):
    id: str = Field(pattern=r"^[A-Za-z0-9_.-]+$")
    source: str
    operation: Literal["field_integral", "field_difference", "energy_budget"]
    columns: dict[str, StrictStr]
    units: dict[str, StrictStr]
    domain: str = Field(min_length=1)
    quantity_kind: Literal["ordinary", "absolute-temperature", "temperature-difference"] = (
        "ordinary"
    )
    field_options: dict = Field(default_factory=dict)
    support: dict | None = None

    @model_serializer(mode="wrap")
    def serialize_field(self, handler):
        result = handler(self)
        if self.support is None:
            result.pop("support", None)
        if self.quantity_kind == "ordinary":
            result.pop("quantity_kind", None)
        return result


class _ScalarResult(_Record):
    calculation_id: StrictStr = Field(min_length=1)
    group: StrictStr = Field(min_length=1)
    field: Literal[
        "mean",
        "sum",
        "value",
        "difference",
        "weighted_mean",
        "weighted_std",
        "weight_sum",
        "area",
        "rate",
        "mean_flux",
        "time_mean",
        "integral",
        "minimum",
        "maximum",
        "start_value",
        "end_value",
        "change",
        "duration",
        "peak_time",
        "first_crossing_time",
        "sample_count",
        "net",
        "positive",
        "negative",
        "absolute",
        "measure",
        "reference_mean",
        "comparison_mean",
        "difference_mean",
        "residual",
        "input",
        "output",
        "source",
        "storage",
    ]


class _ResultComparison(_Record):
    id: str = Field(pattern=r"^[A-Za-z0-9_.-]+$")
    reference: _ScalarResult
    comparison: _ScalarResult
    domain: StrictStr = Field(min_length=1)
    definition_source: StrictStr = Field(min_length=1)
    comparison_scope: StrictStr = Field(min_length=1)
    status: Literal["supported", "unknown", "not-comparable"]
    temperature_reference: float | None = Field(default=None, allow_inf_nan=False)
    support_purpose: Literal["same-state", "event-relative", "event-timing"] | None = None
    fixed_conditions: list[StrictStr] = Field(default_factory=list)

    @model_validator(mode="after")
    def explicit_comparison(self):
        if any(not x.strip() for x in (self.domain, self.comparison_scope, self.definition_source)):
            raise ValueError("Result comparisons require a domain, scope and located definition")
        if self.reference == self.comparison:
            raise ValueError("Result comparisons require distinct references")
        return self

    def engine_arguments(self):
        if self.status != "supported":
            raise ValueError(f"{self.id}: result comparison is {self.status}")
        result = self.model_dump(exclude={"status", "comparison_scope"})
        if self.support_purpose is None:
            result.pop("support_purpose", None)
        if not self.fixed_conditions:
            result.pop("fixed_conditions", None)
        return result

    @model_serializer(mode="wrap")
    def serialize_comparison(self, handler):
        result = handler(self)
        if self.support_purpose is None:
            result.pop("support_purpose", None)
        if not self.fixed_conditions:
            result.pop("fixed_conditions", None)
        return result


class _Input(_Record):
    section_id: str = Field(min_length=1)
    title: str = Field(min_length=1)
    question: str = Field(min_length=1)
    figures: list[_Figure]
    evidence: list[_Evidence] = Field(min_length=1)
    duties: list[_Duty] = Field(min_length=1)
    context: str = ""
    source_files: list[str] = Field(default_factory=list)
    table_calculations: list[_TableCalculation | _FieldCalculation] = Field(default_factory=list)
    result_comparisons: list[_ResultComparison] = Field(default_factory=list)
    style: PublicationStyle = Field(default_factory=PublicationStyle)

    @model_serializer(mode="wrap")
    def serialize_input(self, handler):
        result = handler(self)
        if not self.result_comparisons:
            result.pop("result_comparisons", None)
        return result


class _Paragraph(_Record):
    text: str = Field(min_length=1)
    evidence_ids: list[str]
    figure_ids: list[str]
    inline_math: dict[str, MathNode] = Field(default_factory=dict)
    shared_unit: StrictStr | None = None
    paragraph_id: StrictStr | None = Field(default=None, pattern=r"^[A-Za-z0-9_.-]+$")
    formatting: list[dict] = Field(default_factory=list)
    paragraph_format: dict = Field(default_factory=dict)

    @model_serializer(mode="wrap")
    def serialize_paragraph(self, handler):
        result = handler(self)
        if self.shared_unit is None:
            result.pop("shared_unit", None)
        for key in ("paragraph_id", "formatting", "paragraph_format"):
            if not result.get(key):
                result.pop(key, None)
        return result


class _Draft(_Record):
    title: str = Field(min_length=1)
    paragraphs: list[_Paragraph] = Field(min_length=1)
    captions: dict[str, str]
    evidence_notes: list[StrictStr]
    image_observations: dict[str, Literal["viewed", "author-provided", "not-viewed"]]
    observation_notes: dict[str, StrictStr] = Field(default_factory=dict)
    tables: list[SectionTable] = Field(default_factory=list)
    equations: list[SectionEquation] = Field(default_factory=list)
    block_order: list[StrictStr] = Field(default_factory=list)

    @model_serializer(mode="wrap")
    def serialize_draft(self, handler):
        result = handler(self)
        if not self.block_order:
            result.pop("block_order", None)
        return result

    @field_validator("captions")
    @classmethod
    def nonempty_captions(cls, value):
        if any(not text.strip() for text in value.values()):
            raise ValueError("Every figure requires a nonempty caption")
        return value


def _read(path: Path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _write(path: Path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _fresh(path: Path):
    if path.exists():
        raise FileExistsError(f"Refusing to overwrite author-owned output: {path}")


@contextmanager
def _stage(output: Path):
    _fresh(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".section-", dir=output.parent) as temp:
        staged = Path(temp) / "result"
        staged.mkdir()
        yield staged
        _fresh(output)
        staged.rename(output)


def _load_input(
    path: Path, *, evidence_overrides: dict | None = None, bound_evidence: dict | None = None
) -> _Input:
    raw = _read(path)
    # Manuscript bindings are resolved from their owning section before local validation.
    if bound_evidence:
        raw["evidence"] = [
            record for record in raw.get("evidence", []) if record["id"] not in bound_evidence
        ] + [record for record in bound_evidence.values() if record is not None]
    if evidence_overrides:
        for record in raw.get("evidence", []):
            if record["id"] in evidence_overrides and record.get("kind") != "literature":
                raise ValueError("Shared literature cannot replace non-literature evidence")
        raw["evidence"] = [
            record for record in raw.get("evidence", []) if record["id"] not in evidence_overrides
        ] + [record for record in evidence_overrides.values() if record is not None]
    data = _Input.model_validate(raw)
    calc_ids = [c.id for c in [*data.table_calculations, *data.result_comparisons]]
    if len(calc_ids) != len(set(calc_ids)):
        raise ValueError("Calculation IDs must be unique")
    for calculation in data.table_calculations:
        if calculation.source not in data.source_files:
            data.source_files.append(calculation.source)
        if isinstance(calculation, _FieldCalculation):
            companion = calculation.field_options.get("reference_source")
            if companion:
                name = (Path(calculation.source).parent / companion).as_posix()
                if name not in data.source_files:
                    data.source_files.append(name)
    for record in data.evidence:
        if record.json_value_ref is not None:
            name = record.json_value_ref.source
            if name not in data.source_files:
                data.source_files.append(name)
    for comparison in data.result_comparisons:
        comparison.engine_arguments()
        name = _check_comparison_definition(path.parent, comparison.definition_source)
        if name not in data.source_files:
            data.source_files.append(name)
    asset_ids = [figure.id.casefold() for figure in data.figures]
    if len(asset_ids) != len(set(asset_ids)):
        raise ValueError("Figure IDs must be unique on case-insensitive filesystems")
    for records in (data.figures, data.evidence):
        ids = [item.id for item in records]
        if len(ids) != len(set(ids)):
            raise ValueError("Record IDs must be unique")
    evidence = {item.id for item in data.evidence}
    figures = {item.id for item in data.figures}
    for duty in data.duties:
        _references(duty.evidence_ids, evidence)
        _references(duty.figure_ids, figures)
    for figure in data.figures:
        image_path = (path.parent / figure.path).resolve()
        if image_path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
            raise ValueError("Section figures must be PNG, JPEG or TIFF")
        try:
            with Image.open(image_path) as image:
                image.verify()
            with Image.open(image_path) as image:
                image.load()
        except (OSError, ValueError) as exc:
            raise ValueError(f"Unreadable figure {figure.id}: {image_path}") from exc
    return data


def _references(ids, allowed):
    unknown = set(ids) - set(allowed)
    if unknown:
        raise ValueError(f"Unresolved IDs: {sorted(unknown)}")


def _check_comparison_definition(root: Path, locator: str) -> str:
    match = re.fullmatch(r"(sources/.+):(L[1-9]\d*(?:-L?[1-9]\d*)?)", locator)
    if not match:
        raise ValueError("Comparison definition requires sources/path:Lx-Ly")
    name, lines = match.groups()
    source = (root / name).resolve()
    if not source.is_relative_to(root.resolve()) or not source.is_file():
        raise ValueError("Comparison definition source is missing or outside the package")
    content = source.read_text(encoding="utf-8-sig").splitlines()
    bounds = [int(n) for n in re.findall(r"\d+", lines)]
    if bounds[0] > bounds[-1] or bounds[-1] > len(content):
        raise ValueError("Comparison definition locator is outside source lines")
    if not "\n".join(content[bounds[0] - 1 : bounds[-1]]).strip():
        raise ValueError("Comparison definition locator contains no text")
    return name


def _copy_figures(data: _Input, source_dir: Path, output: Path):
    (output / "figures").mkdir()
    for figure in data.figures:
        source = (source_dir / figure.path).resolve()
        relative = Path("figures") / (figure.id + source.suffix.lower())
        shutil.copyfile(source, output / relative)
        figure.path = relative.as_posix()


def _copy_sources(data: _Input, source_dir: Path, output: Path):
    """Copy only explicitly supplied small review materials, preserving relative locators."""
    for name in data.source_files:
        relative = Path(name)
        if (
            not relative.parts
            or relative.is_absolute()
            or ".." in relative.parts
            or relative.parts[0] != "sources"
        ):
            raise ValueError("source_files must use relative sources/ paths")
        source = source_dir / relative
        if not source.is_file():
            raise ValueError(f"Missing source file: {name}")
        destination = output / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, destination)


def _calculate_sources(data: _Input, output: Path, *, write=True):
    """Compute on copied source tables, before the host begins drafting."""
    from cfdpaper.publication.table_evidence import (
        calculate_declared_table,
        calculate_result_comparison,
    )

    results = []
    for item in data.table_calculations:
        relative = Path(item.source)
        if relative.is_absolute() or ".." in relative.parts or relative.parts[:1] != ("sources",):
            raise ValueError("Calculation sources must use relative sources/ paths")
        result = calculate_declared_table(output / item.source, item.model_dump())
        results.append({**item.model_dump(), **result})
    parents = list(results)
    for item in data.result_comparisons:
        results.append(calculate_result_comparison(parents, **item.engine_arguments()))
    if results and write:
        _write(output / "table-results.json", results)
    parameters = {}
    for record in data.evidence:
        if record.json_value_ref is not None:
            parameters[record.id] = _resolve_metric(record, output, results)
    if parameters and write:
        _write(output / "parameter-results.json", parameters)
    return results


def _resolve_metric(record: _Evidence, root: Path, reports: list) -> dict | None:
    if record.json_value_ref is not None:
        from cfdpaper.publication.parameter_evidence import resolve_json_value

        return resolve_json_value(root, **record.json_value_ref.model_dump())
    if record.result_ref is not None:
        from cfdpaper.publication.table_evidence import resolve_table_result

        return resolve_table_result(reports, **record.result_ref.model_dump())
    return None


TASK = """# Write a figure-grounded subsection with the host AI

Read skills/cfd-evidence-writing/SKILL.md and its linked mechanism-subsections reference.
Both are included in this package; use their section route, not the legacy paragraph CLI.
Read input.json, then actually open every figure when image viewing is available.
Figure descriptions are author-provided observations, not proof that you viewed an image.
Record each figure as viewed, author-provided, or not-viewed in image_observations.
Produce a fresh JSON draft following draft-template.json; do not overwrite author edits.
Use a coherent subsection answering question and every duty, not one isolated caption per figure.
Compare cases, connect observations to defensible transport/source/geometry physics, and distinguish
observations, derived metrics, interpretations and literature. Do not invent measurements,
quantities, citations, boundary conditions, causal proof, or missing image observations.
Keep uncertainty about assumptions, source quality, normalization, and review questions in
evidence_notes, separate from reader-facing manuscript prose. Qualify scientific claims in prose
where the qualification is necessary to interpret them. Source locators are declared inputs;
the package has not independently verified the source or the physics. Explicit source_files
are copied under sources/. Read the actual tables before deriving claims; use table_evidence
helpers for equal-observation CV, partition identities and conditional temperature checks.
Document statistical domain, weighting, units and applicable cases in the evidence descriptions.
Record concrete visible features in observation_notes, keyed by figure ID; a viewed flag alone
does not describe what the image establishes. Never manufacture an observation.
If table-results.json exists, read its source definitions, units, groups and all results before
drafting. These were calculated from the copied CSVs. Missing groups remain missing; use supported
groups and report the specific gap separately. Population CV uses equal records and ddof=0;
partition results describe the supplied rows, not proof of exhaustive or disjoint physical regions.
Temporal results use actual sample spacing and exact sampled window endpoints. A sampled peak or
first qualifying threshold sample is not a solver-native event or interpolated onset. Read the
group's actual time_window; cumulative quantities have change but no second integral or time_mean.
No unit conversion or causal interpretation is performed. Resolve any conflict with supplied
metric JSON before quoting it; do not favor whichever number makes the narrative stronger.
Check free numbers in prose and captions against source values, not only token expansion.
Round each display directly from the original number. Different decimal places can both be valid.
Design plots at their final physical size with common scales and readable labels; source pixels
alone do not ensure document readability. DOCX is a general review layout, not a journal template;
render and inspect every page, reporting the actual renderer and any font substitution separately.
Apply the explicit input style: body defaults are two-character first-line indentation and
0 pt before/after spacing, implemented as paragraph properties rather than leading whitespace.
Author or journal-template overrides take precedence. Captions, headings, tables, equations,
references and image paragraphs retain their own formatting and do not inherit body indentation.
An explicitly selected prose/table-only section may have no figures; do not invent placeholders.
If a supplied figure cannot support the core claim, report the specific evidence or presentation
gap in evidence_notes. A readable or structurally valid figure is not automatic scientific approval.
Check units, boundary conditions, comparison basis and numerical limitations. A sum of
cell-integrated heat rates is not a sum of per-volume heat-release rates: the latter requires
cell-volume weighting. Keep all compared quantities on a consistent control-volume basis.

Token syntax inside paragraph text and captions:
- {{value:evidence_id}} inserts the supplied value, or resolves result_ref from the current
  source table with its declared precision and computed unit. Do not duplicate it as a manual value.
  json_value_ref reads a declared numeric JSON pointer from a copied sources/ configuration;
  parameter-results.json supplies its current value. Read the effective case/stage before using it.
- {{math:local_id}} inserts a paragraph's declared inline_math expression as editable Word math.
  Use explicit sub/sup/fraction nodes for scientific notation, not underscore strings in prose.
- Optional tables/equations use the finite structures in the packaged writing reference.
  {{table:ID}} and {{equation:ID}} reference their identifiers; these are not inline expressions.
  Table/equation IDs are currently printed literally: use publication labels such as 1 or S1,
  not internal slugs such as wall-comparison. Keep every token and object ID consistent.
- {{figure:figure_id}} inserts Figure followed by the declared figure ID.
- {{cite:evidence_id}} inserts a numbered literature citation; kind must be literature.
Use these tokens for supplied quantitative values, figure references and literature citations.
List every paragraph's evidence_ids and figure_ids; tokens must be declared in those lists.
All duty evidence and figures must be covered by paragraph declarations. Supply all captions.
Free scientific prose is allowed; successful structural validation is not semantic approval.
Return only the JSON draft, with no approval claim. A human reviews the assembled section.

Host execution capabilities:
This folder is a provider-neutral file task. A host with local tools can call the existing CLI
to assemble a fresh draft and then export Word. Use two separate commands from a directory:
python -m cfdpaper write . --artifact results-section --package P --draft D --output N
python -m cfdpaper write . --artifact results-section --package N --docx --output W
Replace P/D/N/W with package/draft/fresh-section/fresh-DOCX paths. A web host can return draft JSON
for a local operator to assemble; do not claim it executed calculations, viewed figures or exported
Word unless those tools were actually available and used. No provider SDK is required for this path.
"""


def prepare_section(
    input_path: Path,
    output_dir: Path,
    *,
    evidence_overrides: dict | None = None,
    bound_evidence: dict | None = None,
) -> Path:
    """Copy declared inputs and valid raster figures into a fresh writing package."""
    input_path, output_dir = Path(input_path), Path(output_dir)
    _fresh(output_dir)
    data = _load_input(
        input_path, evidence_overrides=evidence_overrides, bound_evidence=bound_evidence
    )
    with _stage(output_dir) as staged:
        _copy_figures(data, input_path.parent, staged)
        _copy_sources(data, input_path.parent, staged)
        _calculate_sources(data, staged)
        _write(staged / "input.json", data.model_dump())
        skill = Path(__file__).resolve().parents[1] / "skills/cfd-evidence-writing"
        if not skill.is_dir():
            skill = Path(__file__).resolve().parents[3] / "skills/cfd-evidence-writing"
        shutil.copytree(skill, staged / "skills/cfd-evidence-writing")
        (staged / "TASK.md").write_text(TASK, encoding="utf-8")
        _write(
            staged / "draft-template.json",
            {
                "title": data.title,
                "paragraphs": [{"text": "", "evidence_ids": [], "figure_ids": []}],
                "captions": {f.id: f.caption for f in data.figures},
                "evidence_notes": [],
                "image_observations": {f.id: "not-viewed" for f in data.figures},
                "observation_notes": {},
                "tables": [],
                "equations": [],
            },
        )
    return output_dir


def assemble_section(
    package_dir: Path,
    draft_path: Path,
    output_dir: Path,
    *,
    evidence_overrides: dict | None = None,
    bound_evidence: dict | None = None,
) -> Path:
    """Resolve declared references without synthesizing or approving the draft prose."""
    package_dir, draft_path, output_dir = map(Path, (package_dir, draft_path, output_dir))
    _fresh(output_dir)
    data = _load_input(
        package_dir / "input.json",
        evidence_overrides=evidence_overrides,
        bound_evidence=bound_evidence,
    )
    draft = _Draft.model_validate(_read(draft_path))
    evidence = {e.id: e for e in data.evidence}
    figures = {f.id for f in data.figures}
    if set(draft.captions) != figures or set(draft.image_observations) != figures:
        raise ValueError("Every figure requires a caption and image observation status")
    _references(draft.observation_notes, figures)
    if any(not text.strip() for text in draft.observation_notes.values()):
        raise ValueError("Observation notes must not be blank")
    used_evidence, used_figures, citations = set(), set(), []
    reports = _calculate_sources(data, package_dir, write=False)
    resolved_values = {}
    table_ids = [t.table_id for t in draft.tables]
    equation_ids = [e.equation_id for e in draft.equations]
    for ids in (table_ids, equation_ids):
        if len(ids) != len(set(ids)) or any(not re.fullmatch(r"[A-Za-z0-9_.-]+", i) for i in ids):
            raise ValueError("Table and equation IDs must be unique simple identifiers")

    def resolve(text, declared_evidence, declared_figures, shared_unit=None):
        def replace(match):
            kind, identifier = match.group(1), match.group(2)
            if kind in {"table", "equation"}:
                _references([identifier], table_ids if kind == "table" else equation_ids)
                return f"{kind.title()} {identifier}"
            if kind == "figure":
                _references([identifier], declared_figures)
                return f"Figure {identifier}"
            _references([identifier], declared_evidence)
            record = evidence[identifier]
            if kind == "value":
                if record.result_ref is not None or record.json_value_ref is not None:
                    from cfdpaper.publication.table_evidence import (
                        display_unit,
                    )

                    try:
                        resolved = _resolve_metric(record, package_dir, reports)
                    except ValueError as exc:
                        raise ValueError(f"Evidence {identifier}: {exc}") from exc
                    resolved_values[identifier] = resolved
                    unit = display_unit(resolved["unit"])
                    if shared_unit is not None:
                        if shared_unit != unit:
                            raise ValueError(f"Shared unit does not match evidence {identifier}")
                        return resolved["value"]
                    return resolved["value"] + (f" {unit}" if unit not in {"", "1"} else "")
                if record.value is None or not record.value.strip():
                    raise ValueError(f"Missing value for {identifier}")
                if shared_unit is not None:
                    if shared_unit != record.unit:
                        raise ValueError(f"Shared unit does not match evidence {identifier}")
                    return record.value
                return record.value + (f" {record.unit}" if record.unit not in {"", "1"} else "")
            if record.kind != "literature":
                raise ValueError(f"Citation {identifier} must be literature evidence")
            if identifier not in citations:
                citations.append(identifier)
            return f"[{citations.index(identifier) + 1}]"

        resolved = re.sub(
            r"\{\{(value|figure|cite|table|equation):([A-Za-z0-9_.-]+)\}\}", replace, text
        )
        if "{{" in resolved or "}}" in resolved:
            raise ValueError("Unknown or malformed draft token")
        return resolved

    paragraphs = []
    paragraph_counts = {}

    def semantic_runs(text, allowed, images, shared_unit=None, inline_math=None):
        segments, cursor, counts = [], 0, {}
        for match in re.finditer(
            r"(?P<marker>\ue000[^\ue001]+\ue001)?"
            r"(?P<token>\{\{(?P<kind>value|figure|cite|table|equation|math):"
            r"(?P<key>[A-Za-z0-9_.-]+)\}\})",
            text,
        ):
            if match.start() > cursor:
                segments.append(
                    {"text": resolve(text[cursor : match.start()], allowed, images, shared_unit)}
                )
            kind, key = match.group("kind", "key")
            occurrence = counts.get((kind, key), 0)
            counts[kind, key] = occurrence + 1
            segment = {"token": match.group("token"), "span_id": f"{kind}-{key}-{occurrence}"}
            if kind == "math":
                if inline_math is None or key not in inline_math:
                    raise ValueError(f"Unknown inline math {key}")
                expression = inline_math[key].model_dump()

                def walk(node):
                    if node["text"] is not None:
                        node["text"] = resolve(node["text"], allowed, images)
                    for child in node["children"]:
                        walk(child)

                walk(expression)
                segment["math"] = expression
                segment["text"] = math_text(MathNode.model_validate(expression))
            else:
                segment["text"] = (match.group("marker") or "") + resolve(
                    match.group("token"), allowed, images, shared_unit
                )
            segments.append(segment)
            cursor = match.end()
        if cursor < len(text):
            segments.append({"text": resolve(text[cursor:], allowed, images, shared_unit)})
        return segments

    for paragraph in draft.paragraphs:
        if paragraph.shared_unit is not None and not re.search(
            r"(?<!\w)" + re.escape(paragraph.shared_unit) + r"(?!\w)", paragraph.text
        ):
            raise ValueError("A shared unit must appear explicitly in its paragraph")
        _references(paragraph.evidence_ids, evidence)
        _references(paragraph.figure_ids, figures)
        used_evidence.update(paragraph.evidence_ids)
        used_figures.update(paragraph.figure_ids)
        runs = []
        cursor = 0
        for match in re.finditer(r"\{\{math:([A-Za-z0-9_.-]+)\}\}", paragraph.text):
            runs.append(
                {
                    "text": resolve(
                        paragraph.text[cursor : match.start()],
                        paragraph.evidence_ids,
                        paragraph.figure_ids,
                        paragraph.shared_unit,
                    )
                }
            )
            identifier = match.group(1)
            _references([identifier], paragraph.inline_math)
            expression = paragraph.inline_math[identifier].model_dump()

            def resolve_inline(node, allowed=paragraph.evidence_ids, images=paragraph.figure_ids):
                if node["text"] is not None:
                    node["text"] = resolve(node["text"], allowed, images)
                for child in node["children"]:
                    resolve_inline(child, allowed, images)

            resolve_inline(expression)
            runs.append({"math": expression})
            cursor = match.end()
        runs.append(
            {
                "text": resolve(
                    paragraph.text[cursor:],
                    paragraph.evidence_ids,
                    paragraph.figure_ids,
                    paragraph.shared_unit,
                )
            }
        )
        text = "".join(
            run["text"] if "text" in run else math_text(MathNode.model_validate(run["math"]))
            for run in runs
        )
        from cfdpaper.publication.document_blocks import paragraph_identity

        occurrence = paragraph_counts.get(paragraph.text, 0)
        paragraph_counts[paragraph.text] = occurrence + 1
        paragraphs.append(
            {
                **paragraph.model_dump(),
                "text": text,
                "runs": runs,
                "paragraph_id": paragraph.paragraph_id
                or paragraph_identity(paragraph.text, occurrence),
                "source_text": paragraph.text,
                "semantic_runs": semantic_runs(
                    paragraph.text,
                    paragraph.evidence_ids,
                    paragraph.figure_ids,
                    paragraph.shared_unit,
                    paragraph.inline_math,
                ),
            }
        )
    for duty in data.duties:
        if not set(duty.evidence_ids) <= used_evidence or not set(duty.figure_ids) <= used_figures:
            raise ValueError(f"Missing paragraph coverage for duty: {duty.purpose}")
    captions = {key: resolve(value, evidence, figures) for key, value in draft.captions.items()}
    tables = []
    for table in draft.tables:
        if table.after_section_id != data.section_id:
            raise ValueError(f"Table {table.table_id} targets a different section")
        _references(table.evidence_ids, evidence)
        content = table.model_dump()
        content["caption"] = resolve(table.caption, table.evidence_ids, figures)
        content["note"] = resolve(table.note, table.evidence_ids, figures)
        content["columns"] = [resolve(c, table.evidence_ids, figures) for c in table.columns]
        content["rows"] = [
            [resolve(c, table.evidence_ids, figures) for c in row] for row in table.rows
        ]
        tables.append(content)
    equations = []
    for equation in draft.equations:
        _references(equation.evidence_ids, evidence)
        content = equation.model_dump()

        def resolve_node(node, allowed=equation.evidence_ids):
            if node["text"] is not None:
                node["text"] = resolve(node["text"], allowed, figures)
            for child in node["children"]:
                resolve_node(child, allowed)

        resolve_node(content["expression"])
        equations.append(content)
    references = [{"number": n, **evidence[key].model_dump()} for n, key in enumerate(citations, 1)]
    for record in references:
        if record.get("citation_record") is not None:
            from cfdpaper.publication.literature import _reference_label

            record["text"] = _reference_label(record["citation_record"])
    with _stage(output_dir) as staged:
        _copy_figures(data, package_dir, staged)
        _copy_sources(data, package_dir, staged)
        if reports:
            _write(staged / "table-results.json", reports)
        result = {
            "section_id": data.section_id,
            "title": draft.title,
            "paragraphs": paragraphs,
            "figures": [
                {
                    **f.model_dump(),
                    "caption": captions[f.id],
                    "source_caption": draft.captions[f.id],
                    "caption_runs": semantic_runs(draft.captions[f.id], evidence, figures),
                }
                for f in data.figures
            ],
            "evidence": [e.model_dump() for e in data.evidence],
            "resolved_values": resolved_values,
            "style": data.style.model_dump(),
            "tables": tables,
            "equations": equations,
            "references": references,
            "image_observations": draft.image_observations,
            "observation_notes": draft.observation_notes,
            "block_order": draft.block_order,
        }
        _write(staged / "section.json", result)
        lines = [f"## {draft.title}", *[p["text"] for p in paragraphs]]
        for equation in equations:
            model = SectionEquation.model_validate(equation)
            lines.append(f"{math_text(model.expression)}   ({model.equation_id})")
        for table in tables:
            from cfdpaper.publication.export import _markdown_table

            lines.append(_markdown_table(SectionTable.model_validate(table), numbered=True))
            if table["note"]:
                lines.append(table["note"])
        for figure in result["figures"]:
            lines.extend(
                [
                    f"![Figure {figure['id']}]({figure['path']})",
                    f"Figure {figure['id']}. {figure['caption']}",
                ]
            )
        if references:
            lines.extend(
                [
                    "### References",
                    *[
                        f"[{r['number']}] {r['text']}" + reference_source_suffix(r)
                        for r in references
                    ],
                ]
            )
        (staged / "section.md").write_text("\n\n".join(lines) + "\n", encoding="utf-8")
        (staged / "evidence-notes.md").write_text(
            "# Evidence and review notes\n\n" + "\n\n".join(draft.evidence_notes) + "\n",
            encoding="utf-8",
        )
        prompt = (
            "Use the self-contained review-packet directory: review section.md against "
            "input.json, draft.json and figures within that directory. "
            "When supplied, read sources/ and table-results.json; distinguish computed "
            "identities from physical interpretations and check any quoted value "
            "against its source. "
            "Open the images; check coverage, units, normalization and physical inference. "
            'Return JSON {"suggestions": [{"target": "paragraph or figure", '
            '"comment": "finding", "recommendation": "suggested change"}]}. '
            "Suggestions are separate from author decisions; do not claim approval.\n"
        )
        (staged / "review-prompt.md").write_text(prompt, encoding="utf-8")
        packet = staged / "review-packet"
        packet.mkdir()
        _write(packet / "input.json", data.model_dump())
        _write(packet / "draft.json", draft.model_dump())
        shutil.copytree(staged / "figures", packet / "figures")
        _copy_sources(data, staged, packet)
        if (staged / "table-results.json").exists():
            shutil.copyfile(staged / "table-results.json", packet / "table-results.json")
        for name in ("section.md", "review-prompt.md", "evidence-notes.md"):
            shutil.copyfile(staged / name, packet / name)
    return output_dir


def reference_source_suffix(record):
    """Omit only a duplicate display source; retain the record's citation identity."""
    source = record["source"].strip()
    text = record["text"].strip().casefold()
    if (
        "formatted_runs" in record
        or record.get("citation_record") is not None
        or source.casefold() == text
    ):
        return ""
    doi_pattern = r"(?:doi\s*:\s*|https?://(?:dx\.)?doi\.org/)?(10\.\d{4,9}/[^\s<>\"“”]+)"

    def display_doi(value):
        value = value.casefold().rstrip(".,;:!?'’")
        # Parentheses can be part of a DOI. Remove only unmatched closing wrappers.
        while value and value[-1] in ")]}":
            opening = {")": "(", "]": "[", "}": "{"}[value[-1]]
            if value.count(value[-1]) <= value.count(opening):
                break
            value = value[:-1].rstrip(".,;:!?'’")
        return value

    source_match = re.fullmatch(doi_pattern, source, flags=re.IGNORECASE)
    if source_match:
        source_doi = display_doi(source_match[1])
        text_dois = {
            display_doi(match[1])
            for match in re.finditer(r"(?<![\w/])" + doi_pattern, text, flags=re.IGNORECASE)
        }
        if source_doi in text_dois:
            return ""
    return f" — {source}"


def export_section_docx(section_dir: Path, output_path: Path, *, layout="after-text") -> Path:
    """Export editable manuscript text with proportionally bounded embedded images."""
    section_dir, output_path = Path(section_dir), Path(output_path)
    _fresh(output_path)
    if output_path.suffix.lower() != ".docx":
        raise ValueError("DOCX output path must end in .docx")
    if layout not in {"after-text", "near-reference"}:
        raise ValueError("Layout must be after-text or near-reference")
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError(
            "DOCX export requires python-docx; install cfd-paper-agent[docs]"
        ) from exc
    data = _read(section_dir / "section.json")
    config = PublicationStyle.model_validate(data.get("style", {}))
    report_path = output_path.with_suffix(".layout.json")
    _fresh(report_path)
    document = Document()
    for name, size in (
        ("Normal", config.body_pt),
        ("Title", config.title_pt),
        ("Caption", config.caption_pt),
        ("Heading 1", config.heading_pt),
        ("Heading 2", config.heading_pt),
    ):
        style = document.styles[name]
        style.font.name = config.font_family
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in list(fonts.attrib):
            if "theme" in attribute.lower():
                del fonts.attrib[attribute]
        fonts.set(qn("w:eastAsia"), config.font_family)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        style.paragraph_format.widow_control = True
    document.styles["Caption"].font.bold = False
    document.styles["Caption"].font.italic = False
    document.styles["Caption"].paragraph_format.keep_together = True
    document.styles["Normal"].paragraph_format.line_spacing = config.line_spacing
    document.styles["Normal"].paragraph_format.space_after = Pt(config.space_after_pt)
    document.styles["Caption"].paragraph_format.line_spacing = 1.2
    document.styles["Caption"].paragraph_format.space_after = Pt(config.space_after_pt)
    for name in ("Heading 1", "Heading 2"):
        heading = document.styles[name].paragraph_format
        heading.first_line_indent = Pt(0)
        heading.space_before = Pt(8)
        heading.space_after = Pt(4)
        heading.keep_with_next = True
    page = document.sections[0]
    page.page_width, page.page_height = Mm(config.page_width_mm), Mm(config.page_height_mm)
    page.left_margin = page.right_margin = Mm(config.margin_mm)
    page.top_margin = page.bottom_margin = Mm(config.margin_mm)
    document.add_paragraph(data["title"], style="Title")
    placements = []

    def add_figure(figure):
        path = section_dir / figure["path"]
        caption = f"Figure {figure['id']}. {figure['caption']}"
        with Image.open(path) as picture:
            try:
                placement = figure_placement(
                    pixels=picture.size,
                    caption=caption,
                    sizing=FigureSizing.model_validate(figure.get("sizing", {})),
                    style=config,
                )
            except ValueError as exc:
                raise ValueError(f"Figure {figure['id']}: {exc}") from exc
        placements.append({"id": figure["id"], **placement})
        document.add_picture(
            str(path), width=Mm(placement["width_mm"]), height=Mm(placement["height_mm"])
        )
        image_paragraph = document.paragraphs[-1]
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph = document.add_paragraph(
            f"Figure {figure['id']}. {display_scientific_text(figure['caption'])}", style="Caption"
        )
        apply_figure_pagination(
            image_paragraph,
            caption_paragraph,
            config,
            page_break_before=layout == "after-text" and figure.get("placement") != "section-end",
        )

    from cfdpaper.publication.elements import (
        add_equation,
        add_table,
        apply_figure_pagination,
        display_scientific_text,
    )

    placed, placed_tables, placed_equations = set(), set(), set()
    current_section = None

    def add_section_objects(section_id):
        # Multi-section near-reference layout keeps Methods definitions before Results.
        ownership = data.get("section_objects", {}).get(section_id, {})
        for equation in data.get("equations", []):
            key = equation["equation_id"]
            if key in ownership.get("equations", []) and key not in placed_equations:
                add_equation(document, SectionEquation.model_validate(equation))
                placed_equations.add(key)
        for table in data.get("tables", []):
            key = table["table_id"]
            if key in ownership.get("tables", []) and key not in placed_tables:
                add_table(document, SectionTable.model_validate(table), config)
                placed_tables.add(key)

    def add_keywords(section_id):
        if data.get("section_objects", {}).get(section_id, {}).get("role") == "abstract":
            if data.get("keywords"):
                p = document.add_paragraph("Keywords: " + "; ".join(data["keywords"]))
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)

    def add_section_figures(section_id):
        ownership = data.get("section_objects", {})
        ids = (
            ownership.get(section_id, {}).get("figure_ids", [])
            if ownership
            else [figure["id"] for figure in data["figures"]]
        )
        for figure in data["figures"]:
            if (
                figure.get("placement") == "section-end"
                and figure["id"] in ids
                and figure["id"] not in placed
            ):
                add_figure(figure)
                placed.add(figure["id"])

    for paragraph in data["paragraphs"]:
        if "section_heading" in paragraph:
            if current_section is not None:
                add_section_figures(current_section)
            if layout == "near-reference" and current_section is not None:
                add_section_objects(current_section)
            add_keywords(current_section)
            current_section = paragraph.get("section_id")
            level = paragraph.get("level", 1)
            if level not in (1, 2):
                raise ValueError("Manuscript headings must use level 1 or 2")
            document.add_heading(paragraph["section_heading"], level=level)
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(config.body_space_before_pt)
        p.paragraph_format.space_after = Pt(config.body_space_after_pt)
        # Word stores character indentation in hundredths of a character. Keep a
        # point-based fallback for renderers that do not implement firstLineChars.
        p.paragraph_format.first_line_indent = Pt(
            config.body_pt * config.body_first_line_indent_chars
        )
        p._p.get_or_add_pPr().get_or_add_ind().set(
            qn("w:firstLineChars"), str(round(config.body_first_line_indent_chars * 100))
        )
        for run in paragraph.get("runs", [{"text": paragraph["text"]}]):
            if "text" in run:
                p.add_run(display_scientific_text(run["text"]))
            else:
                from docx.oxml import OxmlElement

                from cfdpaper.publication.elements import math_xml

                formula = OxmlElement("m:oMath")
                formula.extend(math_xml(MathNode.model_validate(run["math"])))
                p._p.append(formula)
        p.paragraph_format.keep_together = False
        p.paragraph_format.widow_control = True
        if layout == "near-reference":
            for figure in data["figures"]:
                if (
                    figure["id"] in paragraph["figure_ids"]
                    and figure["id"] not in placed
                    and figure.get("placement") != "section-end"
                ):
                    add_figure(figure)
                    placed.add(figure["id"])
    add_section_figures(current_section)
    add_keywords(current_section)
    if layout == "near-reference" and current_section is not None:
        add_section_objects(current_section)
    for equation in data.get("equations", []):
        if equation["equation_id"] not in placed_equations:
            add_equation(document, SectionEquation.model_validate(equation))
    for table in data.get("tables", []):
        if table["table_id"] not in placed_tables:
            add_table(document, SectionTable.model_validate(table), config)
    for figure in data["figures"]:
        if figure["id"] not in placed:
            add_figure(figure)
    if data["references"]:
        document.add_heading("References", level=2)
        for record in data["references"]:
            p = document.add_paragraph(f"[{record['number']}] ")
            if "formatted_runs" in record:
                for item in record["formatted_runs"]:
                    run = p.add_run(item["text"])
                    run.italic = item.get("italic", False)
                    run.bold = item.get("bold", False)
            else:
                p.add_run(record["text"] + reference_source_suffix(record))
            for run in p.runs:
                run.font.size = Pt(config.reference_pt)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".section-", dir=output_path.parent) as temp:
        staged = Path(temp) / "section.docx"
        document.save(staged)
        staged_report = Path(temp) / "layout.json"
        _write(
            staged_report, {"layout": layout, "style": config.model_dump(), "figures": placements}
        )
        _fresh(output_path)
        _fresh(report_path)
        staged.rename(output_path)
        staged_report.rename(report_path)
    return output_path


def import_section_review(section_dir: Path, review_path: Path, output_path: Path) -> Path:
    """Store reviewer suggestions separately; never revise or approve the manuscript."""
    section_dir, review_path, output_path = map(Path, (section_dir, review_path, output_path))
    _fresh(output_path)
    _read(section_dir / "section.json")
    review = _read(review_path)
    if not isinstance(review, dict) or set(review) != {"suggestions"}:
        raise ValueError("Review must contain only a suggestions array")
    if not isinstance(review["suggestions"], list):
        raise ValueError("suggestions must be an array")
    for suggestion in review["suggestions"]:
        if (
            not isinstance(suggestion, dict)
            or set(suggestion) != {"target", "comment", "recommendation"}
            or any(not isinstance(v, str) or not v.strip() for v in suggestion.values())
        ):
            raise ValueError("Each suggestion requires target, comment and recommendation strings")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("x", encoding="utf-8") as stream:
        json.dump(review, stream, ensure_ascii=False, indent=2)
        stream.write("\n")
    return output_path
