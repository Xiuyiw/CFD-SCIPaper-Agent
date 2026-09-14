"""Small, editable manuscript tables and mathematical expressions."""

from __future__ import annotations

import math
import re
from typing import Literal

from pydantic import BaseModel, ConfigDict, Field, model_validator

from cfdpaper.publication.export import ManuscriptTable


class SectionTable(ManuscriptTable):
    evidence_ids: list[str] = Field(default_factory=list)
    column_widths_mm: list[float] | None = None
    numeric_columns: list[int] = Field(default_factory=list)
    note: str = ""

    @model_validator(mode="after")
    def column_layout(self):
        import math

        if self.column_widths_mm is not None and (
            len(self.column_widths_mm) != len(self.columns)
            or any(not math.isfinite(w) or w <= 0 for w in self.column_widths_mm)
        ):
            raise ValueError("Table column widths must be finite positive widths for every column")
        if any(i < 0 or i >= len(self.columns) for i in self.numeric_columns):
            raise ValueError("Numeric column indices must identify existing columns (zero-based)")
        return self


class MathNode(BaseModel):
    model_config = ConfigDict(extra="forbid")
    kind: Literal[
        "text", "symbol", "row", "sub", "sup", "subsup", "fraction", "sqrt", "dot", "overbar"
    ]
    text: str | None = None
    children: list[MathNode] = Field(default_factory=list)

    @model_validator(mode="after")
    def shape(self):
        if self.kind in {"text", "symbol"}:
            if not self.text or self.children:
                raise ValueError("Text/symbol nodes require text and no children")
        else:
            counts = {
                "sub": 2,
                "sup": 2,
                "subsup": 3,
                "fraction": 2,
                "sqrt": 1,
                "dot": 1,
                "overbar": 1,
            }
            if self.text is not None or not self.children:
                raise ValueError("Composite math nodes require children and no text")
            if self.kind in counts and len(self.children) != counts[self.kind]:
                raise ValueError(f"{self.kind} requires {counts[self.kind]} children")
        return self


class SectionEquation(BaseModel):
    model_config = ConfigDict(extra="forbid")
    equation_id: str = Field(pattern=r"^[A-Za-z0-9_.-]+$")
    expression: MathNode
    evidence_ids: list[str] = Field(default_factory=list)


def math_text(node: MathNode) -> str:
    if node.kind in {"text", "symbol"}:
        return node.text
    values = [math_text(child) for child in node.children]
    if node.kind == "row":
        return "".join(values)
    if node.kind in {"dot", "overbar"}:
        return f"{node.kind}({values[0]})"
    if node.kind == "sqrt":
        return f"√({values[0]})"
    if node.kind == "fraction":
        return f"({values[0]})/({values[1]})"
    if node.kind == "subsup":
        return f"{values[0]}_({values[1]})^({values[2]})"
    return values[0] + ("_(" if node.kind == "sub" else "^(") + values[1] + ")"


def math_xml(node: MathNode):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if node.kind in {"text", "symbol"}:
        # Some OMML importers interpret a multi-letter symbol as a formula-language
        # keyword (e.g. Re -> real part). Separate variable letters, retaining their
        # math styling and the original character sequence.
        if node.kind == "symbol" and len(node.text) > 1 and node.text.isalpha():
            return [
                element
                for character in node.text
                for element in math_xml(MathNode(kind="symbol", text=character))
            ]
        # Older importers escape parentheses before quoting normal-text runs. That
        # exposes backslashes. Emit these punctuation segments as math, not prose.
        if node.kind == "text" and any(char in node.text for char in "()"):
            return [
                element
                for part in re.split(r"([()])", node.text)
                if part
                for element in math_xml(
                    MathNode(kind="symbol" if part in {"(", ")"} else "text", text=part)
                )
            ]
        run = OxmlElement("m:r")
        if node.kind == "text":
            props, normal = OxmlElement("m:rPr"), OxmlElement("m:nor")
            props.append(normal)
            run.append(props)
        text = OxmlElement("m:t")
        text.set(qn("xml:space"), "preserve")
        text.text = node.text
        run.append(text)
        return [run]
    if node.kind == "row":
        return [element for child in node.children for element in math_xml(child)]
    if node.kind in {"dot", "overbar"}:
        accent = OxmlElement("m:acc")
        props, character = OxmlElement("m:accPr"), OxmlElement("m:chr")
        character.set(qn("m:val"), "\u0307" if node.kind == "dot" else "\u0305")
        props.append(character)
        accent.append(props)
        base = OxmlElement("m:e")
        base.extend(math_xml(node.children[0]))
        accent.append(base)
        return [accent]
    tags = {
        "sub": ("sSub", ["e", "sub"]),
        "sup": ("sSup", ["e", "sup"]),
        "subsup": ("sSubSup", ["e", "sub", "sup"]),
        "fraction": ("f", ["num", "den"]),
        "sqrt": ("rad", ["e"]),
    }
    tag, slots = tags[node.kind]
    element = OxmlElement(f"m:{tag}")
    if node.kind == "sqrt":
        props, hidden = OxmlElement("m:radPr"), OxmlElement("m:degHide")
        hidden.set(qn("m:val"), "1")
        props.append(hidden)
        element.append(props)
        element.append(OxmlElement("m:deg"))
    for name, child in zip(slots, node.children, strict=True):
        slot = OxmlElement(f"m:{name}")
        slot.extend(math_xml(child))
        element.append(slot)
    return [element]


def add_equation(document, equation: SectionEquation):
    from docx.oxml import OxmlElement

    paragraph = document.add_paragraph()
    paragraph.alignment = 1
    formula = OxmlElement("m:oMath")
    formula.extend(math_xml(equation.expression))
    paragraph._p.append(formula)
    paragraph.add_run(f"   ({equation.equation_id})")


def add_table(document, table: SectionTable, style):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    available = style.page_width_mm - 2 * style.margin_mm
    widths = table.column_widths_mm or [available / len(table.columns)] * len(table.columns)
    if sum(widths) > available + 0.001:
        raise ValueError(f"Table {table.table_id}: column widths exceed the text area")
    caption = document.add_paragraph(f"Table {table.table_id}. {table.caption}", style="Caption")
    caption.paragraph_format.keep_with_next = True

    # Only bind tables whose conservative wrapped-text estimate occupies at most
    # half a page. Long tables must remain free to paginate with repeated headers.
    def lines(value, width):
        chars = max(1, int(width * 72 / 25.4 / (style.caption_pt * 0.6)))
        return sum(max(1, math.ceil(len(line) / chars)) for line in value.split("\n"))

    height_pt = sum(
        max(lines(value, width) for value, width in zip(values, widths, strict=True))
        * style.caption_pt
        * 1.2
        + 6
        for values in [table.columns, *table.rows]
    )
    height_pt += (lines(caption.text, available) + lines(table.note, available)) * (
        style.caption_pt * 1.2
    )
    compact = height_pt <= (style.page_height_mm - 2 * style.margin_mm) * 72 / 25.4 / 2
    item = document.add_table(rows=1, cols=len(table.columns))
    item.autofit = False
    item.alignment = 1
    props = item._tbl.tblPr
    props.find(qn("w:tblW")).set(qn("w:w"), str(round(sum(widths) / 25.4 * 1440)))
    props.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single" if side in {"top", "bottom"} else "nil")
        border.set(qn("w:sz"), "8")
        borders.append(border)
    props.append(borders)
    for column, width in zip(item.columns, widths, strict=True):
        column.width = Mm(width)
    for row_index, values in enumerate([table.columns, *table.rows]):
        row = item.rows[0] if row_index == 0 else item.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for index, (cell, value, width) in enumerate(zip(row.cells, values, widths, strict=True)):
            cell.width = Mm(width)
            cell.text = value
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.keep_with_next = (
                compact and row_index < len(table.rows)
            ) or row_index == 0
            paragraph.alignment = 2 if index in table.numeric_columns else 0
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.space_before = Pt(3)
            for run in paragraph.runs:
                run.font.size = Pt(style.caption_pt)
                run.font.bold = row_index == 0
            if row_index == 0:
                cell_borders, bottom = OxmlElement("w:tcBorders"), OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                cell_borders.append(bottom)
                cell._tc.get_or_add_tcPr().append(cell_borders)
    if table.note:
        # Keep the note with the final data row, not on an otherwise empty page.
        for cell in item.rows[-1].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
        document.add_paragraph(table.note, style="Caption")
