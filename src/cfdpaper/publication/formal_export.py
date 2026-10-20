"""Formal Word, Markdown and LaTeX views of the same ordered manuscript blocks.

Export does not recalculate scientific facts or rewrite author text.  Managed Word
anchors identify editable blocks and protected bound spans; they do not lock text.
"""

from __future__ import annotations

import copy
import json
import re
import shutil
import tempfile
from difflib import SequenceMatcher
from pathlib import Path

from PIL import Image

from cfdpaper.publication.elements import (
    MathNode,
    SectionEquation,
    SectionTable,
    add_equation,
    add_table,
    display_scientific_text,
    math_xml,
)
from cfdpaper.publication.export import _latex_escape as _escape_tex_ascii
from cfdpaper.publication.export import _markdown_destination
from cfdpaper.publication.style import FigureSizing, PublicationStyle, figure_placement


def _latex_escape(value):
    """Keep common scientific Unicode readable under ordinary PDFLaTeX templates."""
    superscripts = str.maketrans("⁺⁻⁰¹²³⁴⁵⁶⁷⁸⁹", "+-0123456789")
    subscripts = str.maketrans("₊₋₀₁₂₃₄₅₆₇₈₉", "+-0123456789")
    replacements = {
        "−": r"\ensuremath{-}",
        "×": r"\ensuremath{\times}",
        "·": r"\ensuremath{\cdot}",
        "°": r"\textdegree{}",
        "–": "--",
        "—": "---",
        "\u00a0": "~",
        "\u2060": "",
        "μ": r"\ensuremath{\mu}",
        "µ": r"\ensuremath{\mu}",
        "Δ": r"\ensuremath{\Delta}",
        "α": r"\ensuremath{\alpha}",
        "β": r"\ensuremath{\beta}",
        "ρ": r"\ensuremath{\rho}",
        "λ": r"\ensuremath{\lambda}",
    }
    parts = re.split(r"([⁺⁻⁰¹²³⁴⁵⁶⁷⁸⁹]+|[₊₋₀₁₂₃₄₅₆₇₈₉]+)", value)
    result = []
    for part in parts:
        if part and part[0] in "⁺⁻⁰¹²³⁴⁵⁶⁷⁸⁹":
            result.append(r"\textsuperscript{" + part.translate(superscripts) + "}")
        elif part and part[0] in "₊₋₀₁₂₃₄₅₆₇₈₉":
            result.append(r"\textsubscript{" + part.translate(subscripts) + "}")
        else:
            result.extend(replacements.get(char, _escape_tex_ascii(char)) for char in part)
    return "".join(result)


def math_latex(value: MathNode | dict) -> str:
    """Render the existing bounded math tree as genuine TeX, not escaped prose."""
    node = value if isinstance(value, MathNode) else MathNode.model_validate(value)
    if node.kind == "text":
        return r"\text{" + _latex_escape(node.text) + "}"
    if node.kind == "symbol":
        symbols = {
            "α": r"\alpha",
            "β": r"\beta",
            "γ": r"\gamma",
            "δ": r"\delta",
            "Δ": r"\Delta",
            "ε": r"\varepsilon",
            "θ": r"\theta",
            "λ": r"\lambda",
            "μ": r"\mu",
            "ν": r"\nu",
            "π": r"\pi",
            "ρ": r"\rho",
            "σ": r"\sigma",
            "τ": r"\tau",
            "φ": r"\phi",
            "ω": r"\omega",
            "Ω": r"\Omega",
            "∂": r"\partial",
            "∇": r"\nabla",
            "∞": r"\infty",
            "∑": r"\sum",
            "∫": r"\int",
            "×": r"\times",
            "·": r"\cdot",
            "−": "-",
            "≤": r"\leq",
            "≥": r"\geq",
            "≠": r"\neq",
            "≈": r"\approx",
        }
        return " ".join(symbols.get(char, _latex_escape(char)) for char in node.text)
    values = [math_latex(child) for child in node.children]
    if node.kind == "row":
        return " ".join(values)
    if node.kind == "fraction":
        return rf"\frac{{{values[0]}}}{{{values[1]}}}"
    if node.kind == "sqrt":
        return rf"\sqrt{{{values[0]}}}"
    if node.kind in {"dot", "overbar"}:
        command = "dot" if node.kind == "dot" else "overline"
        return f"\\{command}{{{values[0]}}}"
    if node.kind == "subsup":
        return f"{{{values[0]}}}_{{{values[1]}}}^{{{values[2]}}}"
    return f"{{{values[0]}}}" + ("_" if node.kind == "sub" else "^") + f"{{{values[1]}}}"


def _math_leaves(value):
    if isinstance(value, MathNode):
        value = value.model_dump()
    if value["kind"] in {"symbol", "text"}:
        yield value
    else:
        for child in value.get("children", []):
            yield from _math_leaves(child)


def _content(block):
    return block.get("content", {})


def _runs(block):
    content = _content(block)
    if block["kind"] == "reference":
        if content.get("formatted_runs"):
            return content["formatted_runs"]
        if content.get("source"):
            from cfdpaper.publication.section import reference_source_suffix

            return [{"text": _plain(block) + reference_source_suffix(content)}]
    runs = (
        content.get("semantic_runs")
        or content.get("runs")
        or [{"text": block.get("rendered_text", content.get("text", ""))}]
    )
    label = block.get("reference_labels", {}).get("figure")
    if label is None:
        return runs
    if label not in {"Figure", "Fig."}:
        raise ValueError("Figure reference label must be Figure or Fig.")
    return [
        {**run, "text": re.sub(r"^(?:Figure|Fig\.)\s+", label + " ", run["text"])}
        if run.get("token", "").startswith("{{figure:") and "text" in run
        else run
        for run in runs
    ]


def _plain(block):
    return block.get("rendered_text", _content(block).get("text", ""))


def _prefix(block):
    content = _content(block)
    if block["kind"] == "caption":
        label = block.get("reference_labels", {}).get("figure", "Figure")
        if label not in {"Figure", "Fig."}:
            raise ValueError("Figure reference label must be Figure or Fig.")
        return f"{label} {content['id']}. "
    if block["kind"] == "reference":
        return f"[{content['number']}] "
    if block["kind"] == "keywords":
        return "Keywords: "
    return ""


def _math_prose(block, *, latex=False):
    parts = []
    for run in _runs(block):
        if "math" in run:
            parts.append("$" + math_latex(run["math"]) + "$")
        else:
            value = _latex_escape(run["text"]) if latex else run["text"]
            for prop, command, marker in (("italic", "textit", "*"), ("bold", "textbf", "**")):
                if run.get(prop):
                    value = f"\\{command}{{{value}}}" if latex else marker + value + marker
            parts.append(value)
    return "".join(parts)


def _model_content(model, content):
    return model.model_validate(
        {key: value for key, value in content.items() if key in model.model_fields}
    )


def _profile_style(document, template=None):
    style = PublicationStyle.model_validate(document.get("style", {}))
    if template is None:
        size = {"research": style.body_pt, "report": 12, "conference": 10}[document["profile"]]
        style = style.model_copy(update={"body_pt": size})
    return style


def _copy_assets(blocks, source, stage, asset_name):
    assets, result = stage / asset_name, {}
    for block in blocks:
        if block["kind"] != "figure":
            continue
        item = _content(block)
        path = Path(item["path"])
        path = path if path.is_absolute() else source / path
        if not path.is_file() or not path.stat().st_size:
            raise FileNotFoundError(f"Figure {block['id']} is missing or empty: {path}")
        suffix = path.suffix.lower()
        if suffix not in {".png", ".jpg", ".jpeg", ".pdf", ".svg", ".eps"}:
            raise ValueError(f"Unsupported portable figure format: {suffix}")
        assets.mkdir(exist_ok=True)
        name = f"figure-{len(result) + 1}{suffix}"
        shutil.copy2(path, assets / name)
        result[block["id"]] = f"{asset_name}/{name}"
    return result


def _markdown(document, assets):
    lines = [f"# {document['title']}", ""]
    for block in document["blocks"]:
        kind, content = block["kind"], _content(block)
        if kind == "heading":
            lines.append("#" * (content.get("level", 1) + 1) + " " + _plain(block))
        elif kind == "figure":
            lines.append(f"![]({_markdown_destination(assets[block['id']])})")
        elif kind == "equation":
            lines.append("$$\n" + math_latex(content["expression"]) + "\n$$")
        elif kind == "table":

            def cell(value):
                return str(value).replace("|", "\\|").replace("\n", " ")

            rows = [content["columns"], ["---"] * len(content["columns"]), *content["rows"]]
            lines.append(f"Table {content['table_id']}. {content['caption']}")
            lines.append("\n".join("| " + " | ".join(map(cell, row)) + " |" for row in rows))
            if content.get("note"):
                lines.append(content["note"])
        elif kind in {"paragraph", "caption", "keywords", "reference"}:
            lines.append(_prefix(block) + _math_prose(block))
        else:
            raise ValueError(f"Unsupported manuscript block kind: {kind}")
        lines.append("")
    return "\n".join(lines)


def _latex(document, assets, template=None):
    profile = document["profile"]
    options = (
        "10pt,twocolumn" if profile == "conference" else ("12pt" if profile == "report" else "11pt")
    )
    body = []
    figure_environment = "figure*" if profile == "conference" else "figure"
    table_environment = "table*" if profile == "conference" else "table"
    for block in document["blocks"]:
        kind, content = block["kind"], _content(block)
        if kind == "heading":
            command = "subsection" if content.get("level", 1) == 2 else "section"
            body.append(f"\\{command}*{{{_latex_escape(_plain(block))}}}")
        elif kind == "figure":
            path = assets[block["id"]]
            if Path(path).suffix.lower() == ".svg":
                raise ValueError("LaTeX export needs a PNG, JPEG, PDF or EPS figure, not SVG")
            body.append(
                "\n".join(
                    [
                        "\\begin{" + figure_environment + "}[htbp]",
                        r"\centering",
                        r"\includegraphics[width=\linewidth]{\detokenize{" + path + "}}",
                    ]
                )
            )
        elif kind == "equation":
            body.append(
                "\n".join(
                    [
                        r"\begin{equation}",
                        math_latex(content["expression"]),
                        r"\tag{" + _latex_escape(content["equation_id"]) + "}",
                        r"\end{equation}",
                    ]
                )
            )
        elif kind == "table":
            count = len(content["columns"])
            alignment = "".join(
                r">{\raggedleft\arraybackslash}X"
                if i in content.get("numeric_columns", [])
                else r">{\raggedright\arraybackslash}X"
                for i in range(count)
            )
            rows = [
                " & ".join(_latex_escape(str(v)) for v in row) + r" \\"
                for row in [content["columns"], *content["rows"]]
            ]
            body.append(
                "\n".join(
                    [
                        "\\begin{" + table_environment + "}[htbp]",
                        r"\centering",
                        r"\caption*{Table "
                        + _latex_escape(content["table_id"])
                        + ". "
                        + _latex_escape(content["caption"])
                        + "}",
                        f"\\begin{{tabularx}}{{\\linewidth}}{{{alignment}}}",
                        r"\toprule",
                        rows[0],
                        r"\midrule",
                        *rows[1:],
                        r"\bottomrule",
                        r"\end{tabularx}",
                        r"\par\smallskip\noindent\small "
                        + _latex_escape(content.get("note", ""))
                        + r"\par",
                        "\\end{" + table_environment + "}",
                    ]
                )
            )
        elif kind in {"paragraph", "caption", "keywords", "reference"}:
            if kind == "caption":
                body.append(
                    r"\caption*{"
                    + _latex_escape(_prefix(block))
                    + _math_prose(block, latex=True)
                    + "}\n\\end{"
                    + figure_environment
                    + "}"
                )
            else:
                body.append(_latex_escape(_prefix(block)) + _math_prose(block, latex=True))
        else:
            raise ValueError(f"Unsupported manuscript block kind: {kind}")
    body_text = "\n\n".join(body)
    if template:
        text = Path(template).read_text(encoding="utf-8")
        if text.count("%%CFDPAPER_BODY%%") != 1:
            raise ValueError("LaTeX template must contain one %%CFDPAPER_BODY%% marker")
        return text.replace("%%CFDPAPER_TITLE%%", _latex_escape(document["title"])).replace(
            "%%CFDPAPER_BODY%%", body_text
        )
    return "\n".join(
        [
            f"\\documentclass[{options}]{{article}}",
            r"\usepackage[utf8]{inputenc}",
            r"\usepackage[T1]{fontenc}",
            r"\usepackage{amsmath,amssymb,graphicx,booktabs,tabularx,caption,textcomp}",
            r"\usepackage[a4paper,margin=25mm]{geometry}",
            r"\usepackage{indentfirst}",
            r"\setlength{\parindent}{2em}",
            r"\setlength{\parskip}{0pt}",
            r"\title{" + _latex_escape(document["title"]) + "}",
            r"\date{}",
            r"\begin{document}",
            r"\maketitle",
            body_text,
            r"\end{document}",
            "",
        ]
    )


def _set_styles(document, config, template):
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    if template:
        for child in list(document._element.body):
            if child.tag != qn("w:sectPr"):
                document._element.body.remove(child)
        return
    for name, size in (
        ("Normal", config.body_pt),
        ("Title", config.title_pt),
        ("Caption", config.caption_pt),
        ("Heading 1", config.heading_pt),
        ("Heading 2", config.heading_pt),
    ):
        style = document.styles[name]
        style.font.name, style.font.size = config.font_family, Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in list(fonts.attrib):
            if "theme" in attribute.lower():
                del fonts.attrib[attribute]
        for key in ("ascii", "hAnsi", "eastAsia"):
            fonts.set(qn("w:" + key), config.font_family)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        style.paragraph_format.widow_control = True
    normal = document.styles["Normal"].paragraph_format
    normal.first_line_indent = Pt(config.body_first_line_indent_chars * config.body_pt)
    normal.space_before, normal.space_after = Pt(0), Pt(0)
    normal.line_spacing = config.line_spacing
    document.styles["Normal"].element.get_or_add_pPr().get_or_add_ind().set(
        qn("w:firstLineChars"), str(round(config.body_first_line_indent_chars * 100))
    )
    for name in ("Title", "Caption", "Heading 1", "Heading 2"):
        item = document.styles[name].paragraph_format
        item.first_line_indent = Pt(0)
        document.styles[name].element.get_or_add_pPr().get_or_add_ind().set(
            qn("w:firstLineChars"), "0"
        )
        item.space_before, item.space_after = Pt(6), Pt(4)
        if name.startswith("Heading"):
            item.keep_with_next = True
    document.styles["Caption"].font.bold = document.styles["Caption"].font.italic = False
    for page in document.sections:
        page.page_width, page.page_height = Mm(config.page_width_mm), Mm(config.page_height_mm)
        page.left_margin = page.right_margin = Mm(config.margin_mm)
        page.top_margin = page.bottom_margin = Mm(config.margin_mm)


def _word_style_roles(document, template):
    """Use the retained template's existing body/title/caption roles."""
    from docx.enum.style import WD_STYLE_TYPE

    names = {style.name.casefold(): style for style in document.styles}
    roles = {}
    for role, aliases in {
        "Normal": ("Normal",),
        "Title": ("Title", "paper title"),
        "Caption": ("Caption", "figure caption"),
        "Heading 1": ("Heading 1",),
    }.items():
        match = next((names[name.casefold()] for name in aliases if name.casefold() in names), None)
        if match is None:
            raise ValueError(f"Word template is missing its {role} paragraph style")
        roles[role] = match
    if template and "body text" in names:
        body = names["body text"]
        if any(p.style == body for p in document.paragraphs):
            roles["Normal"] = body
    # Existing table and returned-paragraph helpers use these semantic aliases.
    # Derive missing aliases from the original style rather than replacing it.
    for role in ("Title", "Caption"):
        if role.casefold() not in names:
            alias = document.styles.add_style(role, WD_STYLE_TYPE.PARAGRAPH)
            alias.base_style = roles[role]
    return roles


def _style_size(style):
    """Resolve an inherited font size without modifying a template style."""
    visited = set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        if style.font.size is not None:
            return style.font.size.pt
        style = style.base_style
    return None


def _paragraph_overrides(paragraph, values, style_roles=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    if isinstance(values, list):
        props = paragraph._p.get_or_add_pPr()
        for record in values[:1]:
            for name, attrs in record.items():
                if name not in {"pStyle", "jc", "ind", "spacing", "keepNext", "keepLines"}:
                    continue
                if name == "pStyle" and "name" in attrs:
                    paragraph.style = (style_roles or {}).get(attrs["name"], attrs["name"])
                    continue
                old = props.find(qn("w:" + name))
                if old is not None:
                    props.remove(old)
                element = OxmlElement("w:" + name)
                for key, value in attrs.items():
                    element.set(qn("w:" + key), str(value))
                props.append(element)
        return
    allowed = {
        "first_line_indent_pt",
        "left_indent_pt",
        "right_indent_pt",
        "space_before_pt",
        "space_after_pt",
    }
    for key in allowed & values.keys():
        setattr(paragraph.paragraph_format, key.removesuffix("_pt"), Pt(values[key]))
    for key in (
        "alignment",
        "line_spacing",
        "keep_with_next",
        "keep_together",
        "page_break_before",
        "widow_control",
    ):
        if key in values:
            target = paragraph if key == "alignment" else paragraph.paragraph_format
            setattr(target, key, values[key])


def _word_column_width(document):
    """Respect actual template columns rather than treating every page as one column."""
    from docx.oxml.ns import qn

    page = document.sections[-1]
    available = page.page_width.mm - page.left_margin.mm - page.right_margin.mm
    columns = page._sectPr.find(qn("w:cols"))
    if columns is None:
        return available
    explicit = columns.findall(qn("w:col"))
    if explicit and explicit[0].get(qn("w:w")):
        return int(explicit[0].get(qn("w:w"))) * 25.4 / 1440
    count = int(columns.get(qn("w:num"), "1"))
    gap = int(columns.get(qn("w:space"), "720")) * 25.4 / 1440
    return (available - (count - 1) * gap) / count


def _add_text(paragraph, text, formatting, offset, *, reference_size=None, defaults=None):
    """Apply returned character ranges without destroying protected span boundaries."""
    from docx.shared import Pt

    boundaries = {0, len(text)}
    for item in formatting:
        for key in ("start", "end"):
            boundary = item[key] - offset
            if 0 < boundary < len(text):
                boundaries.add(boundary)
    edges = sorted(boundaries)
    for start, end in zip(edges, edges[1:], strict=False):
        run = paragraph.add_run(text[start:end])
        props = {
            **(defaults or {}),
            **next((p for p in formatting if p["start"] <= offset + start < p["end"]), {}),
        }
        for name in ("bold", "italic", "superscript", "subscript"):
            if name in props:
                setattr(run.font, name, props[name])
        if reference_size is not None:
            run.font.size = Pt(reference_size)


def _rebase_formatting(source, target, records):
    """Move display ranges only when their boundaries have an unambiguous mapping."""
    if not source or source == target:
        return records, []
    opcodes = SequenceMatcher(a=source, b=target, autojunk=False).get_opcodes()

    def boundary(position, *, end=False):
        # Exact replacement endpoints are safe; partial styling within changed
        # scientific text is not guessed from nearby digits.
        for tag, a, b, c, d in opcodes:
            if tag == "equal" and a <= position <= b:
                return c + position - a
            if tag != "equal":
                if position == a:
                    return d if tag == "insert" and end else c
                if position == b:
                    return d
                if a < position < b:
                    return None
        return len(target) if position == len(source) else None

    result, conflicts = [], []
    for record in records:
        start, end = boundary(record["start"]), boundary(record["end"], end=True)
        if start is None or end is None or end < start:
            if any(key in record for key in ("bold", "italic", "superscript", "subscript")):
                conflicts.append({"type": "formatting_needs_relocation", "range": record})
            continue
        result.append({**record, "start": start, "end": end})
    return result, conflicts


def _add_comments(document, paragraphs):
    """Restore ordinary anchored comments without importing stale comment identifiers."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree

    root = etree.Element(qn("w:comments"), nsmap={"w": qn("w:comments").split("}")[0][1:]})
    index, conflicts = 0, []
    for paragraph, block in paragraphs:
        for comment in block.get("comments", []):

            def positions(paragraph=paragraph):
                rows, offset = [], 0
                for node in paragraph._p.iter():
                    if node.tag == qn("w:r"):
                        text = "".join(t.text or "" for t in node.iter(qn("w:t")))
                        if text:
                            rows.append((node, offset, offset + len(text)))
                    elif node.tag in {qn("w:t"), qn("m:t")}:
                        offset += len(node.text or "")
                return rows

            def split_at(boundary):
                for run, a, b in positions():
                    if not a < boundary < b:
                        continue
                    if any(node.tag not in {qn("w:rPr"), qn("w:t")} for node in run):
                        return
                    text = "".join(t.text or "" for t in run.iter(qn("w:t")))
                    right = copy.deepcopy(run)
                    for element, value in (
                        (run, text[: boundary - a]),
                        (right, text[boundary - a :]),
                    ):
                        for child in list(element):
                            if child.tag != qn("w:rPr"):
                                element.remove(child)
                        child = OxmlElement("w:t")
                        child.set(qn("xml:space"), "preserve")
                        child.text = value
                        element.append(child)
                    run.addnext(right)
                    return

            start, end = comment.get("start", 0), comment.get("end", 0)
            visible = "".join(
                node.text or ""
                for node in paragraph._p.iter()
                if node.tag in {qn("w:t"), qn("m:t")}
            )
            quote = comment.get("anchor_text", "")
            if quote and visible.count(quote) == 1:
                start, end = visible.index(quote), visible.index(quote) + len(quote)
            elif quote and visible[start:end] != quote:
                start, end = 0, len(visible)
                conflicts.append(
                    {
                        "type": "comment_needs_relocation",
                        "block_id": block["id"],
                        "comment": comment,
                    }
                )
            split_at(end)
            split_at(start)
            rows = positions()
            selected = [run for run, a, b in rows if b > start and a < end]
            if not selected:
                selected = [rows[0][0]] if rows else []
            if not selected:
                selected = [paragraph.add_run("")._r]
            begin, finish, reference = (
                OxmlElement("w:" + tag)
                for tag in ("commentRangeStart", "commentRangeEnd", "commentReference")
            )
            for element in (begin, finish, reference):
                element.set(qn("w:id"), str(index))
            selected[0].addprevious(begin)
            selected[-1].addnext(finish)
            reference_run = OxmlElement("w:r")
            reference_run.append(reference)
            finish.addnext(reference_run)
            item = etree.SubElement(root, qn("w:comment"))
            item.set(qn("w:id"), str(index))
            for key in ("author", "date"):
                if comment.get(key):
                    item.set(qn("w:" + key), str(comment[key]))
            p, r, t = (OxmlElement("w:" + tag) for tag in ("p", "r", "t"))
            t.text = comment.get("text", "")
            r.append(t)
            p.append(r)
            item.append(p)
            index += 1
    if index:
        blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
        part = Part(
            PackURI("/word/comments.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
            blob,
            document.part.package,
        )
        for rid, relation in list(document.part.rels.items()):
            if relation.reltype == RT.COMMENTS:
                document.part.drop_rel(rid)
        document.part.relate_to(part, RT.COMMENTS)
    return conflicts


def _docx(document_data, source, output, template=None):
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt
    except ImportError as exc:
        raise RuntimeError("Formal Word export requires cfd-paper-agent[docs]") from exc
    from cfdpaper.publication.word_return import (
        end_anchor,
        start_anchor,
        wrap_anchor,
        write_word_state,
    )

    document = Document(str(template)) if template else Document()
    config = _profile_style(document_data, template)
    style_roles = _word_style_roles(document, template)
    # Keep profile outside the python-docx object, whose implementation may use slots.
    _set_styles(document, config, template)
    if template:
        # Word 2003 compatibility removes managed rich-text controls on save.
        # Upgrade only this generated view, leaving the retained template intact.
        for setting in document.settings.element.findall(".//" + qn("w:compatSetting")):
            if setting.get(qn("w:name")) == "compatibilityMode":
                if int(setting.get(qn("w:val"), "15")) < 15:
                    setting.set(qn("w:val"), "15")
        page = document.sections[0]
        config = config.model_copy(
            update={
                "page_width_mm": page.page_width.mm,
                "page_height_mm": page.page_height.mm,
                "margin_mm": max(page.left_margin.mm, page.right_margin.mm),
                **{
                    key: _style_size(style_roles[name])
                    for key, name in (
                        ("body_pt", "Normal"),
                        ("title_pt", "Title"),
                        ("heading_pt", "Heading 1"),
                        ("caption_pt", "Caption"),
                    )
                    if _style_size(style_roles[name]) is not None
                },
            }
        )
    column_width = _word_column_width(document)
    document.add_paragraph(document_data["title"], style_roles["Title"])
    state = copy.deepcopy(document_data)
    placements = []
    formatting_conflicts = []
    comment_paragraphs = []
    image_paragraph = None
    references_started = False
    for block in state["blocks"]:
        kind, content = block["kind"], _content(block)
        before_nodes = set(document._element.body)
        if kind == "table":
            table = _model_content(SectionTable, content)
            if table.column_widths_mm and sum(table.column_widths_mm) > column_width:
                factor = column_width / sum(table.column_widths_mm)
                table = table.model_copy(
                    update={"column_widths_mm": [w * factor for w in table.column_widths_mm]}
                )
            table_style = config.model_copy(
                update={"page_width_mm": column_width + 2 * config.margin_mm}
            )
            add_table(document, table, table_style)
            nodes = [node for node in document._element.body if node not in before_nodes]
            for node in nodes:
                for paragraph in node.iter(qn("w:p")):
                    props = paragraph.get_or_add_pPr()
                    ind = props.get_or_add_ind()
                    ind.set(qn("w:firstLine"), "0")
                    ind.set(qn("w:firstLineChars"), "0")
            wrap_anchor(nodes[0], block["id"])
            container = nodes[0].getparent()
            for node in nodes[1:]:
                container.append(node)
            continue
        if kind == "equation":
            add_equation(document, _model_content(SectionEquation, content))
            paragraph = document.paragraphs[-1]
        elif kind == "figure":
            path = Path(content["path"])
            path = path if path.is_absolute() else source / path
            with Image.open(path) as image:
                figure_config = config.model_copy(
                    update={
                        "page_width_mm": column_width + 2 * config.margin_mm,
                        "figure_width_mm": min(config.figure_width_mm, column_width),
                    }
                )
                sizing = {**content.get("sizing", {})}
                returned_layout = block.get("layout", {})
                if isinstance(returned_layout, list) and returned_layout:
                    sizing["target_width_mm"] = returned_layout[0]["cx"] / 36000
                elif returned_layout.get("width_mm"):
                    sizing["target_width_mm"] = returned_layout["width_mm"]
                placement = figure_placement(
                    pixels=image.size,
                    caption=content.get("caption", ""),
                    sizing=FigureSizing.model_validate(sizing),
                    style=figure_config,
                )
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.keep_with_next = True
            returned_layout = block.get("layout", {})
            paragraph.alignment = (
                returned_layout.get("alignment", 1) if isinstance(returned_layout, dict) else 1
            )
            paragraph.add_run().add_picture(
                str(path), width=Mm(placement["width_mm"]), height=Mm(placement["height_mm"])
            )
            placements.append({"id": block["id"], **placement})
            image_paragraph = paragraph
        elif kind == "heading":
            paragraph = document.add_heading(_plain(block), level=content.get("level", 1))
        else:
            if kind == "reference" and not references_started:
                document.add_heading("References", level=1)
                references_started = True
            paragraph = document.add_paragraph(
                style=style_roles["Caption" if kind == "caption" else "Normal"]
            )
            if kind in {"caption", "keywords", "reference"}:
                if not template:
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "0")
            prefix = _prefix(block)
            if prefix:
                paragraph.add_run(prefix)
            start_anchor(paragraph, block["id"])
            span_by_id = {span["id"]: span for span in block.get("spans", [])}
            formatting = block.get("formatting", [])
            runs = _runs(block)
            display = prefix + "".join(
                "".join(node["text"] for node in _math_leaves(run["math"]))
                if "math" in run
                else display_scientific_text(run.get("text", ""))
                for run in runs
            )
            formatting, problems = _rebase_formatting(
                content.get("formatting_source_text", block.get("formatting_source_text")),
                display,
                formatting,
            )
            formatting_conflicts.extend(
                {**problem, "block_id": block["id"]} for problem in problems
            )
            if formatting:
                block["formatting"] = formatting
                block["formatting_source_text"] = display
            visible = []
            offset = len(prefix)
            for run in runs:
                span_id = run.get("span_id")
                key = block["id"] + "::" + span_id if span_id else None
                if key:
                    start_anchor(paragraph, key)
                if "math" in run:
                    formula = OxmlElement("m:oMath")
                    formula.extend(math_xml(MathNode.model_validate(run["math"])))
                    paragraph._p.append(formula)
                    text = "".join(t.text or "" for t in formula.iter(qn("m:t")))
                else:
                    text = display_scientific_text(run.get("text", ""))
                    _add_text(
                        paragraph,
                        text,
                        formatting,
                        offset,
                        reference_size=config.reference_pt if kind == "reference" else None,
                        defaults=run,
                    )
                offset += len(text)
                visible.append(text)
                if key:
                    end_anchor(paragraph, key)
                    if span_id in span_by_id:
                        span_by_id[span_id]["text"] = text
            end_anchor(paragraph, block["id"])
            block["rendered_text"] = "".join(visible)
            _paragraph_overrides(paragraph, block.get("paragraph_format", {}), style_roles)
            comment_paragraphs.append((paragraph, block))
            if kind == "caption":
                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.keep_together = (
                    config.figure_caption_pagination == "keep"
                )
                if image_paragraph is not None:
                    image_paragraph.paragraph_format.keep_with_next = True
                image_paragraph = None
            continue
        start_anchor(paragraph, block["id"])
        end_anchor(paragraph, block["id"])
        _paragraph_overrides(paragraph, block.get("paragraph_format", {}), style_roles)
        comment_paragraphs.append((paragraph, block))
    # Captions already have stable scientific labels. Template list numbering
    # must not add a second Figure number or number an ordinary table note.
    caption_styles = {style_roles["Caption"].style_id, document.styles["Caption"].style_id}
    for element in document._element.body.iter(qn("w:p")):
        props = element.find(qn("w:pPr"))
        paragraph_style = None if props is None else props.find(qn("w:pStyle"))
        if paragraph_style is not None and paragraph_style.get(qn("w:val")) in caption_styles:
            numbers = props.get_or_add_numPr()
            numbers.get_or_add_numId().val = 0
    comment_conflicts = _add_comments(document, comment_paragraphs)
    write_word_state(document, state)
    document.save(output)
    return {
        "style": config.model_dump(),
        "profile": document_data["profile"],
        "template_style_roles": {key: value.name for key, value in style_roles.items()},
        "figures": placements,
        "column_width_mm": column_width,
        "formatting_conflicts": formatting_conflicts,
        "comment_conflicts": comment_conflicts,
        "unmatched_comments": document_data.get("unmatched_comments", []),
    }


def export_formal_manuscript(
    section_dir: Path,
    output_path: Path,
    *,
    profile="research",
    layout="near-reference",
    template_path: Path | None = None,
) -> Path:
    """Export a fresh formal view; never overwrite a template or an author-owned file."""
    from cfdpaper.publication.document_blocks import (
        apply_author_edits,
        build_document,
        source_definition_state,
    )

    source, output = Path(section_dir), Path(output_path)
    suffix = output.suffix.lower()
    if suffix not in {".docx", ".md", ".tex"}:
        raise ValueError("Formal manuscript output must end in .docx, .md or .tex")
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite author-owned output: {output}")
    template = Path(template_path) if template_path is not None else None
    if template and (not template.is_file() or template.suffix.lower() != suffix):
        raise ValueError("Template must be an existing file with the same output format")
    if template and suffix == ".md":
        raise ValueError("Explicit templates are supported for Word and LaTeX output")
    data = json.loads((source / "section.json").read_text(encoding="utf-8"))
    document = build_document(data, profile=profile, layout=layout)
    if (source / "author-edits.json").is_file():
        edits = json.loads((source / "author-edits.json").read_text(encoding="utf-8"))
        apply_author_edits(document, edits)
    if (source / "writing-state.json").is_file():
        document["source_state"] = source_definition_state(
            json.loads((source / "writing-state.json").read_text(encoding="utf-8"))
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    assets = output.with_name(output.stem + "_" + suffix[1:] + "_assets")
    if suffix != ".docx" and assets.exists():
        raise FileExistsError(f"Refusing to overwrite author-owned assets: {assets}")
    report = output.with_suffix(".layout.json")
    if suffix == ".docx" and report.exists():
        raise FileExistsError(f"Refusing to overwrite author-owned layout: {report}")
    with tempfile.TemporaryDirectory(prefix=".formal-", dir=output.parent) as temp:
        stage = Path(temp)
        target = stage / output.name
        if suffix == ".docx":
            result = _docx(document, source, target, template)
            (stage / report.name).write_text(
                json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
            )
        else:
            copied = _copy_assets(document["blocks"], source, stage, assets.name)
            text = (
                _markdown(document, copied)
                if suffix == ".md"
                else _latex(document, copied, template)
            )
            target.write_text(text, encoding="utf-8")
        # Recheck at publication rather than trusting the preflight check.
        if output.exists() or (suffix != ".docx" and assets.exists()):
            raise FileExistsError(f"Refusing to overwrite author-owned output: {output}")
        if (stage / assets.name).exists():
            (stage / assets.name).rename(assets)
        if suffix == ".docx":
            (stage / report.name).rename(report)
        target.rename(output)
    return output
