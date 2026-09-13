import json
import runpy
from pathlib import Path

import pytest
from typer.testing import CliRunner

from cfdpaper.cli import app

runner = CliRunner()


@pytest.mark.parametrize(
    "options,detail",
    [
        ([], "Choose one"),
        (["--manuscript-input", "input.json"], "--output"),
        (["--draft", "drafts.json", "--output", "new"], "--package"),
        (["--approve-final", "--author", "Author"], "author review"),
        (["--section-input", "input.json"], "--manuscript-input"),
        (["--pdf-preview"], "--docx"),
    ],
)
def test_manuscript_action_errors(tmp_path, options, detail):
    result = runner.invoke(app, ["write", str(tmp_path), "--artifact", "manuscript", *options])
    assert result.exit_code != 0
    assert detail in result.stdout + result.stderr


@pytest.mark.parametrize("shared_literature", [False, True])
def test_three_section_cli_and_docx_preserve_structure_and_body_style(tmp_path, shared_literature):
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn

    script = (
        Path(__file__).resolve().parents[1] / "examples/manuscript-workspace/prepare_example.py"
    )
    source, package, assembled = (tmp_path / name for name in ("source", "package", "assembled"))
    runpy.run_path(str(script))["prepare"](source)
    if shared_literature:
        # A supplied analytical note, explicitly not a published scientific reference.
        (source / "excerpts.txt").write_text("The analytical example uses constant properties.")
        (source / "references.json").write_text(
            json.dumps(
                [{"id": "fixture-note", "type": "report", "title": "Analytical example note"}]
            )
        )
        manifest_path = source / "manuscript-input.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        supports = []
        drafts = json.loads((source / "drafts.json").read_text(encoding="utf-8"))
        for sid, draft_name in drafts.items():
            supports.append(
                {
                    "section_id": sid,
                    "evidence_id": "shared-note",
                    "reference_id": "fixture-note",
                    "source": "excerpts.txt",
                    "locator": "paragraph 1",
                    "excerpt": "The analytical example uses constant properties.",
                    "claim": "Constant properties in the supplied analytical example",
                    "role": "method basis",
                    "status": "supported",
                }
            )
            path = source / draft_name
            draft = json.loads(path.read_text(encoding="utf-8"))
            draft["paragraphs"][0]["text"] += (
                " The example assumes constant properties {{cite:shared-note}}."
            )
            draft["paragraphs"][0]["evidence_ids"].append("shared-note")
            path.write_text(json.dumps(draft), encoding="utf-8")
        (source / "literature.json").write_text(
            json.dumps({"bibliography": "references.json", "supports": supports})
        )
        manifest["literature"] = "literature.json"
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    base = ["write", str(source), "--artifact", "manuscript"]
    for options in (
        ["--manuscript-input", str(source / "manuscript-input.json"), "--output", str(package)],
        [
            "--package",
            str(package),
            "--draft",
            str(source / "drafts.json"),
            "--output",
            str(assembled),
        ],
        [
            "--package",
            str(assembled),
            "--docx",
            "--layout",
            "near-reference",
            "--output",
            str(tmp_path / "manuscript.docx"),
        ],
    ):
        result = runner.invoke(app, base + options)
        assert result.exit_code == 0, result.stdout + result.stderr
    data = json.loads((assembled / "section.json").read_text(encoding="utf-8"))
    assert sum("section_heading" in p for p in data["paragraphs"]) == 3
    doc = docx.Document(tmp_path / "manuscript.docx")
    assert len([p for p in doc.paragraphs if p.style.name == "Heading 1"]) == 3
    assert len(doc.tables) == 1
    assert len(doc.inline_shapes) == 2
    elements = list(doc.element.body)
    next_heading = next(p._p for p in doc.paragraphs if p.text == "2. Hydraulic response")
    assert elements.index(doc.tables[0]._tbl) < elements.index(next_heading)
    text = "\n".join(p.text for p in doc.paragraphs)
    if shared_literature:
        shared = [r for r in data["references"] if r["source"] == "reference:fixture-note"]
        assert len(shared) == 1
        assert text.count("Analytical example note") == 1
        assert text.count(f"constant properties [{shared[0]['number']}]") == 3
    assert "6.4 Pa" in text and "19.2 Pa" in text
    assert "Figure 1" in text and "Figure 2" in text
    body = next(p for p in doc.paragraphs if p.text.startswith("The reference considers"))
    assert body.paragraph_format.space_before.pt == 0
    assert body.paragraph_format.space_after.pt == 0
    assert body._p.pPr.ind.get(qn("w:firstLineChars")) == "200"
    heading = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
    assert not heading._p.xpath("./w:pPr/w:ind[@w:firstLineChars]")
