"""The public synthetic example uses the product's connected research workflow."""

import json
import runpy
import shutil
from pathlib import Path

import numpy as np
import pytest
from docx import Document

from cfdpaper.publication.research_workspace import build_research_workspace


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def scalar(data, section, identity):
    key = json.dumps([section, identity], separators=(",", ":"))
    return data["resolved_values"][key]["raw_value"]


@pytest.fixture(scope="module")
def example(tmp_path_factory):
    path = Path(__file__).resolve().parents[1] / "examples/surface-research/run_example.py"
    module = runpy.run_path(str(path))
    root = tmp_path_factory.mktemp("surface-example") / "demo"
    return root, module["run"](root)


def test_example_builds_original_faces_complementary_figures_and_bound_docx(example):
    root, candidate = example
    assert candidate == root / "candidate"
    results = candidate / "workspace/sections/results"
    assembled = read(candidate / "manuscript/section.json")
    for case in ("A", "B"):
        source = root / f"materials/sources/{case}.npz"
        assert source.read_bytes() == (results / "sources" / source.name).read_bytes()
        with np.load(source, allow_pickle=False) as arrays:
            values, weights = arrays["face_value"], arrays["area_m2"]
            assert values.shape == weights.shape == (6,)
            assert arrays["vertices_m"].shape == (12, 3)
            np.testing.assert_array_equal(arrays["face_starts"], np.arange(6) * 5)
            np.testing.assert_allclose(weights, np.array([1, 2, 1, 1, 2, 1]) * 1e-4)
            expected = np.average(values, weights=weights)
            assert scalar(assembled, "results", f"{case}-mean") == pytest.approx(expected)
            assert scalar(assembled, "conclusion", f"{case}-mean") == pytest.approx(expected)
    assert scalar(assembled, "results", "A-exceedance-1") == 0.5
    assert scalar(assembled, "results", "B-exceedance-1") == 0.125
    assert scalar(assembled, "results", "A-exceedance-2") == 0
    assert scalar(assembled, "results", "B-exceedance-2") == 0.125
    for figure in ("field", "distribution"):
        for extension in ("svg", "pdf", "png", "tiff"):
            assert (results / f"figures/{figure}.{extension}").stat().st_size > 0
    assert (results / "figures/plot.py").is_file()
    document = Document(candidate / "manuscript.docx")
    assert len(document.inline_shapes) == 2
    assert len(document.tables) == 1
    assert len(document.tables[0].columns) == 5
    assert [p.text for p in document.paragraphs if p.style.name == "Heading 1"] == [
        "Methods",
        "Results",
        "Discussion",
        "Conclusions",
    ]
    text = "\n".join(p.text for p in document.paragraphs).replace("\xa0", " ")
    assert "synthetic" in text.lower()
    assert "{{" not in text
    assert "335.000 K" in text


@pytest.mark.parametrize("change", ["value", "threshold"])
def test_copied_example_refreshes_without_rewriting_author_drafts(example, tmp_path, change):
    root, candidate = example
    work = tmp_path / "relocated-workspace"
    shutil.copytree(candidate / "workspace", work)
    originals = {path: path.read_bytes() for path in (root / "materials/sources").iterdir()}
    drafts = {name: (work / name).read_bytes() for name in read(work / "drafts.json").values()}
    script = work / "sections/results/figures/plot.py"
    script.write_text(
        script.read_text(encoding="utf-8").replace("'height_mm': 110", "'height_mm': 115"),
        encoding="utf-8",
    )
    author_script = script.read_bytes()
    if change == "value":
        source = work / "sections/results/sources/A.npz"
        with np.load(source, allow_pickle=False) as bundle:
            arrays = {key: bundle[key].copy() for key in bundle.files}
        arrays["face_value"] += 2.0
        np.savez(source, **arrays)
    else:
        study_path = work / "sections/results/study.json"
        study = read(study_path)
        study["thresholds"][0] = 340.0
        study_path.write_text(json.dumps(study), encoding="utf-8")
    refreshed = build_research_workspace(work, tmp_path / "updated", baseline_dir=candidate)
    assembled = read(refreshed / "manuscript/section.json")
    if change == "value":
        assert scalar(assembled, "results", "A-mean") == 337.0
        assert scalar(assembled, "conclusion", "A-mean") == 337.0
        original = read(candidate / "manuscript/section.json")
        assert scalar(assembled, "results", "A-std") == scalar(original, "results", "A-std")
    else:
        assert scalar(assembled, "results", "A-exceedance-1") == 0.125
        header = Document(refreshed / "manuscript.docx").tables[0].cell(0, 3).text
        assert "340.000 K" in header.replace("\xa0", " ")
        assert "thresholds" in read(refreshed / "refresh.json")["definition_changes"]
    assert read(refreshed / "refresh.json")["interpretation_review_required"] is True
    assert read(refreshed / "refresh.json")["figure_recipes"][0]["status"] == "executed"
    assert (refreshed / "workspace/sections/results/figures/plot.py").read_bytes() == author_script
    for name, content in drafts.items():
        assert (refreshed / "workspace" / name).read_bytes() == content
    assert all(path.read_bytes() == content for path, content in originals.items())


def test_example_refuses_existing_output(example):
    root, _ = example
    path = Path(__file__).resolve().parents[1] / "examples/surface-research/run_example.py"
    with pytest.raises(FileExistsError):
        runpy.run_path(str(path))["run"](root)


def test_prepare_only_leaves_four_empty_host_drafts(tmp_path):
    path = Path(__file__).resolve().parents[1] / "examples/surface-research/run_example.py"
    workspace = runpy.run_path(str(path))["run"](tmp_path / "fresh", prepare_only=True)
    mapping = read(workspace / "drafts.json")
    assert set(mapping) == {"methods", "results", "discussion", "conclusion"}
    for name in mapping.values():
        assert all(not paragraph["text"] for paragraph in read(workspace / name)["paragraphs"])
    assert not (workspace.parent / "candidate").exists()
