"""Connected synthetic workspace lifecycle, not a scientific validation claim."""

import json
import shutil

import numpy as np
import pytest
from docx import Document
from PIL import Image, ImageDraw

from cfdpaper.publication.figure_tasks import import_figure_task, prepare_figure_task
from cfdpaper.publication.research_workspace import (
    _figure_metadata,
    attach_imported_figure,
    build_research_workspace,
    prepare_surface_workspace,
)


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write(path, data):
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")


def study(root):
    (root / "sources").mkdir(parents=True)
    (root / "sources/definition.md").write_text(
        "Synthetic common-wall face values, physical areas, and final steady support.\n",
        encoding="utf-8",
    )
    vertices = [[0, 0, 0], [1, 0, 0], [1, 1, 0], [0, 1, 0], [4, 0, 0], [4, 1, 0]]
    for index, case in enumerate(("A", "B")):
        np.savez(
            root / f"sources/{case}.npz",
            face_value=[300.0 + 10 * index, 400.0 + 10 * index],
            area_m2=[1.0, 3.0],
            vertices_m=vertices,
            faces_connectivity=[4, 0, 1, 2, 3, 4, 1, 4, 5, 2],
            face_starts=[0, 5],
        )
    data = {
        "study_id": "surface-comparison",
        "title": "Synthetic surface comparison",
        "question": "How do these declared wall observations differ?",
        "time_support": "Final saved steady result",
        "common_conditions": ["Same wall domain", "Same prescribed heat load"],
        "definition": {"source": "sources/definition.md", "text": "Area-weighted face temperature"},
        "domain": "test wall",
        "quantity": "Temperature",
        "unit": "K",
        "coordinate_unit": "m",
        "projection_axes": [0, 1],
        "thresholds": [373.15],
        "arrays": {
            "values": "face_value",
            "weights": "area_m2",
            "vertices": "vertices_m",
            "connectivity": "faces_connectivity",
            "starts": "face_starts",
        },
        "cases": [
            {"id": case, "label": f"Case {case}", "source": f"sources/{case}.npz"}
            for case in ("A", "B")
        ],
    }
    write(root / "study.json", data)
    return root / "study.json"


def fill_drafts(workspace):
    for sid, name in read(workspace / "drafts.json").items():
        path = workspace / name
        draft, section = read(path), read(path.parent / "input.json")
        ids = [item["id"] for item in section["evidence"]]
        figures = [item["id"] for item in section["figures"]]
        text = (
            f"Author's distinct {sid} wording. "
            "A has {{value:A-mean}} with spread {{value:A-std}} and area {{value:A-exceedance-1}}; "
            "B has {{value:B-mean}}."
        )
        if figures:
            text += " {{figure:field}} and {{figure:distribution}}; {{table:surface-summary}}."
        draft["paragraphs"] = [{"text": text, "evidence_ids": ids, "figure_ids": figures}]
        draft["image_observations"] = {key: "author-provided" for key in figures}
        draft["captions"] = {key: "Author caption for " + key for key in figures}
        write(path, draft)


@pytest.mark.parametrize("selected", [["field"], ["distribution", "field"], []])
def test_surface_workspace_selects_publication_figures_without_losing_diagnostics(
    tmp_path, selected
):
    materials = tmp_path / "materials"
    materials.mkdir()
    source = study(materials)
    data = read(source)
    data["selected_figures"] = selected
    write(source, data)
    workspace = prepare_surface_workspace(source, tmp_path / "author")
    result = read(workspace / "sections/results/input.json")
    assert [figure["id"] for figure in result["figures"]] == selected
    assert result["duties"][0]["figure_ids"] == selected
    contract = next(
        item
        for item in read(workspace / "manuscript-input.json")["spine"]["sections"]
        if item["section_id"] == "results"
    )
    assert contract["required_figure_ids"] == selected
    assert set(read(workspace / "sections/results/draft.json")["captions"]) == set(selected)
    assert {
        item["id"] for item in read(workspace / "research-context.json")["selected_figures"]
    } == set(selected)
    assert (workspace / "sections/results/figures/distribution-source-data.csv").is_file()
    assert (workspace / "sections/results/figures/field-source-data.csv").is_file()


@pytest.mark.parametrize("selected", [["not-a-figure"], ["field", "field"], "field"])
def test_surface_workspace_rejects_invalid_figure_selection_before_output(tmp_path, selected):
    materials = tmp_path / "materials"
    materials.mkdir()
    source = study(materials)
    data = read(source)
    data["selected_figures"] = selected
    write(source, data)
    output = tmp_path / "author"
    with pytest.raises(ValueError, match="selected_figures"):
        prepare_surface_workspace(source, output)
    assert not output.exists()


def test_combined_diagnostic_recipe_refreshes_bound_text_and_preserves_author_assets(tmp_path):
    from cfdpaper.publication.research_workspace import attach_research_diagnostics

    materials = tmp_path / "materials"
    materials.mkdir()
    workspace = prepare_surface_workspace(study(materials), tmp_path / "author")
    (materials / "sources/heat.csv").write_text("heat,area\n10,1\n-2,1\n", encoding="utf-8")
    (materials / "sources/budget.csv").write_text(
        "input,output,source\n20,12,0\n", encoding="utf-8"
    )
    config = {
        "field_calculations": [
            {
                "id": "heat-path",
                "source": "sources/heat.csv",
                "operation": "field_integral",
                "columns": {"value": "heat", "measure": "area"},
                "units": {"value": "W/m2", "measure": "m2"},
                "domain": "declared wall",
                "field_options": {"association": "face", "integrand": "density"},
            },
            {
                "id": "column",
                "source": "sources/budget.csv",
                "operation": "energy_budget",
                "columns": {"input": "input", "output": "output", "source": "source"},
                "units": {"input": "W", "output": "W", "source": "W"},
                "domain": "steady column",
                "field_options": {"steady_state": True},
            },
        ],
        "panels": [
            {
                "id": "path",
                "kind": "signed_flux",
                "title": "Signed wall exchange",
                "sign_convention": "Positive points outward.",
                "series": [{"calculation_id": "heat-path", "label": "Wall"}],
            },
            {
                "id": "budget",
                "kind": "budget",
                "view": "decomposition",
                "title": "Column balance",
                "series": [{"calculation_id": "column", "label": "Column"}],
            },
        ],
    }
    write(materials / "diagnostics.json", config)
    attached = attach_research_diagnostics(
        workspace,
        materials / "diagnostics.json",
        tmp_path / "attached",
        combine=True,
        figure_id="heat-readout",
        panel_ids=["path", "budget"],
    )
    input_path = attached / "sections/results/input.json"
    section = read(input_path)
    for evidence in section["evidence"]:
        if evidence["id"] in {"heat-path-all-net", "A-mean"}:
            evidence["result_ref"]["places"] = 2
    write(input_path, section)
    fill_drafts(attached)
    path = attached / "sections/results/draft.json"
    draft = read(path)
    draft["paragraphs"][0]["text"] += " Net transfer is {{value:heat-path-all-net}}."
    draft["captions"]["heat-readout"] = "Author's current signed-transfer caption."
    write(path, draft)
    first = build_research_workspace(attached, tmp_path / "first")
    author = first / "workspace"
    local_script = author / "sections/results/figures/diagnostics/plot.py"
    local_script.write_text(
        local_script.read_text(encoding="utf-8") + "\n# Author style note\n", encoding="utf-8"
    )
    author_script = local_script.read_bytes()
    (author / "sections/results/sources/heat.csv").write_text(
        "heat,area\n20,1\n-2,1\n", encoding="utf-8"
    )
    updated = build_research_workspace(author, tmp_path / "updated", baseline_dir=first)
    reports = read(updated / "workspace/sections/results/table-results.json")
    assert next(r for r in reports if r["id"] == "heat-path")["groups"][0]["result"]["net"] == 18
    output = read(updated / "manuscript/section.json")
    assert any("Net transfer is 18.00 W." in p.get("text", "") for p in output["paragraphs"])
    assert any(f["caption"] == draft["captions"]["heat-readout"] for f in output["figures"])
    assert (
        updated / "workspace/sections/results/figures/diagnostics/plot.py"
    ).read_bytes() == author_script
    selected = read(updated / "workspace/sections/results/input.json")["figures"]
    assert [f["id"] for f in selected] == ["field", "distribution", "heat-readout"]
    current_evidence = read(updated / "workspace/sections/results/input.json")["evidence"]
    assert {
        e["id"]: e["result_ref"]["places"]
        for e in current_evidence
        if e["id"] in {"heat-path-all-net", "A-mean"}
    } == {"heat-path-all-net": 2, "A-mean": 2}
    recipe = next(
        r
        for r in read(updated / "workspace/figure-recipes.json")
        if r["id"] == "physical-diagnostics"
    )
    assert "sections/results/figures/diagnostics/heat-readout.png" in recipe["outputs"]
    assert (
        updated / "workspace/sections/results/figures/diagnostics/heat-readout.png"
    ).read_bytes() != (
        first / "workspace/sections/results/figures/diagnostics/heat-readout.png"
    ).read_bytes()


def test_diagnostic_metadata_reads_registered_secondary_bundle_without_replacing_caption(tmp_path):
    from cfdpaper.publication.research_workspace import _diagnostic_metadata

    root = tmp_path / "sections/results"
    bundle = root / "figures/height"
    bundle.mkdir(parents=True)
    (bundle / "height.png").write_bytes(b"artifact location only; no visual validation")
    write(
        root / "input.json",
        {"figures": [{"id": "height", "path": "old.png", "caption": "Author caption"}]},
    )
    write(
        bundle / "diagnostic-figures.json",
        {
            "figures": [
                {
                    "id": "height",
                    "path": "height.png",
                    "caption": "Generated caption",
                    "sizing": {"source_width_mm": 160},
                }
            ]
        },
    )
    write(
        tmp_path / "figure-recipes.json",
        [{"outputs": ["sections/results/figures/height/diagnostic-figures.json"]}],
    )
    _diagnostic_metadata(tmp_path)
    figure = read(root / "input.json")["figures"][0]
    assert figure["path"] == "figures/height/height.png"
    assert figure["sizing"] == {"source_width_mm": 160}
    assert figure["caption"] == "Author caption"


def test_attach_field_diagnostics_uses_existing_research_and_refresh(tmp_path):
    from cfdpaper.publication.research_workspace import attach_research_diagnostics

    materials = tmp_path / "materials"
    materials.mkdir()
    workspace = prepare_surface_workspace(study(materials), tmp_path / "author")
    fill_drafts(workspace)
    (materials / "sources/heat.csv").write_text("heat,area\n20,1\n-4,3\n", encoding="utf-8")
    config = {
        "field_calculations": [
            {
                "id": "heat-path",
                "source": "sources/heat.csv",
                "operation": "field_integral",
                "columns": {"value": "heat", "measure": "area"},
                "units": {"value": "W/m2", "measure": "m2"},
                "domain": "test wall",
                "field_options": {"association": "face", "integrand": "density"},
            }
        ],
        "panels": [
            {
                "id": "heat-path-figure",
                "kind": "signed_flux",
                "title": "Wall exchange",
                "sign_convention": "Positive heat leaves the solid.",
                "series": [{"calculation_id": "heat-path", "label": "Case A"}],
            }
        ],
        "mechanism_questions": ["Does net flow conceal reverse exchange?"],
    }
    write(materials / "diagnostics.json", config)
    attached = attach_research_diagnostics(
        workspace, materials / "diagnostics.json", tmp_path / "attached"
    )
    result = read(attached / "sections/results/input.json")
    assert any(e["id"] == "heat-path-all-net" for e in result["evidence"])
    attached_reports = read(attached / "sections/results/table-results.json")
    assert any(r["id"] == "heat-path" for r in attached_reports)
    method_context = read(attached / "sections/methods/manuscript-context.json")
    assert method_context["evidence_bindings"]["heat-path-all-net"] == "results/heat-path-all-net"
    summary = read(attached / "research-context.json")["result_summary"]
    heat = next(r for r in summary if r["id"] == "heat-path")
    assert heat["groups"][0]["result"]["net"] == 8
    assert "csv_records" not in heat["groups"][0]
    assert heat["groups"][0]["full_result"].startswith("sections/results/table-results.json#/")
    assert (attached / "sections/results/figures/diagnostics/plot.py").is_file()
    assert read(workspace / "drafts.json") == read(attached / "drafts.json")
    first = build_research_workspace(attached, tmp_path / "first")
    author = first / "workspace"
    (author / "sections/results/sources/heat.csv").write_text(
        "heat,area\n30,1\n-4,3\n", encoding="utf-8"
    )
    second = build_research_workspace(author, tmp_path / "second", baseline_dir=first)
    reports = read(second / "workspace/sections/results/table-results.json")
    assert next(r for r in reports if r["id"] == "heat-path")["groups"][0]["result"]["net"] == 18
    statuses = {
        s["recipe_id"]: s["status"] for s in read(second / "refresh.json")["figure_recipes"]
    }
    assert statuses["surface-study"] == "unchanged"
    assert statuses["physical-diagnostics"] == "executed"
    assert read(author / "sections/methods/draft.json") == read(
        second / "workspace/sections/methods/draft.json"
    )
    config_path = second / "workspace/sections/results/diagnostics.json"
    changed = read(config_path)
    changed["field_calculations"][0]["source"] = "sources/relocated-heat.csv"
    write(config_path, changed)
    new_source = config_path.parent / "sources/relocated-heat.csv"
    new_source.write_text("heat,area\n30,1\n-4,3\n", encoding="utf-8")
    third = build_research_workspace(second / "workspace", tmp_path / "third", baseline_dir=second)
    assert "diagnostics" in read(third / "refresh.json")["definition_changes"]
    recipe = next(
        r
        for r in read(third / "workspace/figure-recipes.json")
        if r["id"] == "physical-diagnostics"
    )
    assert "sections/results/sources/relocated-heat.csv" in recipe["inputs"]
    assert "sections/results/sources/heat.csv" not in recipe["inputs"]
    continued = third / "workspace"
    (continued / "sections/results/sources/relocated-heat.csv").write_text(
        "heat,area\n40,1\n-4,3\n", encoding="utf-8"
    )
    fourth = build_research_workspace(continued, tmp_path / "fourth", baseline_dir=third)
    statuses = {
        s["recipe_id"]: s["status"] for s in read(fourth / "refresh.json")["figure_recipes"]
    }
    assert statuses["physical-diagnostics"] == "executed"
    assert statuses["surface-study"] == "unchanged"


@pytest.fixture(scope="module")
def assets(tmp_path_factory):
    root = tmp_path_factory.mktemp("research-assets")
    source = study(root / "materials")
    prepared = prepare_surface_workspace(source, root / "prepared")
    filled = root / "filled"
    shutil.copytree(prepared, filled)
    fill_drafts(filled)
    built = build_research_workspace(filled, root / "baseline")
    return source, prepared, filled, built


@pytest.fixture(scope="module")
def figure_candidate(tmp_path_factory):
    root = tmp_path_factory.mktemp("figure-candidate")
    (root / "context.md").write_text("Synthetic interpretation, not physical validation.\n")
    write(
        root / "task.json",
        {
            "figure_id": "mechanism",
            "kind": "schematic",
            "purpose": "Synthetic mechanism drawing",
            "claim_ceiling": "Conceptual interpretation only",
            "sources": [
                {
                    "id": "context",
                    "path": "context.md",
                    "role": "context",
                    "description": "Synthetic context",
                }
            ],
            "labels": [{"text": "Synthetic pathway", "unit": ""}],
            "final_width_mm": 100,
        },
    )
    task = prepare_figure_task(root / "task.json", root / "task")
    returned = root / "returned"
    (returned / "artwork").mkdir(parents=True)
    (returned / "artwork/figure.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg"><rect width="100" height="40"/></svg>'
    )
    preview = Image.new("RGB", (800, 400), "white")
    ImageDraw.Draw(preview).rectangle((20, 20, 780, 380), outline="black", width=3)
    preview.save(returned / "artwork/figure.png")
    (returned / "artwork/legend.txt").write_text("A synthetic diagram, not a field plot.\n")
    write(
        returned / "delivery.json",
        {
            "figure_id": "mechanism",
            "editable_sources": ["artwork/figure.svg"],
            "preview": "artwork/figure.png",
            "caption": "Conceptual synthesis of {{figure:results/field}}; no causal validation.",
            "title": "Synthetic interpretation",
            "legend": "Solid outline indicates a synthetic object.",
            "body_reference": (
                "The interpretation in {{figure:mechanism}} is constrained by "
                "{{figure:results/field}}, without establishing transport causation."
            ),
            "exports": ["artwork/legend.txt"],
            "sizing": {"source_width_mm": 100, "minimum_source_font_pt": 10},
        },
    )
    return import_figure_task(task, returned / "delivery.json", root / "imported")


def clone(assets, tmp_path):
    baseline = assets[3]
    root = tmp_path / "working"
    shutil.copytree(baseline / "workspace", root)
    return root, baseline


def snapshot(root):
    return {p.relative_to(root).as_posix(): p.read_bytes() for p in root.rglob("*") if p.is_file()}


def scalar(data, sid, identity):
    return data["resolved_values"][json.dumps([sid, identity], separators=(",", ":"))]["raw_value"]


def test_prepare_has_real_fields_definitions_and_empty_host_templates(assets):
    source, prepared, _, _ = assets
    for name in ("A.npz", "B.npz", "definition.md"):
        assert (prepared / "sections/results/sources" / name).read_bytes() == (
            source.parent / "sources" / name
        ).read_bytes()
    assert read(prepared / "sections/results/study.json")["unit"] == "K"
    assert (prepared / "sections/results/figures/plot.py").is_file()
    assert (prepared / "sections/results/figures/field.svg").is_file()
    assert (prepared / "sections/results/figures/distribution-source-data.csv").is_file()
    manifest = read(prepared / "manuscript-input.json")
    assert [entry["section_id"] for entry in manifest["sections"]] == [
        "methods",
        "results",
        "discussion",
        "conclusion",
    ]
    for name in read(prepared / "drafts.json").values():
        assert read(prepared / name)["paragraphs"][0]["text"] == ""
    section = read(prepared / "sections/results/input.json")
    fraction = next(item for item in section["evidence"] if item["id"] == "A-exceedance-1")
    assert fraction["result_ref"]["percentage"] is True
    assert (prepared / "RESEARCH.md").is_file()
    assert not (prepared / "manuscript.docx").exists()


def test_figure_metadata_normalizes_workspace_alias_before_relative_paths(assets, tmp_path):
    work, _ = clone(assets, tmp_path)
    section_path = work / "sections/results/input.json"
    before = read(section_path)
    alias = work / "sections" / ".."
    assert alias != alias.resolve()

    _figure_metadata(alias)

    assert read(section_path) == before
    assert [figure["path"] for figure in before["figures"]] == [
        "figures/field.png",
        "figures/distribution.png",
    ]


def test_build_native_word_table_and_cross_section_values(assets):
    built = assets[3]
    data = read(built / "manuscript/section.json")
    assert scalar(data, "results", "A-mean") == 375
    assert scalar(data, "conclusion", "A-mean") == 375
    assert scalar(data, "results", "A-exceedance-1") == 0.75
    document = Document(built / "manuscript.docx")
    assert len(document.tables) == 1
    assert "75.000" in document.tables[0].rows[1].cells[3].text
    assert "%" in document.tables[0].rows[1].cells[3].text
    assert len(document.inline_shapes) == 2
    assert [figure["caption"] for figure in data["figures"]] == [
        "Author caption for field",
        "Author caption for distribution",
    ]


def test_source_update_refreshes_images_tables_bindings_but_preserves_author_text(assets, tmp_path):
    work, old = clone(assets, tmp_path)
    before = snapshot(old)
    path = work / "sections/results/sources/A.npz"
    with np.load(path) as bundle:
        arrays = {key: bundle[key].copy() for key in bundle.files}
    arrays["face_value"][0] = 380
    np.savez(path, **arrays)
    drafts = {name: (work / name).read_bytes() for name in read(work / "drafts.json").values()}
    result = build_research_workspace(work, tmp_path / "updated", baseline_dir=old)
    data = read(result / "manuscript/section.json")
    assert scalar(data, "results", "A-mean") == 395
    assert scalar(data, "discussion", "A-mean") == 395
    assert scalar(data, "conclusion", "A-mean") == 395
    assert scalar(data, "results", "A-exceedance-1") == 1
    assert "100.000" in Document(result / "manuscript.docx").tables[0].rows[1].cells[3].text
    for name, raw in drafts.items():
        assert (result / "workspace" / name).read_bytes() == raw
    assert (result / "workspace/sections/results/figures/field.png").read_bytes() != (
        old / "workspace/sections/results/figures/field.png"
    ).read_bytes()
    assert read(result / "refresh.json")["figure_recipes"][0]["status"] == "executed"
    assert read(result / "refresh.json")["interpretation_review_required"] is True
    assert read(result / "manuscript/changes.json")["baseline_available"] is True
    assert snapshot(old) == before


def test_author_script_style_and_caption_remain_the_refresh_entry(assets, tmp_path):
    work, old = clone(assets, tmp_path)
    script = work / "sections/results/figures/plot.py"
    script.write_text(
        script.read_text().replace("'height_mm': 110", "'height_mm': 115"), encoding="utf-8"
    )
    author_script = script.read_bytes()
    result = build_research_workspace(work, tmp_path / "restyled", baseline_dir=old)
    assert (result / "workspace/sections/results/figures/plot.py").read_bytes() == author_script
    summary = read(result / "workspace/sections/results/figures/surface-figures.json")
    assert summary["style"]["height_mm"] == 115
    assert read(result / "refresh.json")["interpretation_review_required"] is False
    assert (
        read(result / "manuscript/section.json")["figures"][0]["caption"]
        == "Author caption for field"
    )


def test_definition_update_requires_interpretation_review_not_rewritten_prose(assets, tmp_path):
    work, old = clone(assets, tmp_path)
    path = work / "sections/results/study.json"
    data = read(path)
    data["thresholds"] = [450]
    data["time_support"] = "Declared alternate support, requires host reconciliation"
    write(path, data)
    draft = (work / "sections/discussion/draft.json").read_bytes()
    result = build_research_workspace(work, tmp_path / "definition-update", baseline_dir=old)
    refresh = read(result / "refresh.json")
    assert refresh["interpretation_review_required"] is True
    assert set(refresh["definition_changes"]) == {"thresholds", "time_support"}
    assert scalar(read(result / "manuscript/section.json"), "results", "A-exceedance-1") == 0
    assert "450.000" in Document(result / "manuscript.docx").tables[0].rows[0].cells[3].text
    assert (result / "workspace/sections/discussion/draft.json").read_bytes() == draft
    assert "Interpretation review required" in (result / "CHANGES.md").read_text(encoding="utf-8")


def test_unchanged_relocated_workspace_preserves_current_figures_and_drafts(assets, tmp_path):
    work, old = clone(assets, tmp_path)
    preserved = snapshot(work)
    result = build_research_workspace(work, tmp_path / "relocated", baseline_dir=old)
    assert read(result / "refresh.json")["figure_recipes"][0]["status"] == "unchanged"
    assert read(result / "manuscript/changes.json")["changes"] == []
    for name, raw in preserved.items():
        if name.endswith("draft.json") or "/figures/" in name:
            assert (result / "workspace" / name).read_bytes() == raw


@pytest.mark.parametrize("failure", ["script", "empty-draft"])
def test_failure_does_not_publish_partial_build_or_modify_inputs(assets, tmp_path, failure):
    work, old = clone(assets, tmp_path)
    if failure == "script":
        (work / "sections/results/figures/plot.py").write_text(
            "raise RuntimeError('fixture failure')\n"
        )
    else:
        path = work / "sections/methods/draft.json"
        data = read(path)
        data["paragraphs"][0]["text"] = ""
        write(path, data)
    original = snapshot(work)
    output = tmp_path / "failed"
    with pytest.raises((ValueError, RuntimeError)):
        build_research_workspace(work, output, baseline_dir=old)
    assert not output.exists()
    assert snapshot(work) == original


def test_prepare_and_build_refuse_overwrite_or_nested_output(assets, tmp_path):
    source, prepared, _, old = assets
    with pytest.raises(FileExistsError):
        prepare_surface_workspace(source, prepared)
    with pytest.raises(ValueError, match="outside"):
        build_research_workspace(old / "workspace", old / "workspace/nested")
    with pytest.raises(ValueError, match="outside"):
        build_research_workspace(old / "workspace", old)


def test_imported_figure_builds_with_global_numbering_and_preserved_sources(
    assets, figure_candidate, tmp_path
):
    work, baseline = clone(assets, tmp_path)
    discussion_input = work / "sections/discussion/input.json"
    section = read(discussion_input)
    section["evidence"] = [record for record in section["evidence"] if record["id"] != "A-mean"]
    write(discussion_input, section)
    original, imported_original = snapshot(work), snapshot(figure_candidate)
    old_draft = read(work / "sections/discussion/draft.json")
    attached = attach_imported_figure(
        work,
        figure_candidate,
        tmp_path / "attached",
        section_id="discussion",
        evidence_ids=["A-mean", "study-definition"],
        image_observation="viewed",
        observation_note="The caller inspected this synthetic outline.",
        paragraph_index=0,
    )
    new_draft = read(attached / "sections/discussion/draft.json")
    assert new_draft["paragraphs"][1:] == old_draft["paragraphs"]
    assert new_draft["image_observations"]["mechanism"] == "viewed"
    assert new_draft["observation_notes"]["mechanism"] == (
        "The caller inspected this synthetic outline."
    )
    for name, content in original.items():
        if name not in {
            "manuscript-input.json",
            "sections/discussion/input.json",
            "sections/discussion/draft.json",
        }:
            assert (attached / name).read_bytes() == content
    result = build_research_workspace(attached, tmp_path / "candidate", baseline_dir=baseline)
    numbering = read(result / "manuscript/numbering.json")
    assert numbering["sections"]["discussion"]["figure"] == {"mechanism": "3"}
    data = read(result / "manuscript/section.json")
    figure = data["figures"][2]
    assert figure["id"] == "3"
    assert figure["caption"] == "Conceptual synthesis of Figure 1; no causal validation."
    assert figure["sizing"]["target_width_mm"] == 100
    assert any(
        p.get("text", "").startswith("The interpretation in Figure 3")
        and "Figure 1" in p["text"]
        and p["figure_ids"] == ["3"]
        for p in data["paragraphs"]
    )
    document = Document(result / "manuscript.docx")
    assert len(document.inline_shapes) == 3
    assert document.inline_shapes[2].width.mm == pytest.approx(100)
    assert any(
        p.text.startswith("Figure 3. Conceptual synthesis of Figure 1") for p in document.paragraphs
    )
    assert read(result / "refresh.json")["figure_recipes"][0]["status"] == "unchanged"
    section = read(attached / "sections/discussion/input.json")
    expected_names = {
        "task.json",
        "delivery.json",
        "sources/context.md",
        "artwork/figure.svg",
        "artwork/figure.png",
        "artwork/legend.txt",
    }
    assert set(section["source_files"]) == {
        "sources/figure-tasks/mechanism/" + name for name in expected_names
    }
    for root in (attached, result / "workspace", result / "manuscript"):
        sources = root / "sections/discussion/sources/figure-tasks/mechanism"
        assert not (sources / "skills").exists()
        for name in expected_names:
            assert (sources / name).read_bytes() == imported_original[name]
    assert snapshot(work) == original
    assert snapshot(figure_candidate) == imported_original


@pytest.mark.parametrize("changed_definition", ["thresholds", "definition"])
def test_imported_figure_review_lists_caption_without_changed_paragraph_values(
    assets, figure_candidate, tmp_path, changed_definition
):
    work, old = clone(assets, tmp_path)
    attached = attach_imported_figure(
        work,
        figure_candidate,
        tmp_path / "attached",
        section_id="discussion",
        evidence_ids=["A-mean", "study-definition"],
    )
    baseline = build_research_workspace(attached, tmp_path / "baseline", baseline_dir=old)
    working = tmp_path / "definition-update"
    shutil.copytree(baseline / "workspace", working)
    draft_path = "sections/discussion/draft.json"
    author_draft = (working / draft_path).read_bytes()
    caption = read(working / draft_path)["captions"]["mechanism"]
    artwork_path = "sections/discussion/sources/figure-tasks/mechanism/artwork/figure.png"
    artwork = (working / artwork_path).read_bytes()
    study_path = working / "sections/results/study.json"
    data = read(study_path)
    if changed_definition == "thresholds":
        data["thresholds"] = [450]
    else:
        data["definition"]["text"] += "; revised observation interpretation"
    write(study_path, data)

    result = build_research_workspace(working, tmp_path / "updated", baseline_dir=baseline)
    refresh = read(result / "refresh.json")
    assert refresh["definition_changes"] == [changed_definition]
    assert refresh["interpretation_review_required"] is True
    assert "discussion" in refresh["interpretation_review_sections"]
    passages = read(result / "manuscript/changes.json")["affected_passages"]["discussion"]
    mechanism = next(p for p in passages if any(f["id"] == "mechanism" for f in p["figures"]))
    assert mechanism["result_changes"] == []
    assert mechanism["figures"][0]["caption"] == caption
    for root in (working, baseline / "workspace", result / "workspace", result / "manuscript"):
        assert (root / draft_path).read_bytes() == author_draft
        assert (root / artwork_path).read_bytes() == artwork
    for root in (result, result / "manuscript"):
        report = (root / "CHANGES.md").read_text(encoding="utf-8")
        assert "- Figure mechanism:" in report
        assert caption in report


def test_imported_figure_defaults_to_append_and_not_viewed(assets, figure_candidate, tmp_path):
    work, _ = clone(assets, tmp_path)
    path = work / "sections/discussion/draft.json"
    old_draft = read(path)
    attached = attach_imported_figure(
        work,
        figure_candidate,
        tmp_path / "attached",
        section_id="discussion",
        evidence_ids=[],
    )
    draft = read(attached / "sections/discussion/draft.json")
    assert draft["paragraphs"][:-1] == old_draft["paragraphs"]
    assert draft["image_observations"]["mechanism"] == "not-viewed"
    assert "mechanism" not in draft["observation_notes"]
    with pytest.raises(ValueError, match="already exists"):
        attach_imported_figure(
            attached,
            figure_candidate,
            tmp_path / "duplicate",
            section_id="discussion",
            evidence_ids=[],
        )
    with pytest.raises(FileExistsError):
        attach_imported_figure(
            work, figure_candidate, attached, section_id="discussion", evidence_ids=[]
        )
    with pytest.raises(ValueError, match="outside"):
        attach_imported_figure(
            work, figure_candidate, work / "nested", section_id="discussion", evidence_ids=[]
        )


def test_explicit_figure_update_retains_identity_author_sources_and_body(
    assets, figure_candidate, tmp_path
):
    work, old = clone(assets, tmp_path)
    attached = attach_imported_figure(
        work,
        figure_candidate,
        tmp_path / "attached",
        section_id="discussion",
        evidence_ids=["A-mean"],
    )
    baseline = build_research_workspace(attached, tmp_path / "baseline", baseline_dir=old)
    author = baseline / "workspace"
    local = author / "sections/discussion/sources/figure-tasks/mechanism/artwork/figure.svg"
    local.write_text(local.read_text() + "<!-- author source -->")
    draft_path = author / "sections/discussion/draft.json"
    draft = read(draft_path)
    draft["paragraphs"][-1]["text"] += " Author interpretation retained."
    write(draft_path, draft)
    original = snapshot(author)
    changed = tmp_path / "new-artwork"
    shutil.copytree(figure_candidate, changed)
    metadata = read(changed / "delivery.json")
    metadata["caption"] = "Updated explanation of the same synthetic pathway."
    metadata["body_reference"] = "Do not duplicate {{figure:mechanism}} in the body."
    write(changed / "delivery.json", metadata)
    updated = attach_imported_figure(
        author,
        changed,
        tmp_path / "updated",
        section_id="discussion",
        evidence_ids=["A-mean"],
        replace_existing=True,
    )
    assert snapshot(author) == original
    updated_draft = read(updated / "sections/discussion/draft.json")
    assert updated_draft["paragraphs"] == draft["paragraphs"]
    assert updated_draft["captions"]["mechanism"] == metadata["caption"]
    assert updated_draft["image_observations"]["mechanism"] == "not-viewed"
    assert local.relative_to(author).as_posix() in snapshot(updated)
    assert (updated / local.relative_to(author)).read_bytes() == local.read_bytes()
    section = read(updated / "sections/discussion/input.json")
    assert len(section["figures"]) == 1
    assert section["figures"][0]["id"] == "mechanism"
    assert "mechanism-update-1/" in section["figures"][0]["path"]
    result = build_research_workspace(updated, tmp_path / "candidate", baseline_dir=baseline)
    assert read(result / "manuscript/numbering.json")["sections"]["discussion"]["figure"] == {
        "mechanism": "3"
    }
    assert all(s["status"] == "unchanged" for s in read(result / "refresh.json")["figure_recipes"])
    assert "discussion" in read(result / "manuscript/changes.json")["affected_passages"]


def test_figure_update_requires_existing_identity(assets, figure_candidate, tmp_path):
    work, _ = clone(assets, tmp_path)
    with pytest.raises(ValueError, match="existing"):
        attach_imported_figure(
            work,
            figure_candidate,
            tmp_path / "bad-update",
            section_id="discussion",
            evidence_ids=[],
            replace_existing=True,
        )


def test_figure_attachment_extends_explicit_author_block_order(assets, figure_candidate, tmp_path):
    work, old = clone(assets, tmp_path)
    path = work / "sections/discussion/draft.json"
    draft = read(path)
    for index, paragraph in enumerate(draft["paragraphs"]):
        paragraph["paragraph_id"] = f"author-p-{index}"
    previous_order = ["paragraph:" + p["paragraph_id"] for p in draft["paragraphs"]]
    draft["block_order"] = previous_order
    write(path, draft)
    attached = attach_imported_figure(
        work,
        figure_candidate,
        tmp_path / "attached",
        section_id="discussion",
        evidence_ids=["A-mean"],
        paragraph_index=0,
    )
    updated = read(attached / "sections/discussion/draft.json")
    assert updated["block_order"] == [
        "paragraph:figure-mechanism-intro",
        "figure:mechanism",
        *previous_order,
    ]
    build_research_workspace(attached, tmp_path / "candidate", baseline_dir=old)


def test_question_context_reads_current_values_definitions_and_selected_passages(assets, tmp_path):
    from cfdpaper.publication.research_workspace import research_question_context

    work, _ = clone(assets, tmp_path)
    path = work / "sections/results/sources/A.npz"
    with np.load(path) as source:
        arrays = {key: source[key] for key in source.files}
    arrays["face_value"] = arrays["face_value"] + 4
    np.savez(path, **arrays)
    before = snapshot(work)
    context = research_question_context(
        work,
        question="Does the mean describe the whole wall?",
        section_id="discussion",
        evidence_ids=["A-mean", "study-definition"],
    )
    assert snapshot(work) == before
    metric = context["evidence"][0]
    assert metric["resolved"]["raw_value"] == pytest.approx(379)
    assert metric["owner"] == "results/A-mean"
    assert metric["calculation"]["operation"] == "weighted_population"
    definition = context["evidence"][1]["source_excerpt"]
    assert definition["first_line"] == 1
    assert "Synthetic common-wall" in definition["text"]
    assert context["passages"]
    assert "vertices_m" not in json.dumps(context)
    assert 'face_value": [' not in json.dumps(context)
    with pytest.raises(ValueError, match="evidence"):
        research_question_context(
            work, question="Question", section_id="discussion", evidence_ids=["absent"]
        )


def test_question_context_keeps_a_declared_report_reading(assets, tmp_path):
    from cfdpaper.publication.research_workspace import research_question_context

    work, _ = clone(assets, tmp_path)
    path = work / "sections/discussion/input.json"
    section = read(path)
    section["evidence"].append(
        {
            "id": "reported",
            "kind": "metric",
            "value": "42.5",
            "unit": "W",
            "text": "Declared report reading",
            "source": "report row 4",
        }
    )
    write(path, section)
    reading = research_question_context(
        work, question="What was reported?", section_id="discussion", evidence_ids=["reported"]
    )["evidence"][0]
    assert reading["declared_value"] == "42.5" and reading["declared_unit"] == "W"
    assert reading["resolved"] == {} and reading["source_excerpt"] is None


@pytest.mark.parametrize(
    "failure",
    ["missing-token", "unknown-evidence", "section", "index", "preview", "observation", "source"],
)
def test_imported_figure_failure_preserves_inputs(assets, figure_candidate, tmp_path, failure):
    work, _ = clone(assets, tmp_path)
    imported = tmp_path / "imported"
    shutil.copytree(figure_candidate, imported)
    kwargs = {"section_id": "discussion", "evidence_ids": ["A-mean"]}
    metadata = read(imported / "delivery.json")
    if failure == "missing-token":
        metadata["body_reference"] = "Figure 3 is not a stable reference."
        write(imported / "delivery.json", metadata)
    elif failure == "unknown-evidence":
        kwargs["evidence_ids"] = ["not-an-evidence-id"]
    elif failure == "section":
        kwargs["section_id"] = "not-a-section"
    elif failure == "index":
        kwargs["paragraph_index"] = -1
    elif failure == "preview":
        (imported / metadata["preview"]).write_bytes(b"not an image")
    elif failure == "observation":
        kwargs["image_observation"] = "scientifically-approved"
    elif failure == "source":
        (imported / "sources/context.md").unlink()
    before, imported_before = snapshot(work), snapshot(imported)
    output = tmp_path / "failed"
    with pytest.raises((ValueError, OSError)):
        attach_imported_figure(work, imported, output, **kwargs)
    assert not output.exists()
    assert snapshot(work) == before
    assert snapshot(imported) == imported_before
