import json
import runpy
import shutil
from pathlib import Path

import numpy as np
import pytest
from docx import Document

pytest.importorskip("shapely")

from cfdpaper.publication.research_workspace import build_research_workspace


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def scalar(assembled, identity, section="results"):
    key = json.dumps([section, identity], separators=(",", ":"))
    return assembled["resolved_values"][key]["raw_value"]


@pytest.fixture(scope="module")
def example(tmp_path_factory):
    script = Path(__file__).resolve().parents[1] / "examples/physical-diagnostics/run_example.py"
    root = tmp_path_factory.mktemp("physical-example") / "demo"
    return root, runpy.run_path(str(script))["run"](root)


def test_public_draft_contains_region_signed_and_paired_evidence(example):
    root, candidate = example
    assembled = read(candidate / "manuscript/section.json")
    assert scalar(assembled, "A-mean") == pytest.approx(340.625)
    assert scalar(assembled, "B-mean") == pytest.approx(346.875)
    assert scalar(assembled, "region-A-left-weighted_mean") == pytest.approx(347.5)
    assert scalar(assembled, "region-B-left-weighted_mean") == pytest.approx(339.1666666667)
    assert scalar(assembled, "region-A-left-measure") == pytest.approx(0.0003)
    area_key = json.dumps(["results", "region-A-left-measure"], separators=(",", ":"))
    assert float(assembled["resolved_values"][area_key]["value"]) == pytest.approx(0.0003)
    assert scalar(assembled, "flux-A-all-positive") == pytest.approx(20)
    assert scalar(assembled, "flux-B-all-positive") == pytest.approx(20)
    assert scalar(assembled, "flux-A-all-net") == pytest.approx(10)
    assert scalar(assembled, "flux-B-all-net") == pytest.approx(4)
    assert scalar(assembled, "paired-all-difference_mean") == pytest.approx(6.25)
    assert scalar(assembled, "budget-all-storage") == 10
    assert scalar(assembled, "budget-all-residual") == 10
    doc = Document(candidate / "manuscript.docx")
    assert len(doc.inline_shapes) == 5
    assert len(doc.tables) == 2
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "{{" not in text
    assert "constructed" in text.lower()
    assert "mixing" in text
    results = candidate / "workspace/sections/results"
    for name in ("regions", "transfers", "paired-response"):
        for extension in ("svg", "pdf", "png", "tiff"):
            assert (results / f"figures/diagnostics/{name}.{extension}").stat().st_size > 0
    for path in (root / "materials/sources").iterdir():
        assert path.read_bytes() == (results / "sources" / path.name).read_bytes()


def test_refresh_preserves_author_script_and_updates_bound_diagnostics(example, tmp_path):
    _, candidate = example
    workspace = tmp_path / "working"
    shutil.copytree(candidate / "workspace", workspace)
    script = workspace / "sections/results/figures/diagnostics/plot.py"
    script.write_text(
        script.read_text(encoding="utf-8") + "\n# author local styling\n", encoding="utf-8"
    )
    author_script = script.read_bytes()
    drafts = {
        name: (workspace / name).read_bytes() for name in read(workspace / "drafts.json").values()
    }
    source = workspace / "sections/results/sources/A.npz"
    with np.load(source, allow_pickle=False) as data:
        arrays = {key: data[key].copy() for key in data.files}
    arrays["face_value"] += 2
    np.savez(source, **arrays)
    refreshed = build_research_workspace(workspace, tmp_path / "updated", baseline_dir=candidate)
    assembled = read(refreshed / "manuscript/section.json")
    assert scalar(assembled, "A-mean") == pytest.approx(342.625)
    assert scalar(assembled, "region-A-left-weighted_mean") == pytest.approx(349.5)
    assert scalar(assembled, "paired-all-difference_mean") == pytest.approx(4.25)
    assert scalar(assembled, "flux-A-all-net") == pytest.approx(10)
    assert (
        refreshed / "workspace/sections/results/figures/diagnostics/plot.py"
    ).read_bytes() == author_script
    assert all(
        (refreshed / "workspace" / name).read_bytes() == content for name, content in drafts.items()
    )
    assert read(refreshed / "refresh.json")["interpretation_review_required"] is True
