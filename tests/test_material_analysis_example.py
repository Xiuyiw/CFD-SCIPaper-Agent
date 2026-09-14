"""Replay the public synthetic example through the real analysis and writing APIs."""

import csv
import json
import runpy
from pathlib import Path

import pytest

EXAMPLE = Path(__file__).resolve().parents[1] / "examples/material-analysis/run_example.py"


def test_default_replay_binds_six_values_in_native_table(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn

    output = tmp_path / "table-example"
    run = runpy.run_path(str(EXAMPLE))["run"]
    document_path = run(output)
    proposal = json.loads((output / "recorded-host-proposal.json").read_text(encoding="utf-8"))
    candidate = proposal["candidates"][0]
    assert candidate["presentation"] == "table"
    assert candidate["presentation_reason"]
    assert {m["result_ref"]["field"] for m in candidate["metrics"]} == {"area", "rate", "mean_flux"}
    assert all("value" not in m and "unit" not in m for m in candidate["metrics"])
    section = json.loads((output / "section/section.json").read_text(encoding="utf-8"))
    assert section["figures"] == []
    assert not (output / "section-input/plot_analysis.py").exists()
    assert not (output / "section-input/figures").exists()
    expected = {
        "area-Reference": 5,
        "area-Modified": 10,
        "rate-Reference": 50,
        "rate-Modified": 70,
        "mean_flux-Reference": 10,
        "mean_flux-Modified": 7,
    }
    assert {k: v["raw_value"] for k, v in section["resolved_values"].items()} == expected
    for value in section["resolved_values"].values():
        assert value["csv_records"] in ([2, 3], [4, 5])
    assert "Table 1" in section["paragraphs"][0]["text"]
    assert "{{" not in json.dumps(section["tables"])
    assert section["tables"][0]["table_id"] == "1"
    document = docx.Document(document_path)
    assert len(document.tables) == 1
    assert len(document.inline_shapes) == 0
    assert [[c.text for c in r.cells] for r in document.tables[0].rows][1:] == [
        ["Reference", "5.00\u00a0m²", "50.00\u00a0W", "10.00\u00a0W\u00a0m⁻²"],
        ["Modified", "10.00\u00a0m²", "70.00\u00a0W", "7.00\u00a0W\u00a0m⁻²"],
    ]
    body = next(p for p in document.paragraphs if p.text.startswith("The modified"))
    assert body.text == body.text.lstrip()
    assert body._p.pPr.find(qn("w:ind")).get(qn("w:firstLineChars")) == "200"
    spacing = body._p.pPr.find(qn("w:spacing"))
    assert spacing.get(qn("w:before")) == spacing.get(qn("w:after")) == "0"
    caption = next(p for p in document.paragraphs if p.text.startswith("Table 1."))
    assert caption._p.pPr.find(qn("w:ind")) is None
    with pytest.raises(FileExistsError):
        run(output)


def test_optional_plot_replay_preserves_complementary_quantities(tmp_path):
    docx = pytest.importorskip("docx")
    output = tmp_path / "plot-example"
    document_path = runpy.run_path(str(EXAMPLE))["run"](output, presentation="plot")
    section = json.loads((output / "section/section.json").read_text(encoding="utf-8"))
    assert len(section["figures"]) == 3
    assert section["tables"] == []
    assert len(docx.Document(document_path).inline_shapes) == 3
    assert (output / "section-input/plot_analysis.py").is_file()
    for index, field in enumerate(("area", "rate", "mean_flux"), 1):
        stem = output / f"section-input/figures/analysis-{index}"
        for extension in ("svg", "pdf", "png", "tiff"):
            assert stem.with_suffix(f".{extension}").stat().st_size > 0
        assert "<text" in stem.with_suffix(".svg").read_text(encoding="utf-8")
        with Path(f"{stem}-source-data.csv").open(encoding="utf-8", newline="") as stream:
            rows = list(csv.DictReader(stream))
        assert len(rows) == 2
        assert {r["field"] for r in rows} == {field}
        assert {r["group"] for r in rows} == {"Reference", "Modified"}
