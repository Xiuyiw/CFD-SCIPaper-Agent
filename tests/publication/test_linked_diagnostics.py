"""Synthetic integration evidence, not validation of a physical CFD case."""

import json
import runpy
import shutil
from copy import deepcopy
from pathlib import Path

import pytest

from cfdpaper.analysis_suggestions import compile_analysis, prepare_analysis
from cfdpaper.publication.analysis_section import _points, build_analysis_section
from cfdpaper.publication.manuscript import _change_report, assemble_manuscript, prepare_manuscript
from cfdpaper.publication.section import assemble_section, export_section_docx


def write(path, value):
    path.write_text(json.dumps(value), encoding="utf-8")
    return path


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def test_public_linked_example_runs_with_known_contrasting_responses(tmp_path):
    script = Path(__file__).parents[2] / "examples/spatial-diagnostics/run_linked_example.py"
    section = runpy.run_path(str(script))["run"](tmp_path / "demo")
    results = read(section / "section.json")["resolved_values"]
    assert results["mean-change"]["raw_value"] == pytest.approx(-2.2)
    assert results["spread-change"]["raw_value"] == pytest.approx(2.4122475728)
    text = (section / "section.md").read_text(encoding="utf-8")
    assert text.count("-2.200") == 2
    assert text.count("2.412") == 2


def test_targeted_context_keeps_figure_and_detects_unrounded_sign_change():
    before = {
        "sections": {
            "wall": {
                "input": {"figures": [{"id": "map", "path": "field.png", "caption": "Original"}]},
                "draft": {
                    "paragraphs": [
                        {
                            "text": "Author explanation with {{value:delta}}.",
                            "evidence_ids": ["delta"],
                            "figure_ids": ["map"],
                        }
                    ],
                    "captions": {"map": "Author caption"},
                },
                "evidence": {
                    "delta": {
                        "resolved": {
                            "raw_value": -0.0001,
                            "value": "0.000",
                            "unit": "K",
                            "calculation_id": "contrast",
                            "group": "all",
                            "field": "difference",
                        }
                    }
                },
            }
        }
    }
    after = deepcopy(before)
    after["sections"]["wall"]["evidence"]["delta"]["resolved"]["raw_value"] = 0.0001
    untouched = deepcopy(after)
    report = _change_report(before, after)
    passage = report["affected_passages"]["wall"][0]
    assert passage["result_changes"][0]["difference_sign_changed"] is True
    assert passage["figures"] == [{"id": "map", "path": "field.png", "caption": "Author caption"}]
    assert after == untouched
    after["sections"]["wall"]["evidence"]["delta"]["resolved"]["unit"] = "Pa"
    assert (
        _change_report(before, after)["affected_passages"]["wall"][0]["result_changes"][0][
            "difference_sign_changed"
        ]
        is None
    )


def inputs(tmp_path):
    raw = tmp_path / "raw"
    raw.mkdir()
    for name, values in (("a", (300, 320)), ("b", (300, 310))):
        (raw / f"{name}.csv").write_text(
            f"facet,T,area\none,{values[0]},1\ntwo,{values[1]},3\n", encoding="utf-8"
        )
    (raw / "method.md").write_text(
        "Synthetic wall facet temperatures in K, areas in m2.\n"
        "Identical wall coverage and boundary conditions; two fixed configurations.\n",
        encoding="utf-8",
    )
    package = tmp_path / "analysis"
    prepare_analysis(raw, package)
    calculations = [
        {
            "id": name,
            "source": f"sources/{name}.csv",
            "operation": "weighted_population",
            "columns": {"value": "T", "weight": "area"},
            "units": {"value": "K", "weight": "m2"},
            "weight_kind": "area",
            "quantity_kind": "absolute-temperature",
            "domain": "Whole heated wall",
            "definition_source": {"path": "sources/method.md", "locator": "L1-L2"},
            "comparison": {"status": "supported", "scope": "Same wall and boundary conditions"},
            "member_id": ["facet"],
            "expected_members": [["one"], ["two"]],
            "interpretation_limits": ["Synthetic element means, not subfacet fluctuations"],
        }
        for name in ("a", "b")
    ]
    candidate = {
        "id": "wall",
        "title": "Wall temperature contrast",
        "question": "How do the declared wall means differ?",
        "rationale": "Compare weighted outputs without copying rounded intermediate values.",
        "presentation": "table",
        "calculations": calculations,
        "result_comparisons": [
            {
                "id": "contrast",
                "reference": {"calculation_id": "a", "group": "all", "field": "weighted_mean"},
                "comparison": {"calculation_id": "b", "group": "all", "field": "weighted_mean"},
                "domain": "Whole heated wall",
                "definition_source": "sources/method.md:L1-L2",
                "comparison_scope": "B minus A at the stated wall and boundary conditions",
                "status": "supported",
            }
        ],
        "metrics": [
            {
                "id": "delta",
                "text": "Difference in area-weighted wall means",
                "result_ref": {"calculation_id": "contrast", "group": "all", "field": "difference"},
            }
        ],
        "interpretation_limits": ["This difference does not establish a cooling mechanism"],
    }
    return package, candidate


def compile_candidate(package, candidate, out):
    return compile_analysis(
        package, write(package / "proposal.json", {"candidates": [candidate]}), "wall", out
    )


def draft():
    return {
        "title": "Wall temperature contrast",
        "paragraphs": [
            {
                "text": "The signed change in the wall mean is {{value:delta}}.",
                "evidence_ids": ["delta"],
                "figure_ids": [],
            }
        ],
        "captions": {},
        "image_observations": {},
        "evidence_notes": [],
        "tables": [
            {
                "table_id": "contrast",
                "caption": "Declared wall comparison",
                "columns": ["Quantity", "Change"],
                "rows": [["Area-weighted mean", "{{value:delta}}"]],
                "after_section_id": "wall",
                "evidence_ids": ["delta"],
            }
        ],
    }


def test_comparison_propagates_to_prose_table_plot_inputs_and_relocates(tmp_path):
    package, candidate = inputs(tmp_path)
    compiled = compile_candidate(package, candidate, tmp_path / "compiled")
    payload = read(compiled)
    points = _points(payload, compiled.parent)
    assert points[0]["raw_value"] == -7.5
    assert [p["raw_value"] for p in points[0]["supporting_results"]] == [315, 307.5]
    assert "source" not in points[0]
    assert (
        read(compiled.parent / "table-results.json")[-1]["groups"][0]["result"]["relative_change"]
        is None
    )
    writing = build_analysis_section(compiled, tmp_path / "section-input")
    path = write(tmp_path / "draft.json", draft())
    first = assemble_section(writing, path, tmp_path / "first")
    assert (first / "section.md").read_text().count("-7.500") == 2
    moved = tmp_path / "moved"
    shutil.move(writing, moved)
    (moved / "sources/b.csv").write_text("facet,T,area\none,300,1\ntwo,340,3\n")
    updated = assemble_section(moved, path, tmp_path / "updated")
    assert (updated / "section.md").read_text().count("15.000") == 2
    assert (first / "section.md").read_text().count("-7.500") == 2
    pytest.importorskip("docx")
    from docx import Document

    document = Document(export_section_docx(updated, tmp_path / "updated.docx"))
    assert "15.000" in document.tables[0].cell(1, 1).text


@pytest.mark.parametrize("status", ["unknown", "not-comparable"])
def test_comparison_requires_declared_comparability(tmp_path, status):
    package, candidate = inputs(tmp_path)
    candidate["result_comparisons"][0]["status"] = status
    with pytest.raises(ValueError, match=status):
        compile_candidate(package, candidate, tmp_path / "compiled")


def test_definition_locator_is_checked_not_just_retained(tmp_path):
    package, candidate = inputs(tmp_path)
    candidate["result_comparisons"][0]["definition_source"] = "sources/method.md:L99"
    with pytest.raises(ValueError, match="outside"):
        compile_candidate(package, candidate, tmp_path / "compiled")


def test_cross_section_binding_keeps_both_upstream_sources_and_updates(tmp_path):
    package, candidate = inputs(tmp_path)
    compiled = compile_candidate(package, candidate, tmp_path / "compiled")
    writing = build_analysis_section(compiled, tmp_path / "section-input")
    write(tmp_path / "result-draft.json", draft())
    summary = tmp_path / "summary"
    summary.mkdir()
    write(
        summary / "input.json",
        {
            "section_id": "summary",
            "title": "Summary",
            "question": "What changed?",
            "figures": [],
            "evidence": [],
            "duties": [
                {"purpose": "Summarize signed change", "evidence_ids": ["shared"], "figure_ids": []}
            ],
        },
    )
    write(
        summary / "draft.json",
        {
            "title": "Summary",
            "paragraphs": [
                {
                    "text": "The signed change is {{value:shared}}.",
                    "evidence_ids": ["shared"],
                    "figure_ids": [],
                }
            ],
            "captions": {},
            "image_observations": {},
            "evidence_notes": [],
        },
    )
    manifest = write(
        tmp_path / "manuscript.json",
        {
            "title": "Synthetic linked results",
            "spine": {
                "topic_id": "wall",
                "central_claim_id": "delta",
                "sections": [
                    {
                        "section_id": sid,
                        "role": role,
                        "title": sid,
                        "purpose": "Use current mean difference",
                        "required_claim_ids": [eid],
                    }
                    for sid, role, eid in (
                        ("summary", "abstract", "shared"),
                        ("wall", "results", "delta"),
                    )
                ],
            },
            "sections": [
                {
                    "section_id": "wall",
                    "input": writing.relative_to(tmp_path).as_posix() + "/input.json",
                },
                {
                    "section_id": "summary",
                    "input": "summary/input.json",
                    "evidence_bindings": {"shared": "wall/delta"},
                },
            ],
        },
    )
    drafts = write(
        tmp_path / "drafts.json", {"wall": "result-draft.json", "summary": "summary/draft.json"}
    )
    prepared = prepare_manuscript(manifest, tmp_path / "prepared")
    first = assemble_manuscript(prepared, drafts, tmp_path / "first")
    resolved = read(first / "section.json")["resolved_values"]['["summary","shared"]']
    assert "source" not in resolved
    assert [p["source"] for p in resolved["supporting_results"]] == [
        "sections/wall/sources/a.csv",
        "sections/wall/sources/b.csv",
    ]
    packet = read(first / "sections/summary/review-packet/bound-evidence.json")["shared"]
    assert packet["result_comparisons"][0]["id"] == "contrast"
    assert all(
        (first / "sections/summary/review-packet" / p["source"]).is_file()
        for p in packet["resolved"]["supporting_results"]
    )
    (first / "sections/wall/sources/b.csv").write_text("facet,T,area\none,300,1\ntwo,340,3\n")
    updated = assemble_manuscript(first, first / "drafts.json", tmp_path / "updated")
    assert "signed change is 15.000 K" in (updated / "manuscript.md").read_text()
    changes = read(updated / "changes.json")
    assert set(changes["affected_sections"]) == {"wall", "summary"}
    assert changes["affected_passages"]["summary"][0]["paragraph"] == 1
    context = changes["affected_passages"]["summary"][0]
    assert "{{value:shared}}" in context["text"]
    contrast = context["result_changes"][0]
    assert contrast["before"]["raw_value"] == -7.5
    assert contrast["after"]["raw_value"] == 15
    assert contrast["difference_sign_changed"] is True
    assert "Difference sign changed" in (updated / "CHANGES.md").read_text(encoding="utf-8")
    assert (first / "sections/summary/draft.json").read_bytes() == (
        updated / "sections/summary/draft.json"
    ).read_bytes()
