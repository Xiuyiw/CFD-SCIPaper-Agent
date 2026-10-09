"""Computed, bound and relocated public transient writing example."""

import json
import runpy
import shutil
from pathlib import Path

import pytest
from PIL import Image, ImageChops

from cfdpaper.publication.section import assemble_section, export_section_docx

EXAMPLE = Path(__file__).resolve().parents[2] / "examples/transient-writing"


def test_prepare_only_is_portable_without_recorded_answer(tmp_path):
    module = runpy.run_path(str(EXAMPLE / "run_example.py"))
    package = module["prepare"](tmp_path / "fresh")
    assert (package / "TASK.md").is_file()
    assert "file-reading web chat" in (package / "README.md").read_text(encoding="utf-8")
    assert not list(package.rglob("draft.json"))
    assert "invented" in (package / "sources/method.md").read_text(encoding="utf-8").lower()
    assert (package / "skills/cfd-evidence-writing/references/temporal-subsections.md").is_file()
    assert (package / "sources/history.csv").read_bytes() == (
        EXAMPLE / "inputs/history.csv"
    ).read_bytes()


def test_temporal_values_bind_and_survive_relocation_with_native_docx(tmp_path):
    docx = pytest.importorskip("docx")
    module = runpy.run_path(str(EXAMPLE / "run_example.py"))
    original = tmp_path / "original"
    section = module["run"](original)
    data = json.loads((section / "section.json").read_text(encoding="utf-8"))
    expected = {
        "A-max": 420,
        "B-max": 410,
        "A-peak": 6,
        "B-peak": 8,
        "A-crossing": 6,
        "B-crossing": 8,
        "A-mean": 4487.5 / 12,
        "B-mean": 370,
        "A-integral": 126,
        "B-integral": 126,
        "A-ledger": 128,
        "B-ledger": 124,
    }
    assert set(data["resolved_values"]) == set(expected)
    for name, value in expected.items():
        bound = data["resolved_values"][name]
        assert bound["raw_value"] == pytest.approx(value)
        assert bound["source"] == "sources/history.csv"
        assert bound["csv_records"]
    reports = json.loads((original / "writing/table-results.json").read_text(encoding="utf-8"))
    assert all(group["result"]["sample_count"] == 6 for r in reports for group in r["groups"])
    ledger = next(r for r in reports if r["id"] == "heat-ledger")
    assert all(g["result"]["integral"] is None for g in ledger["groups"])
    moved = Path(shutil.move(str(original), tmp_path / "moved"))
    again = assemble_section(
        moved / "writing", moved / "recorded-tutorial-draft.json", tmp_path / "reassembled"
    )
    rebuilt = json.loads((again / "section.json").read_text(encoding="utf-8"))
    assert rebuilt["resolved_values"] == data["resolved_values"]
    figure = moved / "writing/sources/figure"
    runpy.run_path(str(figure / "plot_history.py"))["run"](
        figure / "source-data.csv", tmp_path / "redrawn"
    )
    with (
        Image.open(figure / "history.png") as before,
        Image.open(tmp_path / "redrawn/history.png") as after,
    ):
        assert ImageChops.difference(before.convert("RGB"), after.convert("RGB")).getbbox() is None
    assert "<text" in (figure / "history.svg").read_text(encoding="utf-8")
    assert all(
        (figure / f"history.{ext}").stat().st_size > 0 for ext in ("svg", "pdf", "png", "tiff")
    )
    assert all(
        c["title_legend_gap_px"] > 0
        for c in json.loads((figure / "layout-checks.json").read_text(encoding="utf-8"))
    )
    document = docx.Document(
        export_section_docx(again, tmp_path / "again.docx", layout="near-reference")
    )
    assert len(document.inline_shapes) == 1 and len(document.tables) == 1
    assert document.inline_shapes[0].width / 36000 == pytest.approx(160)
    table_text = " ".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert "126.0 J" in table_text.replace("\u00a0", " ")
    paragraph = next(p for p in document.paragraphs if p.text.startswith("The invented histories"))
    assert paragraph.paragraph_format.space_before.pt == 0
    assert paragraph.paragraph_format.space_after.pt == 0
    from docx.oxml.ns import qn

    assert paragraph._p.pPr.ind.get(qn("w:firstLineChars")) == "200"
    assert document.styles["Normal"].font.name == "Times New Roman"
