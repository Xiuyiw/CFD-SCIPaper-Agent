"""Synthetic lifecycle checks: these fixtures are not AI-produced scientific prose."""

import json
import shutil
from pathlib import Path

import pytest
from PIL import Image

from cfdpaper.publication.manuscript import assemble_manuscript, prepare_manuscript


def write(path, value):
    path.write_text(json.dumps(value), encoding="utf-8")
    return path


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def fixture(tmp_path):
    sections, contracts, drafts = [], [], {}
    for index, sid in enumerate(("methods", "response", "transport"), 1):
        root = tmp_path / sid
        root.mkdir()
        (root / "sources").mkdir()
        (root / "sources" / "values.csv").write_text("case,value\na,1\na,3\n")
        Image.new("RGB", (60, 40), (index * 60, 20, 40)).save(root / "plot.png")
        data = {
            "section_id": sid,
            "title": sid.title(),
            "question": f"Purpose of {sid}?",
            "figures": [
                {
                    "id": "same",
                    "path": "plot.png",
                    "caption": sid,
                    "description": "Synthetic image for software testing.",
                }
            ],
            "evidence": [
                {
                    "id": "metric",
                    "kind": "metric",
                    "text": "Fixture mean",
                    "source": "values.csv",
                    "result_ref": {
                        "calculation_id": "mean",
                        "group": "a",
                        "field": "mean",
                        "places": 2,
                    },
                },
                {
                    "id": "ref",
                    "kind": "literature",
                    "text": "Shared test reference",
                    "source": "doi:synthetic-shared",
                },
                {
                    "id": "unique",
                    "kind": "literature",
                    "text": f"Reference for {sid}",
                    "source": f"doi:synthetic-{sid}",
                },
            ],
            "duties": [
                {
                    "purpose": "Describe the fixture",
                    "evidence_ids": ["metric"],
                    "figure_ids": ["same"],
                }
            ],
            "table_calculations": [
                {
                    "id": "mean",
                    "source": "sources/values.csv",
                    "operation": "population",
                    "columns": {"value": "value"},
                    "units": {"value": "K"},
                    "domain": "Synthetic rows",
                    "group_by": "case",
                }
            ],
        }
        write(root / "input.json", data)
        sections.append({"section_id": sid, "input": f"{sid}/input.json"})
        contracts.append(
            {
                "section_id": sid,
                "role": "methods" if index == 1 else "results",
                "title": sid.title(),
                "purpose": f"Assigned {sid} responsibility",
                "required_claim_ids": ["metric"],
                "required_figure_ids": ["same"],
            }
        )
        draft = {
            "title": sid.title(),
            "paragraphs": [
                {
                    "text": "Synthetic {{value:metric}}: {{figure:same}}, {{table:same}}, "
                    "{{equation:same}}, {{cite:unique}} and {{cite:ref}}. "
                    "Literal Figure 1 / Table 1 / Equation 1 / [1] / 12.3 stays. "
                    "Inline {{math:x}}.",
                    "evidence_ids": ["metric", "ref", "unique"],
                    "figure_ids": ["same"],
                    "inline_math": {
                        "x": {
                            "kind": "sub",
                            "children": [
                                {"kind": "symbol", "text": "T"},
                                {"kind": "text", "text": "{{cite:ref}}"},
                            ],
                        },
                        "unused": {"kind": "text", "text": "{{cite:unique}}"},
                    },
                }
            ],
            "captions": {"same": "Synthetic caption {{figure:same}} {{cite:ref}}"},
            "image_observations": {"same": "author-provided"},
            "evidence_notes": [f"Independent review note for {sid}"],
            "tables": [
                {
                    "table_id": "same",
                    "caption": sid + " {{cite:ref}}",
                    "columns": ["Metric", "Value"],
                    "rows": [["Mean", "{{value:metric}}"]],
                    "after_section_id": sid,
                    "evidence_ids": ["metric", "ref"],
                }
            ],
            "equations": [
                {
                    "equation_id": "same",
                    "expression": {
                        "kind": "fraction",
                        "children": [
                            {"kind": "text", "text": "{{value:metric}}"},
                            {"kind": "text", "text": "2"},
                        ],
                    },
                    "evidence_ids": ["metric"],
                }
            ],
        }
        write(root / "draft.json", draft)
        drafts[sid] = f"{sid}/draft.json"
    manifest = {
        "title": "Synthetic three-section manuscript",
        "spine": {"topic_id": "synthetic", "central_claim_id": "metric", "sections": contracts},
        "sections": sections,
        "context": "Public synthetic software fixture",
        "terms": {"mean": "equal-record arithmetic mean"},
    }
    return write(tmp_path / "input.json", manifest), write(tmp_path / "drafts.json", drafts)


def shared_literature_fixture(tmp_path):
    source, drafts = fixture(tmp_path)
    (tmp_path / "excerpt.txt").write_text("Synthetic study: the mean is defined over equal rows.")
    write(
        tmp_path / "bibliography.json",
        [
            {
                "id": "paper",
                "type": "article-journal",
                "title": "Synthetic reference",
                "DOI": "10.1234/test",
            },
            {
                "id": "alias",
                "type": "article-journal",
                "title": "Synthetic reference",
                "DOI": "https://doi.org/10.1234/TEST",
            },
        ],
    )
    write(
        tmp_path / "literature.json",
        {
            "bibliography": "bibliography.json",
            "supports": [
                {
                    "section_id": sid,
                    "evidence_id": "ref",
                    "reference_id": "alias" if sid == "transport" else "paper",
                    "source": "excerpt.txt",
                    "locator": "paragraph 1",
                    "excerpt": "the mean is defined over equal rows",
                    "claim": "Definition of the fixture mean",
                    "role": "method basis",
                    "status": "supported",
                }
                for sid in ("methods", "response", "transport")
            ],
        },
    )
    manifest = read(source)
    manifest["literature"] = "literature.json"
    write(source, manifest)
    return source, drafts


def test_shared_bibliography_aliases_and_portable_supported_citations(tmp_path):
    source, drafts = shared_literature_fixture(tmp_path)
    package = prepare_manuscript(source, tmp_path / "package")
    output = assemble_manuscript(package, drafts, tmp_path / "candidate")
    references = read(output / "section.json")["references"]
    shared = [r for r in references if r["source"] == "doi:10.1234/test"]
    assert len(shared) == 1
    assert shared[0]["text"] != "Definition of the fixture mean"
    reference_text = (
        (output / "manuscript.md").read_text(encoding="utf-8").split("## References")[1]
    )
    assert reference_text.lower().count("doi:10.1234/test") == 1
    for sid in ("methods", "response", "transport"):
        support = read(output / "sections" / sid / "literature-support.json")
        assert support["supports"][0]["reference_id"] == "paper"
        assert "literature-support.json" in (output / "sections" / sid / "TASK.md").read_text()
    moved = tmp_path / "moved"
    shutil.copytree(output, moved)
    for directory in (
        package,
        output,
        tmp_path / "methods",
        tmp_path / "response",
        tmp_path / "transport",
    ):
        shutil.rmtree(directory)
    for name in ("bibliography.json", "literature.json", "excerpt.txt"):
        (tmp_path / name).unlink()
    resumed = assemble_manuscript(moved, moved / "drafts.json", tmp_path / "resumed")
    assert read(resumed / "section.json")["references"] == references


def test_detached_review_includes_original_literature_excerpts(tmp_path):
    source, drafts = shared_literature_fixture(tmp_path)
    package = prepare_manuscript(source, tmp_path / "package")
    output = assemble_manuscript(package, drafts, tmp_path / "candidate")
    detached = tmp_path / "detached"
    shutil.copytree(output / "sections/response/review-packet", detached)
    support = read(detached / "literature-support.json")
    assert len(support["records"]) == 1
    for item in support["supports"]:
        assert item["excerpt"] in (detached / item["source"]).read_text(encoding="utf-8")
    assert "literature-support.json" in (detached / "review-prompt.md").read_text()


def test_csl_reference_formatting_keeps_global_identity_and_word_runs(tmp_path, monkeypatch):
    from docx import Document

    from cfdpaper.publication.section import export_section_docx

    source, drafts = shared_literature_fixture(tmp_path)
    # This fixture also has manual references. Give all of them explicit metadata;
    # the formatter must not invent bibliography details from their source labels.
    records = read(tmp_path / "bibliography.json")
    records.extend(
        {"id": sid, "type": "article-journal", "title": sid, "DOI": f"synthetic-{sid}"}
        for sid in ("methods", "response", "transport")
    )
    write(tmp_path / "bibliography.json", records)
    style = tmp_path / "numeric.csl"
    style.write_text("<style/>", encoding="utf-8")
    data = read(source)
    data["citation_style"] = style.name
    write(source, data)
    calls = []

    def format_entries(metadata, path):
        assert path.read_text() == "<style/>"
        calls.append([r["id"] for r in metadata])
        return [
            {
                "id": r["id"],
                "text": f"Formatted {r['title']}",
                "runs": [{"text": "Formatted "}, {"text": r["title"], "italic": True}],
            }
            for r in metadata
        ]

    monkeypatch.setattr(
        "cfdpaper.publication.manuscript.format_numeric_bibliography", format_entries
    )
    package = prepare_manuscript(source, tmp_path / "package")
    output = assemble_manuscript(package, drafts, tmp_path / "candidate")
    assert calls == [["methods", "paper", "response", "transport"]]
    assert (output / "citation-style.csl").read_text() == "<style/>"
    references = read(output / "section.json")["references"]
    assert len(references) == 4
    assert references[1]["source"] == "doi:10.1234/test"
    text = (output / "manuscript.md").read_text(encoding="utf-8")
    assert "[2] Formatted Synthetic reference" in text
    assert " — doi:" not in text
    doc = Document(export_section_docx(output, tmp_path / "paper.docx"))
    para = next(p for p in doc.paragraphs if p.text == "[2] Formatted Synthetic reference")
    assert any(r.italic and r.text == "Synthetic reference" for r in para.runs)
    resumed = assemble_manuscript(output, output / "drafts.json", tmp_path / "resumed")
    assert read(resumed / "section.json")["references"] == references


def test_csl_requires_metadata_for_manual_references(tmp_path, monkeypatch):
    source, drafts = shared_literature_fixture(tmp_path)
    data = read(source)
    data["citation_style"] = "numeric.csl"
    write(source, data)
    (tmp_path / "numeric.csl").write_text("<style/>")
    package = prepare_manuscript(source, tmp_path / "package")
    with pytest.raises(ValueError, match="shared bibliographic metadata"):
        assemble_manuscript(package, drafts, tmp_path / "candidate")
    assert not (tmp_path / "candidate").exists()


@pytest.mark.parametrize("withdrawal", ["unsupported", "needs-review", "deleted"])
def test_withdrawing_shared_support_prevents_stale_citation(tmp_path, withdrawal):
    source, drafts = shared_literature_fixture(tmp_path)
    package = prepare_manuscript(source, tmp_path / "package")
    output = assemble_manuscript(package, drafts, tmp_path / "candidate")
    path = output / "literature/literature.json"
    data = read(path)
    if withdrawal == "deleted":
        data["supports"].pop(0)
    else:
        data["supports"][0]["status"] = withdrawal
    write(path, data)
    with pytest.raises(ValueError, match="Unresolved IDs"):
        assemble_manuscript(output, output / "drafts.json", tmp_path / "rejected")
    assert not (tmp_path / "rejected").exists()


def test_shared_literature_cannot_replace_metric(tmp_path):
    source, _ = shared_literature_fixture(tmp_path)
    path = tmp_path / "literature.json"
    data = read(path)
    data["supports"][0]["evidence_id"] = "metric"
    write(path, data)
    with pytest.raises(ValueError, match="non-literature"):
        prepare_manuscript(source, tmp_path / "package")


def test_role_guidance_is_routed_and_included(tmp_path):
    source, _ = fixture(tmp_path)
    data = read(source)
    for record, role in zip(
        data["spine"]["sections"], ("introduction", "abstract", "conclusion"), strict=True
    ):
        record["role"] = role
    write(source, data)
    package = prepare_manuscript(source, tmp_path / "package")
    for sid, ref in (
        ("methods", "literature-sections.md"),
        ("response", "abstract-conclusions.md"),
        ("transport", "abstract-conclusions.md"),
    ):
        path = package / "sections" / sid
        assert ref in (path / "TASK.md").read_text()
        assert (path / "skills/cfd-evidence-writing/references" / ref).is_file()


def test_three_sections_portable_methods_math_numbering_and_notes(tmp_path):
    source, drafts = fixture(tmp_path)
    package = prepare_manuscript(source, tmp_path / "package")
    method = package / "sections/methods"
    assert "methods-sections.md" in (method / "TASK.md").read_text()
    reference = "skills/cfd-evidence-writing/references/methods-sections.md"
    assert (method / reference).read_bytes() == (
        Path(__file__).resolve().parents[2] / reference
    ).read_bytes()
    shared = read(method / "manuscript-context.json")
    assert len(shared["spine"]["sections"]) == 3
    assert shared["terms"]["mean"] == "equal-record arithmetic mean"
    # Original source assets are no longer needed after preparation.
    for sid in ("methods", "response", "transport"):
        (tmp_path / sid / "plot.png").unlink()
        (tmp_path / sid / "sources/values.csv").unlink()
    output = assemble_manuscript(package, drafts, tmp_path / "output")
    result, numbering = read(output / "section.json"), read(output / "numbering.json")
    assert [p["section_heading"] for p in result["paragraphs"] if "section_heading" in p] == [
        "Methods",
        "Response",
        "Transport",
    ]
    assert all("text" not in p for p in result["paragraphs"] if "section_heading" in p)
    assert [f["id"] for f in result["figures"]] == ["1", "2", "3"]
    assert [t["table_id"] for t in result["tables"]] == ["1", "2", "3"]
    assert [e["equation_id"] for e in result["equations"]] == ["1", "2", "3"]
    assert len({(output / f["path"]).read_bytes() for f in result["figures"]}) == 3
    assert len(result["references"]) == 4
    assert len(result["evidence"]) == 9
    for index, sid in enumerate(("methods", "response", "transport"), 1):
        assert result["paragraphs"][2 * index - 2]["section_id"] == sid
        assert result["section_objects"][sid] == {
            "role": "methods" if sid == "methods" else "results",
            "tables": [str(index)],
            "equations": [str(index)],
        }
        paragraph = result["paragraphs"][2 * index - 1]
        assert f"Figure {index}, Table {index}, Equation {index}" in paragraph["text"]
        assert "Literal Figure 1 / Table 1 / Equation 1 / [1] / 12.3 stays." in paragraph["text"]
        assert paragraph["runs"][1]["math"]["kind"] == "sub"
        assert paragraph["runs"][1]["math"]["children"][1]["text"] == "[2]"
        assert numbering["sections"][sid]["cite"]["ref"] == 2
        assert (output / f"sections/{sid}/draft.json").read_bytes() == (
            tmp_path / sid / "draft.json"
        ).read_bytes()
        assert read(output / f"sections/{sid}/section.json")["figures"][0]["id"] == "same"
        assert f"note for {sid}" in (output / f"notes/{sid}.md").read_text()
    assert "Independent review note" not in (output / "manuscript.md").read_text(encoding="utf-8")
    assert "\ue000" not in (output / "section.json").read_text(encoding="utf-8")


def test_reordered_spine_rebinds_objects_and_shared_citations(tmp_path):
    source, drafts = fixture(tmp_path)
    original = prepare_manuscript(source, tmp_path / "original-package")
    original_output = assemble_manuscript(original, drafts, tmp_path / "original")
    data = read(source)
    data["spine"]["sections"].reverse()
    write(source, data)
    package = prepare_manuscript(source, tmp_path / "reordered-package")
    output = assemble_manuscript(package, drafts, tmp_path / "reordered")
    before, after = read(original_output / "numbering.json"), read(output / "numbering.json")
    assert before["sections"]["transport"]["figure"]["same"] == "3"
    assert after["sections"]["transport"]["figure"]["same"] == "1"
    assert after["sections"]["methods"]["table"]["same"] == "3"
    assert after["sections"]["methods"]["equation"]["same"] == "3"
    assert before["sections"]["transport"]["cite"]["unique"] == 4
    assert after["sections"]["transport"]["cite"]["unique"] == 1
    assert read(output / "section.json")["paragraphs"][1]["text"].startswith(
        "Synthetic 2.00 K: Figure 1, Table 1, Equation 1, [1] and [2]."
    )


def test_recomputes_current_packaged_source_and_preserves_old_output(tmp_path):
    source, drafts = fixture(tmp_path)
    package = prepare_manuscript(source, tmp_path / "package")
    old = assemble_manuscript(package, drafts, tmp_path / "old")
    old_bytes = (old / "section.json").read_bytes()
    (package / "sections/response/sources/values.csv").write_text("case,value\na,2\na,6\n")
    write(package / "sections/response/table-results.json", [])
    output = assemble_manuscript(package, drafts, tmp_path / "new")
    result = read(output / "section.json")
    assert "4.00 K" in result["paragraphs"][3]["text"]
    assert result["tables"][1]["rows"][0][1] == "4.00 K"
    assert result["equations"][1]["expression"]["children"][0]["text"] == "4.00 K"
    key = read(output / "numbering.json")["sections"]["response"]["evidence"]["metric"]
    assert result["resolved_values"][key]["raw_value"] == 4.0
    assert (output / result["resolved_values"][key]["source"]).is_file()
    assert (old / "section.json").read_bytes() == old_bytes
    with pytest.raises(FileExistsError):
        assemble_manuscript(package, drafts, old)
    with pytest.raises(FileExistsError):
        prepare_manuscript(source, package)


@pytest.mark.parametrize("change", ["duplicate", "unknown", "wrong-owner", "figure-owner"])
def test_bad_spine_or_section_identity_leaves_no_package(tmp_path, change):
    source, _ = fixture(tmp_path)
    data = read(source)
    if change == "duplicate":
        data["sections"].append(data["sections"][0])
    elif change == "unknown":
        data["sections"][0]["section_id"] = "unknown"
    elif change == "wrong-owner":
        own = read(tmp_path / "methods/input.json")
        own["evidence"][0]["id"] = "different"
        own["duties"][0]["evidence_ids"] = ["different"]
        write(tmp_path / "methods/input.json", own)
    else:
        own = read(tmp_path / "methods/input.json")
        own["figures"][0]["id"] = "different"
        own["duties"][0]["figure_ids"] = ["different"]
        write(tmp_path / "methods/input.json", own)
    write(source, data)
    with pytest.raises(ValueError):
        prepare_manuscript(source, tmp_path / "package")
    assert not (tmp_path / "package").exists()


@pytest.mark.parametrize("change", ["missing", "extra", "coverage", "table-reference", "bad-token"])
def test_bad_drafts_fail_without_partial_output(tmp_path, change):
    source, drafts = fixture(tmp_path)
    package = prepare_manuscript(source, tmp_path / "package")
    if change in {"missing", "extra"}:
        mapping = read(drafts)
        if change == "missing":
            del mapping["methods"]
        else:
            mapping["extra"] = "methods/draft.json"
        write(drafts, mapping)
    else:
        path = tmp_path / "methods/draft.json"
        draft = read(path)
        if change == "coverage":
            draft["paragraphs"][0]["evidence_ids"].remove("metric")
        elif change == "table-reference":
            draft["paragraphs"][0]["text"] = draft["paragraphs"][0]["text"].replace(
                "{{table:same}}", "table"
            )
        else:
            draft["paragraphs"][0]["text"] += " {{figure:unknown}}"
        write(path, draft)
    with pytest.raises(ValueError):
        assemble_manuscript(package, drafts, tmp_path / "output")
    assert not (tmp_path / "output").exists()


def test_cross_section_objects_rebind_after_reordering(tmp_path):
    source, drafts = fixture(tmp_path)
    path = tmp_path / "response/draft.json"
    draft = read(path)
    draft["paragraphs"][0]["text"] += (
        " Defined in {{equation:methods/same}}; compare {{table:transport/same}}"
        " and {{figure:transport/same}}. Literal Equation 1 stays."
    )
    write(path, draft)
    package = prepare_manuscript(source, tmp_path / "package")
    output = assemble_manuscript(package, drafts, tmp_path / "first")
    text = read(output / "section.json")["paragraphs"][3]["text"]
    assert "Defined in Equation 1; compare Table 3 and Figure 3" in text
    assert "Literal Equation 1 stays" in text
    manifest = read(package / "manuscript-input.json")
    manifest["spine"]["sections"].reverse()
    write(package / "manuscript-input.json", manifest)
    output = assemble_manuscript(package, drafts, tmp_path / "reordered")
    text = read(output / "section.json")["paragraphs"][3]["text"]
    assert "Defined in Equation 3; compare Table 1 and Figure 1" in text
    assert "Literal Equation 1 stays" in text
    refs = read(output / "numbering.json")["sections"]["response"]["cross_references"]
    assert refs[0] == {"kind": "equation", "section_id": "methods", "id": "same", "number": "3"}
    assert "{{equation:methods/same}}" in (output / "sections/response/draft.json").read_text()


@pytest.mark.parametrize("target", ["missing/same", "methods/missing"])
def test_unknown_cross_section_object_is_not_guessed(tmp_path, target):
    source, drafts = fixture(tmp_path)
    path = tmp_path / "response/draft.json"
    draft = read(path)
    draft["paragraphs"][0]["text"] += " {{equation:" + target + "}}"
    write(path, draft)
    package = prepare_manuscript(source, tmp_path / "package")
    with pytest.raises(ValueError, match="Unknown cross-section"):
        assemble_manuscript(package, drafts, tmp_path / "output")
    assert not (tmp_path / "output").exists()


def test_markdown_keeps_objects_with_their_section(tmp_path):
    source, drafts = fixture(tmp_path)
    package = prepare_manuscript(source, tmp_path / "package")
    output = assemble_manuscript(package, drafts, tmp_path / "output")
    markdown = (output / "manuscript.md").read_text(encoding="utf-8")
    assert markdown.index("   (1)") < markdown.index("## Response")
    assert markdown.index("![Figure 1]") < markdown.index("## Response")


def test_assembled_manuscript_can_move_edit_and_reassemble_without_originals(tmp_path):
    original = tmp_path / "original-project"
    original.mkdir()
    source, drafts = fixture(original)
    author_mapping = drafts.read_bytes()
    prepared = prepare_manuscript(source, original / "prepared")
    assembled = assemble_manuscript(prepared, drafts, original / "assembled")
    moved = Path(shutil.copytree(assembled, tmp_path / "moved-manuscript"))
    before = read(moved / "section.json")
    kept_drafts = {
        sid: (moved / f"sections/{sid}/draft.json").read_bytes() for sid in ("methods", "transport")
    }
    source_bytes = {
        sid: (original / sid / "sources/values.csv").read_bytes()
        for sid in ("methods", "response", "transport")
    }
    shutil.rmtree(original)
    path = moved / "sections/response/draft.json"
    edited = read(path)
    edited["paragraphs"][0]["text"] += " Author revision of this section only."
    write(path, edited)
    continued = assemble_manuscript(moved, moved / "drafts.json", tmp_path / "continued")
    after = read(continued / "section.json")
    assert after["paragraphs"][3]["text"].endswith("Author revision of this section only.")
    assert after["references"] == before["references"]
    assert after["resolved_values"] == before["resolved_values"]
    assert after["tables"] == before["tables"]
    assert after["equations"] == before["equations"]
    for index, sid in enumerate(("methods", "response", "transport"), 1):
        assert (continued / f"sections/{sid}/sources/values.csv").read_bytes() == source_bytes[sid]
        paragraph = after["paragraphs"][2 * index - 1]
        expected = f"Synthetic 2.00 K: Figure {index}, Table {index}, Equation {index}"
        assert expected in paragraph["text"]
        assert paragraph["runs"][1]["math"]["children"][1]["text"] == "[2]"
        assert "Literal Figure 1 / Table 1 / Equation 1 / [1] / 12.3 stays." in paragraph["text"]
        if sid in kept_drafts:
            assert (continued / f"sections/{sid}/draft.json").read_bytes() == kept_drafts[sid]
            assert paragraph == before["paragraphs"][2 * index - 1]
    assert (moved / "author-drafts.json").read_bytes() == author_mapping


def test_moved_assembled_manuscript_preserves_cross_section_tokens(tmp_path):
    original = tmp_path / "original-project"
    original.mkdir()
    source, drafts = fixture(original)
    path = original / "response/draft.json"
    draft = read(path)
    draft["paragraphs"][0]["text"] += (
        " Defined in {{equation:methods/same}}; compare {{table:transport/same}}"
        " and {{figure:transport/same}}. Literal Equation 1 stays."
    )
    write(path, draft)
    prepared = prepare_manuscript(source, original / "prepared")
    assembled = assemble_manuscript(prepared, drafts, original / "assembled")
    moved = Path(shutil.copytree(assembled, tmp_path / "moved-manuscript"))
    shutil.rmtree(original)
    manifest = read(moved / "manuscript-input.json")
    manifest["spine"]["sections"].reverse()
    write(moved / "manuscript-input.json", manifest)
    continued = assemble_manuscript(moved, moved / "drafts.json", tmp_path / "continued")
    paragraph = read(continued / "section.json")["paragraphs"][3]["text"]
    assert "Defined in Equation 3; compare Table 1 and Figure 1" in paragraph
    assert "Literal Equation 1 stays" in paragraph
    assert "{{equation:methods/same}}" in (continued / "sections/response/draft.json").read_text()
    references = read(continued / "numbering.json")["sections"]["response"]["cross_references"]
    assert references[0] == {
        "kind": "equation",
        "section_id": "methods",
        "id": "same",
        "number": "3",
    }
    for sid in ("methods", "response", "transport"):
        context = read(continued / f"sections/{sid}/manuscript-context.json")
        assert [section["section_id"] for section in context["spine"]["sections"]] == [
            "transport",
            "response",
            "methods",
        ]


def test_assembled_manuscript_carries_continuation_inputs_tasks_and_skills(tmp_path):
    source, drafts = fixture(tmp_path)
    prepared = prepare_manuscript(source, tmp_path / "prepared")
    output = assemble_manuscript(prepared, drafts, tmp_path / "assembled")
    assert read(output / "drafts.json") == {
        sid: f"sections/{sid}/draft.json" for sid in ("methods", "response", "transport")
    }
    continuation = (output / "CONTINUE.md").read_text(encoding="utf-8")
    for name in ("manuscript.md", "manuscript-input.json", "drafts.json"):
        assert name in continuation
    for sid in ("methods", "response", "transport"):
        local = output / "sections" / sid
        assert read(local / "input.json")["section_id"] == sid
        assert "manuscript-context.json" in (local / "TASK.md").read_text(encoding="utf-8")
        context = read(local / "manuscript-context.json")
        assert context["section"]["section_id"] == sid
        assert len(context["spine"]["sections"]) == 3
        skill = "skills/cfd-evidence-writing/SKILL.md"
        assert (local / skill).read_bytes() == (prepared / "sections" / sid / skill).read_bytes()
    methods_reference = "skills/cfd-evidence-writing/references/methods-sections.md"
    assert (output / "sections/methods" / methods_reference).is_file()


def bound_fixture(tmp_path):
    source, drafts = fixture(tmp_path)
    data, mapping = read(source), read(drafts)
    for sid, role in (("summary", "abstract"), ("conclusions", "conclusion")):
        root = tmp_path / sid
        root.mkdir()
        write(
            root / "input.json",
            {
                "section_id": sid,
                "title": sid.title(),
                "question": "What follows from the result?",
                "figures": [],
                "evidence": [],
                "duties": [{"purpose": "Summarize", "evidence_ids": ["shared"], "figure_ids": []}],
            },
        )
        write(
            root / "draft.json",
            {
                "title": sid.title(),
                "paragraphs": [
                    {
                        "text": "Mean: {{value:shared}}.",
                        "evidence_ids": ["shared"],
                        "figure_ids": [],
                    }
                ],
                "captions": {},
                "image_observations": {},
                "evidence_notes": [],
            },
        )
        mapping[sid] = f"{sid}/draft.json"
        data["sections"].append(
            {
                "section_id": sid,
                "input": f"{sid}/input.json",
                "evidence_bindings": {"shared": "response/metric"},
            }
        )
        contract = {
            "section_id": sid,
            "role": role,
            "title": sid.title(),
            "purpose": "Summarize current evidence",
            "required_claim_ids": ["shared"],
        }
        data["spine"]["sections"].insert(
            0 if sid == "summary" else len(data["spine"]["sections"]), contract
        )
    data["keywords"] = ["Synthetic transport", "Evidence-linked writing"]
    return write(source, data), write(drafts, mapping)


def test_bound_source_updates_summaries_without_editing_other_drafts(tmp_path):
    source, drafts = bound_fixture(tmp_path)
    prepared = prepare_manuscript(source, tmp_path / "prepared")
    first = assemble_manuscript(prepared, drafts, tmp_path / "first")
    assert not read(first / "changes.json")["baseline_available"]
    assert "Mean: 2.00 K." in (first / "manuscript.md").read_text(encoding="utf-8")
    assert not (first / "sections/summary/sources").exists()
    working = tmp_path / "working"
    shutil.copytree(first, working)
    # Local rendered snapshots must never override their owner's current calculation.
    local = read(working / "sections/summary/input.json")
    local["evidence"][0]["value"] = "9999"
    write(working / "sections/summary/input.json", local)
    (working / "sections/response/sources/values.csv").write_text("case,value\na,5\na,7\n")
    output = assemble_manuscript(working, working / "drafts.json", tmp_path / "second")
    result = read(output / "section.json")
    text = (output / "manuscript.md").read_text(encoding="utf-8")
    assert text.count("Mean: 6.00 K.") == 2
    assert "9999" not in text
    change = read(output / "changes.json")
    assert set(change["affected_sections"]) == {"response", "summary", "conclusions"}
    assert change["unchanged_sections"] == ["methods", "transport"]
    assert change["affected_passages"]["summary"][0]["paragraph"] == 1
    for sid in ("methods", "transport", "summary", "conclusions"):
        assert (first / f"sections/{sid}/draft.json").read_bytes() == (
            output / f"sections/{sid}/draft.json"
        ).read_bytes()
    value = result["resolved_values"][json.dumps(["summary", "shared"], separators=(",", ":"))]
    assert value["source"] == "sections/response/sources/values.csv"
    assert value["owner_evidence"] == "metric"
    assert value["raw_value"] == 6
    assert (first / "sections/response/sources/values.csv").read_text() == "case,value\na,1\na,3\n"
    third = assemble_manuscript(output, output / "drafts.json", tmp_path / "third")
    assert read(third / "changes.json")["changes"] == []


@pytest.mark.parametrize(
    "target,message",
    [
        ("summary/shared", "self evidence"),
        ("missing/metric", "Unknown"),
        ("response/missing", "Unknown bound"),
        ("response/ref", "shared literature"),
        ("conclusions/shared", "chain"),
        ("response", "owner-section"),
    ],
)
def test_invalid_bindings_do_not_make_candidates(tmp_path, target, message):
    source, _ = bound_fixture(tmp_path)
    data = read(source)
    data["sections"][-2]["evidence_bindings"]["shared"] = target
    write(source, data)
    with pytest.raises(ValueError, match=message):
        prepare_manuscript(source, tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_dependency_cycle_and_author_evidence_overwrite_rejected(tmp_path):
    source, _ = bound_fixture(tmp_path)
    data = read(source)
    data["sections"][1]["depends_on"] = ["summary"]
    write(source, data)
    with pytest.raises(ValueError, match="cycle"):
        prepare_manuscript(source, tmp_path / "cycle")
    data["sections"][1]["depends_on"] = []
    write(source, data)
    summary = tmp_path / "summary/input.json"
    raw = read(summary)
    raw["evidence"] = [
        {
            "id": "shared",
            "kind": "metric",
            "text": "Author metric",
            "source": "author",
            "value": "9",
        }
    ]
    write(summary, raw)
    with pytest.raises(ValueError, match="author-owned"):
        prepare_manuscript(source, tmp_path / "overwrite")


def test_removed_binding_cannot_keep_old_numeric_copy(tmp_path):
    source, drafts = bound_fixture(tmp_path)
    prepared = prepare_manuscript(source, tmp_path / "prepared")
    first = assemble_manuscript(prepared, drafts, tmp_path / "first")
    manifest = read(first / "manuscript-input.json")
    next(s for s in manifest["sections"] if s["section_id"] == "summary")["evidence_bindings"] = {}
    write(first / "manuscript-input.json", manifest)
    with pytest.raises(ValueError):
        assemble_manuscript(first, first / "drafts.json", tmp_path / "invalid")


def test_keywords_follow_abstract_in_markdown_and_docx(tmp_path):
    from docx import Document

    from cfdpaper.publication.section import export_section_docx

    source, drafts = bound_fixture(tmp_path)
    prepared = prepare_manuscript(source, tmp_path / "prepared")
    output = assemble_manuscript(prepared, drafts, tmp_path / "assembled")
    text = (output / "manuscript.md").read_text(encoding="utf-8")
    assert text.index("Mean: 2.00 K.") < text.index("Keywords:") < text.index("## Methods")
    docx = export_section_docx(output, tmp_path / "manuscript.docx", layout="near-reference")
    paragraphs = Document(docx).paragraphs
    index = next(i for i, p in enumerate(paragraphs) if p.text.startswith("Keywords:"))
    assert paragraphs[index - 1].text == "Mean: 2.00 K."
    assert paragraphs[index + 1].text == "Methods"
    assert paragraphs[index].paragraph_format.first_line_indent.pt == 0


def test_table_only_binding_keeps_source_and_standalone_review_materials(tmp_path):
    source, drafts = bound_fixture(tmp_path)
    data = read(source)
    entry = next(s for s in data["sections"] if s["section_id"] == "summary")
    entry["evidence_bindings"]["table-mean"] = "response/metric"
    write(source, data)
    path = tmp_path / "summary/draft.json"
    draft = read(path)
    draft["tables"] = [
        {
            "table_id": "summary-values",
            "caption": "Fixture mean",
            "columns": ["Metric", "Value"],
            "rows": [["Mean", "{{value:table-mean}}"]],
            "after_section_id": "summary",
            "evidence_ids": ["table-mean"],
        }
    ]
    write(path, draft)
    prepared = prepare_manuscript(source, tmp_path / "prepared")
    output = assemble_manuscript(prepared, drafts, tmp_path / "out")
    combined = read(output / "section.json")
    key = json.dumps(["summary", "table-mean"], separators=(",", ":"))
    assert combined["resolved_values"][key]["source"] == "sections/response/sources/values.csv"
    packet = output / "sections/summary/review-packet"
    bound = read(packet / "bound-evidence.json")["table-mean"]
    assert (packet / bound["resolved"]["source"]).read_text() == "case,value\na,1\na,3\n"
    assert bound["definition"]["result_ref"]["calculation_id"] == "mean"
    assert bound["calculations"][0]["units"] == {"value": "K"}
    assert "bound-evidence.json" in (packet / "review-prompt.md").read_text()
