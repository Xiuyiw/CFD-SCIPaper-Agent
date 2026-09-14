"""Portable weighted-statistics example: real calculation, figure and document wiring."""

import json
import runpy
import shutil
from pathlib import Path

import pytest
from PIL import Image, ImageChops

from cfdpaper.publication.section import assemble_section, export_section_docx

EXAMPLE = Path(__file__).resolve().parents[2] / "examples/spatial-diagnostics"


def test_fresh_host_package_does_not_contain_recorded_answers(tmp_path):
    module = runpy.run_path(str(EXAMPLE / "run_example.py"))
    output = module["prepare"](tmp_path / "fresh")
    assert (output / "materials/host-task.md").is_file()
    assert not list(output.rglob("draft.json"))
    assert not list(output.rglob("proposal.json"))
    assert (output / "materials/sources/wall.csv").read_bytes() == (
        EXAMPLE / "inputs/wall.csv"
    ).read_bytes()


def test_weighted_example_survives_relocation_and_keeps_document_format(tmp_path):
    docx = pytest.importorskip("docx")
    module = runpy.run_path(str(EXAMPLE / "run_example.py"))
    original = tmp_path / "demo"
    module["run"](original)
    moved = Path(shutil.move(str(original), tmp_path / "relocated"))
    section = assemble_section(
        moved / "section-input/writing", moved / "recorded-host-draft.json", tmp_path / "again"
    )
    text = (section / "section.md").read_text(encoding="utf-8")
    assert all(value in text for value in ("45.60", "43.40", "2.939", "5.352"))
    assert "**Table 1." in text and "Figure 1" in text
    volume = json.loads((moved / "volume-check/table-results.json").read_text())
    assert volume[0]["weight_kind"] == "volume"
    groups = volume[0]["groups"]
    assert [g["result"]["weighted_mean"] for g in groups] == pytest.approx([0.4, 0.4])
    assert [g["result"]["weighted_std"] for g in groups] == pytest.approx(
        [0.21213203435596423, 0.07071067811865477]
    )
    assert "remains pending" not in (moved / "section-input/writing/TASK.md").read_text()
    data = json.loads((moved / "section-input/writing/input.json").read_text())
    assert "Custom figure production remains pending" not in data["context"]

    figure = moved / "figure-delivery/artwork"
    runpy.run_path(str(figure / "plot_wall.py"))["run"](
        figure / "source-data.csv", tmp_path / "redrawn"
    )
    with (
        Image.open(figure / "wall.png") as before,
        Image.open(tmp_path / "redrawn/wall.png") as after,
    ):
        assert ImageChops.difference(before.convert("RGB"), after.convert("RGB")).getbbox() is None
    assert "<text" in (figure / "wall.svg").read_text(encoding="utf-8")
    assert all((figure / f"wall.{ext}").stat().st_size for ext in ("pdf", "png", "tiff"))

    document = docx.Document(
        export_section_docx(section, tmp_path / "example.docx", layout="near-reference")
    )
    assert len(document.tables) == 1 and len(document.inline_shapes) == 1
    assert document.inline_shapes[0].width / 36000 == pytest.approx(160)
    assert any(p.text.startswith("Table 1.") for p in document.paragraphs)
    paragraph = next(p for p in document.paragraphs if p.text.startswith("The Modified field"))
    body = paragraph.paragraph_format
    assert body.space_before.pt == 0 and body.space_after.pt == 0
    assert body.first_line_indent.pt > 0
    from docx.oxml.ns import qn

    assert paragraph._p.pPr.ind.get(qn("w:firstLineChars")) == "200"
