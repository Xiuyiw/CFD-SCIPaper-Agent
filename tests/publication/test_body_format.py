"""Body paragraph formatting must not leak into other manuscript elements."""

import json
import zipfile

import pytest
from PIL import Image, ImageDraw

from cfdpaper.publication.section import (
    assemble_section,
    export_section_docx,
    prepare_section,
)
from cfdpaper.publication.style import PublicationStyle


def sample_section(root, style=None, *, with_figure=True):
    """Public synthetic content, also usable for a real renderer preview."""
    root.mkdir(parents=True)
    picture = Image.new("RGB", (800, 160), "#edf2f4")
    ImageDraw.Draw(picture).text(
        (30, 65), "Synthetic figure - paragraph formatting sample", fill="#263238"
    )
    picture.save(root / "figure.png")
    source = {
        "section_id": "s1",
        "title": "Synthetic manuscript formatting sample",
        "question": "Does body formatting remain separate from other elements?",
        "figures": [
            {
                "id": "1",
                "path": "figure.png",
                "caption": "Synthetic formatting illustration.",
                "description": "A labelled placeholder, not scientific results.",
            }
        ],
        "evidence": [
            {
                "id": "ref",
                "kind": "literature",
                "text": "Synthetic reference for formatting only",
                "source": "Synthetic source",
            }
        ],
        "duties": [{"purpose": "Formatting", "evidence_ids": ["ref"], "figure_ids": ["1"]}],
        "style": style or {},
    }
    if not with_figure:
        source["figures"] = []
        source["duties"][0]["figure_ids"] = []
    source_path = root / "source.json"
    source_path.write_text(json.dumps(source), encoding="utf-8")
    package = prepare_section(source_path, root / "package")
    draft = {
        "title": source["title"],
        "paragraphs": [
            {
                "text": "First paragraph: the indentation is a paragraph property, not spaces. "
                "This synthetic text is long enough to wrap onto subsequent lines so the "
                "first-line offset can be checked in a rendered document. "
                "The same body formatting applies to editable inline math {{math:q}} and "
                "to the reference {{cite:ref}}. Figure {{figure:1}} is placed separately.",
                "inline_math": {"q": {"kind": "symbol", "text": "q"}},
                "evidence_ids": ["ref"],
                "figure_ids": ["1"],
            },
            {
                "text": "Second paragraph: its first line should have the same offset as "
                "the first paragraph. No added paragraph gap is expected under the defaults. "
                "Captions, headings, equations, table cells and image paragraphs retain "
                "their own formatting and must not inherit body indentation.",
                "evidence_ids": [],
                "figure_ids": [],
            },
        ],
        "captions": {"1": "Synthetic formatting illustration."},
        "evidence_notes": [],
        "image_observations": {"1": "author-provided"},
        "tables": [
            {
                "table_id": "1",
                "caption": "Synthetic table caption",
                "columns": ["Item", "Value"],
                "rows": [["Example", "1"]],
                "after_section_id": "s1",
                "note": "Synthetic table note.",
            }
        ],
        "equations": [{"equation_id": "1", "expression": {"kind": "symbol", "text": "q"}}],
    }
    if not with_figure:
        draft["captions"] = {}
        draft["image_observations"] = {}
        draft["paragraphs"][0]["figure_ids"] = []
        draft["paragraphs"][0]["text"] = "A table-only discussion with {{cite:ref}}."
        draft["paragraphs"][0].pop("inline_math")
    draft_path = root / "draft.json"
    draft_path.write_text(json.dumps(draft), encoding="utf-8")
    return assemble_section(package, draft_path, root / "section")


@pytest.mark.parametrize("layout", ["after-text", "near-reference"])
@pytest.mark.parametrize(
    "style, chars, before, after",
    [
        ({}, 2, 0, 0),
        (
            {
                "body_first_line_indent_chars": 0,
                "body_space_before_pt": 6,
                "body_space_after_pt": 12,
            },
            0,
            6,
            12,
        ),
        ({"body_first_line_indent_chars": 1.5, "body_pt": 12}, 1.5, 0, 0),
    ],
)
def test_body_defaults_and_explicit_template_overrides(
    tmp_path, layout, style, chars, before, after
):
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn

    section = sample_section(tmp_path / "sample", style)
    output = export_section_docx(section, tmp_path / "body.docx", layout=layout)
    document = docx.Document(output)
    bodies = [
        p for p in document.paragraphs if p.text.startswith(("First paragraph", "Second paragraph"))
    ]
    assert len(bodies) == 2
    for paragraph in bodies:
        assert paragraph.text == paragraph.text.lstrip()
        assert "\t" not in paragraph.text
        spacing = paragraph._p.pPr.find(qn("w:spacing"))
        assert spacing.get(qn("w:before")) == str(before * 20)
        assert spacing.get(qn("w:after")) == str(after * 20)
        indent = paragraph._p.pPr.find(qn("w:ind"))
        assert indent.get(qn("w:firstLineChars")) == str(round(chars * 100))
        assert indent.get(qn("w:firstLine")) == str(round(chars * style.get("body_pt", 11) * 20))
    assert bodies[0]._p.xpath("./m:oMath")  # Inline math remains in the body paragraph.
    assert document.styles["Normal"].paragraph_format.first_line_indent is None
    assert document.styles["Caption"].paragraph_format.space_after.pt == 8
    report = json.loads(output.with_suffix(".layout.json").read_text(encoding="utf-8"))
    assert report["style"]["body_first_line_indent_chars"] == chars
    assert report["style"]["body_space_before_pt"] == before
    assert report["style"]["body_space_after_pt"] == after


def test_body_changes_leave_non_body_xml_and_figure_sizing_unchanged(tmp_path):
    docx = pytest.importorskip("docx")
    snapshots = []
    for name, overrides in (
        ("default", {}),
        ("template", {"body_first_line_indent_chars": 0, "body_space_after_pt": 12}),
    ):
        section = sample_section(tmp_path / name, overrides)
        output = export_section_docx(section, tmp_path / f"{name}.docx")
        document = docx.Document(output)
        non_body = [
            p
            for p in document.paragraphs
            if not p.text.startswith(("First paragraph", "Second paragraph"))
        ]
        assert any(p.style.name == "Title" for p in non_body)
        assert any(p.style.name == "Heading 2" for p in non_body)
        assert any(p.style.name == "Caption" for p in non_body)
        assert any(p._p.xpath(".//wp:inline") for p in non_body)
        assert any(p._p.xpath("./m:oMath") for p in non_body)
        assert any(p.text.startswith("[1]") for p in non_body)
        for paragraph in non_body:
            assert not paragraph._p.xpath("./w:pPr/w:ind")
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        assert not paragraph._p.xpath("./w:pPr/w:ind")
                        assert paragraph.paragraph_format.space_before.pt == 3
                        assert paragraph.paragraph_format.space_after.pt == 3
        with zipfile.ZipFile(output) as archive:
            styles = archive.read("word/styles.xml")
        report = json.loads(output.with_suffix(".layout.json").read_text(encoding="utf-8"))
        snapshots.append(
            (
                [p._p.xml for p in non_body],
                [table._tbl.xml for table in document.tables],
                styles,
                report["figures"],
            )
        )
    assert snapshots[0] == snapshots[1]


@pytest.mark.parametrize(
    "field", ["body_first_line_indent_chars", "body_space_before_pt", "body_space_after_pt"]
)
@pytest.mark.parametrize("value", [-1, float("nan"), float("inf")])
def test_invalid_body_format_is_rejected(field, value):
    with pytest.raises(ValueError):
        PublicationStyle(**{field: value})


def test_zero_figure_table_section_prepares_assembles_and_exports(tmp_path):
    docx = pytest.importorskip("docx")
    section = sample_section(tmp_path / "table-only", with_figure=False)
    output = export_section_docx(section, tmp_path / "table-only.docx")
    document = docx.Document(output)
    assert len(document.inline_shapes) == 0
    assert len(document.tables) == 1
    assert document.paragraphs[1].text == "A table-only discussion with [1]."
    assert document.paragraphs[1].paragraph_format.first_line_indent.pt == 22
    assert not list((section / "figures").iterdir())
    assert json.loads((section / "section.json").read_text())["figures"] == []


def test_zero_figures_do_not_permit_dangling_figure_references(tmp_path):
    sample_section(tmp_path / "sample", with_figure=False)
    source_path = tmp_path / "sample/source.json"
    source = json.loads(source_path.read_text())
    source["duties"][0]["figure_ids"] = ["missing"]
    source_path.write_text(json.dumps(source), encoding="utf-8")
    with pytest.raises(ValueError, match="Unresolved IDs"):
        prepare_section(source_path, tmp_path / "invalid-package")
