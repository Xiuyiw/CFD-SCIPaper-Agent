"""Bounded Word return examples use synthetic text and public OOXML structures."""

from copy import deepcopy

import pytest

docx = pytest.importorskip("docx")
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from cfdpaper.publication.word_return import (  # noqa: E402
    end_anchor,
    merge_word_blocks,
    read_word_return,
    start_anchor,
    wrap_anchor,
    write_word_state,
)


def block(identifier, text, **extra):
    return {
        "id": identifier,
        "kind": "paragraph",
        "section_id": "results",
        "local_id": identifier,
        "text": text,
        "rendered_text": text,
        "spans": [],
        **extra,
    }


def make_document(tmp_path, blocks=None):
    blocks = blocks or [block("p1", "A plain result."), block("p2", "Another result.")]
    document = docx.Document()
    for item in blocks:
        paragraph = document.add_paragraph()
        start_anchor(paragraph, item["id"])
        if item["spans"]:
            span = item["spans"][0]
            before, after = item["text"].split(span["token"])
            paragraph.add_run(before)
            start_anchor(paragraph, item["id"] + "::" + span["id"])
            paragraph.add_run(span["text"])
            end_anchor(paragraph, item["id"] + "::" + span["id"])
            paragraph.add_run(after)
        else:
            paragraph.add_run(item["rendered_text"])
        end_anchor(paragraph, item["id"])
    write_word_state(document, {"schema_version": 1, "baseline_id": "base-1", "blocks": blocks})
    path = tmp_path / "return.docx"
    document.save(path)
    return document, path


def paragraphs(document):
    return [Paragraph(p, document._body) for p in document._element.xpath(".//w:body//w:p")]


def test_open_save_split_runs_retains_semantics_and_unmapped_metadata(tmp_path):
    document, path = make_document(tmp_path)
    document = docx.Document(path)
    run = paragraphs(document)[0].runs[0]
    run.text = "A plain "
    second = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "result."
    second.append(text)
    run._r.addnext(second)
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"] == []
    assert result["baseline_id"] == "base-1"
    assert [b["text"] for b in result["blocks"]] == ["A plain result.", "Another result."]
    assert not document._element.xpath(".//w:dataBinding | .//w:lock")


def test_free_text_and_emphasis_are_recovered_without_scientific_edit(tmp_path):
    document, path = make_document(tmp_path)
    paragraphs(document)[0].runs[0].text = "An author result."
    paragraphs(document)[0].runs[0].bold = True
    paragraphs(document)[1].runs[0].italic = True
    paragraphs(document)[1].runs[0].font.superscript = True
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"] == []
    assert result["blocks"][0]["text"] == "An author result."
    assert result["blocks"][0]["formatting"][0]["bold"] is True
    assert result["blocks"][1]["formatting"][0]["superscript"] is True


def protected_block():
    return block(
        "p1",
        "Temperature is {{value:t}} K.",
        rendered_text="Temperature is 300 K.",
        spans=[{"id": "t", "token": "{{value:t}}", "text": "300"}],
    )


def test_unchanged_protected_span_restores_token_after_free_text_edit(tmp_path):
    document, path = make_document(tmp_path, [protected_block()])
    paragraphs(document)[0].runs[0].text = "Measured temperature is "
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"] == []
    assert result["blocks"][0]["text"] == "Measured temperature is {{value:t}} K."


def test_hand_edited_bound_value_is_a_located_conflict(tmp_path):
    document, path = make_document(tmp_path, [protected_block()])
    paragraphs(document)[0]._p.xpath(".//w:sdt//w:t")[0].text = "999"
    document.save(path)
    result = read_word_return(path)
    conflict = result["conflicts"][0]
    assert conflict["type"] == "protected_span_modified"
    assert conflict["block_id"] == "p1"
    assert conflict["span_id"] == "t"
    assert conflict["author"] == "999"
    assert result["blocks"][0]["text"] == protected_block()["text"]


def test_three_way_parallel_blocks_and_current_bound_value_are_kept():
    base = [protected_block(), block("p2", "Other result.")]
    author, current = deepcopy(base), deepcopy(base)
    author[0]["text"] = "Measured temperature is {{value:t}} K."
    author[0]["rendered_text"] = "Measured temperature is 300 K."
    current[0]["spans"][0]["text"] = "301"
    current[0]["rendered_text"] = "Temperature is 301 K."
    current[1]["text"] = current[1]["rendered_text"] = "Current result."
    result = merge_word_blocks(base, author, current)
    assert result["conflicts"] == []
    assert result["blocks"][0]["text"] == author[0]["text"]
    assert result["blocks"][0]["rendered_text"] == "Measured temperature is 301 K."
    assert result["blocks"][1]["text"] == "Current result."


def test_same_block_overlap_preserves_both_versions_and_deletion_conflicts():
    base = [block("p1", "The result is high.")]
    author = [block("p1", "The result is low.")]
    current = [block("p1", "The result is moderate.")]
    result = merge_word_blocks(base, author, current)
    assert result["conflicts"][0]["type"] == "overlapping_edit"
    assert result["conflicts"][0]["author"]["text"] == author[0]["text"]
    assert result["conflicts"][0]["current"]["text"] == current[0]["text"]
    assert merge_word_blocks(base, [], current)["conflicts"][0]["type"] == "delete_modify"


def test_same_block_nonoverlapping_edits_merge():
    base = [block("p1", "Alpha has one result; beta has another result.")]
    author = [block("p1", "Alpha reports one result; beta has another result.")]
    current = [block("p1", "Alpha has one result; beta has two results.")]
    result = merge_word_blocks(base, author, current)
    assert result["conflicts"] == []
    assert result["blocks"][0]["text"] == "Alpha reports one result; beta has two results."


def test_comments_remain_annotations_with_block_and_text_range(tmp_path):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from lxml import etree

    document, path = make_document(tmp_path)
    run = paragraphs(document)[0].runs[0]._r
    for tag, insert in [("commentRangeStart", run.addprevious), ("commentRangeEnd", run.addnext)]:
        marker = OxmlElement("w:" + tag)
        marker.set(qn("w:id"), "0")
        insert(marker)
    reference, marker = OxmlElement("w:r"), OxmlElement("w:commentReference")
    marker.set(qn("w:id"), "0")
    reference.append(marker)
    run.getnext().addnext(reference)
    root, comment = OxmlElement("w:comments"), OxmlElement("w:comment")
    comment.set(qn("w:id"), "0")
    comment.set(qn("w:author"), "A")
    paragraph, comment_run, text = OxmlElement("w:p"), OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "Check this interpretation."
    comment_run.append(text)
    paragraph.append(comment_run)
    comment.append(paragraph)
    root.append(comment)
    part = Part(
        PackURI("/word/comments.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
        etree.tostring(root),
        document.part.package,
    )
    document.part.relate_to(part, RT.COMMENTS)
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"] == []
    assert result["comments"][0]["author"] == "A"
    assert result["comments"][0]["text"] == "Check this interpretation."
    assert result["comments"][0]["block_id"] == "p1"
    assert result["comments"][0]["anchor_text"] == "A plain result."
    assert result["blocks"][0]["text"] == "A plain result."


def test_unaccepted_revisions_do_not_become_accepted_text(tmp_path):
    document, path = make_document(tmp_path)
    run = paragraphs(document)[0].runs[0]
    insertion = OxmlElement("w:ins")
    insertion.set(qn("w:id"), "42")
    insertion.set(qn("w:author"), "A")
    run._r.addprevious(insertion)
    insertion.append(run._r)
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"][0]["type"] == "unresolved_revision"
    assert result["conflicts"][0]["block_id"] == "p1"
    assert "<w:ins" in result["conflicts"][0]["xml"]
    assert result["blocks"][0]["text"] == "A plain result."


@pytest.mark.parametrize("change", ["missing", "duplicate"])
def test_missing_and_duplicate_anchors_are_real_conflicts(tmp_path, change):
    document, path = make_document(tmp_path)
    anchor = document._element.xpath(".//w:sdt")[0]
    if change == "missing":
        anchor.getparent().remove(anchor)
    else:
        anchor.addnext(deepcopy(anchor))
    document.save(path)
    result = read_word_return(path)
    assert any(
        c["type"] == f"{change}_anchor" and c["block_id"] == "p1" for c in result["conflicts"]
    )


def test_unanchored_plain_paragraph_gets_neighbor_identity_and_moves_keep_ids(tmp_path):
    document, path = make_document(tmp_path)
    inserted = document.add_paragraph("A new author observation.")
    paragraphs(document)[0]._p.getparent().getparent().addnext(inserted._p)
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"] == []
    assert [b["text"] for b in result["blocks"]] == [
        "A plain result.",
        "A new author observation.",
        "Another result.",
    ]
    new = result["blocks"][1]
    assert new["id"] not in {"p1", "p2"}
    assert new["section_id"] == "results"
    assert new["neighbors"] == {"before": "p1", "after": "p2"}
    first, last = document._element.xpath(".//w:body/w:sdt")
    first.addprevious(last)
    document.save(path)
    moved = read_word_return(path)
    assert [b["id"] for b in moved["blocks"] if b["id"] in {"p1", "p2"}] == ["p2", "p1"]


def test_formula_content_change_is_not_flattened_to_free_prose(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph()
    start_anchor(paragraph, "eq1")
    formula = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "q"
    run.append(text)
    formula.append(run)
    paragraph._p.append(formula)
    end_anchor(paragraph, "eq1")
    original = block("eq1", "q", kind="equation")
    write_word_state(document, {"schema_version": 1, "baseline_id": "b", "blocks": [original]})
    text.text = "Q"
    path = tmp_path / "math.docx"
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"][0]["type"] == "object_modified"
    assert result["conflicts"][0]["object_kind"] == "math"
    assert result["blocks"][0]["text"] == "q"


def test_field_change_is_located_and_unmodified_field_is_kept(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph()
    start_anchor(paragraph, "p1")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " ADDIN CITATION example ")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "[1]"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    end_anchor(paragraph, "p1")
    state = {"schema_version": 1, "baseline_id": "b", "blocks": [block("p1", "[1]")]}
    write_word_state(document, state)
    path = tmp_path / "field.docx"
    document.save(path)
    assert read_word_return(path)["conflicts"] == []
    field.set(qn("w:instr"), " ADDIN CITATION different ")
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"][0]["object_kind"] == "field"


def test_simple_table_header_labels_recover_but_rows_and_structure_conflict(tmp_path):
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [["Item", "Value"], ["A", "one"]], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    wrap_anchor(table, "t1")
    original = block(
        "t1", "", kind="table", content={"columns": ["Item", "Value"], "rows": [["A", "one"]]}
    )
    write_word_state(document, {"schema_version": 1, "baseline_id": "b", "blocks": [original]})
    path = tmp_path / "table.docx"
    document.save(path)
    assert read_word_return(path)["conflicts"] == []
    table.cell(0, 1).text = "Description"
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"] == []
    assert result["blocks"][0]["content"]["columns"] == ["Item", "Description"]
    table.cell(1, 1).text = "two"
    document.save(path)
    assert read_word_return(path)["conflicts"][0]["type"] == "table_structure_or_text_modified"
    table.add_row()
    document.save(path)
    assert read_word_return(path)["conflicts"][0]["type"] == "table_structure_or_text_modified"


def test_inline_image_resize_preserves_content_but_replacement_conflicts(tmp_path):
    from PIL import Image

    first, second = tmp_path / "first.png", tmp_path / "second.png"
    Image.new("RGB", (20, 10), "white").save(first)
    Image.new("RGB", (20, 10), "black").save(second)
    document = docx.Document()
    paragraph = document.add_paragraph()
    start_anchor(paragraph, "f1")
    picture = paragraph.add_run().add_picture(str(first))
    end_anchor(paragraph, "f1")
    original = block("f1", "", kind="figure")
    write_word_state(document, {"schema_version": 1, "baseline_id": "b", "blocks": [original]})
    path = tmp_path / "picture.docx"
    document.save(path)
    assert read_word_return(path)["conflicts"] == []
    picture.width, picture.height = picture.width * 2, picture.height * 2
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"] == []
    assert result["blocks"][0]["layout"][0]["cx"] == picture.width
    picture._inline.xpath(".//a:blip")[0].set(
        qn("r:embed"), document.part.get_or_add_image(str(second))[0]
    )
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"][0]["object_kind"] == "image"


def test_caption_generated_prefix_is_not_imported_as_semantic_text(tmp_path):
    original = block("cap1", "An original caption.", kind="caption")
    document = docx.Document()
    paragraph = document.add_paragraph()
    start_anchor(paragraph, "cap1")
    paragraph.add_run("Figure 1. ")
    paragraph.add_run(original["text"])
    end_anchor(paragraph, "cap1")
    write_word_state(document, {"schema_version": 1, "baseline_id": "b", "blocks": [original]})
    paragraph.runs[1].text = "An author caption."
    path = tmp_path / "caption.docx"
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"] == []
    assert result["blocks"][0]["text"] == "An author caption."
    assert result["blocks"][0]["rendered_text"] == "An author caption."
    paragraph.runs[1].bold = True
    document.save(path)
    formatted = read_word_return(path)["blocks"][0]
    assert formatted["formatting_source_text"] == "Figure 1. An author caption."
    paragraph.runs[0].text = "Figure 99. "
    document.save(path)
    assert read_word_return(path)["conflicts"][0]["type"] == "generated_label_modified"


def test_default_formatting_is_not_an_author_override_after_open_save(tmp_path):
    _, path = make_document(tmp_path)
    result = read_word_return(path)
    assert not any("formatting" in b or "paragraph_format" in b for b in result["blocks"])


def test_enter_inside_plain_sdt_assigns_new_paragraph_identity(tmp_path):
    from cfdpaper.publication.document_blocks import stable_block_id

    document, path = make_document(tmp_path)
    first = paragraphs(document)[0]
    inserted = OxmlElement("w:p")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "An inserted paragraph inside the content control."
    run.append(text)
    inserted.append(run)
    first._p.addnext(inserted)
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"] == []
    assert result["blocks"][0]["id"] == "p1"
    assert result["blocks"][0]["text"] == "A plain result."
    assert result["blocks"][1]["text"] == text.text
    assert result["blocks"][1]["neighbors"] == {"before": "p1", "after": "p2"}
    assert result["blocks"][1]["id"] == stable_block_id(
        "results", "paragraph", result["blocks"][1]["local_id"]
    )


def test_math_run_splitting_is_not_a_formula_edit(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph()
    start_anchor(paragraph, "eq1")
    formula = OxmlElement("m:oMath")
    run, text = OxmlElement("m:r"), OxmlElement("m:t")
    text.text = "ab"
    run.append(text)
    formula.append(run)
    paragraph._p.append(formula)
    end_anchor(paragraph, "eq1")
    write_word_state(
        document,
        {"schema_version": 1, "baseline_id": "b", "blocks": [block("eq1", "ab", kind="equation")]},
    )
    text.text = "a"
    second_run, second_text = OxmlElement("m:r"), OxmlElement("m:t")
    second_text.text = "b"
    second_run.append(second_text)
    formula.append(second_run)
    path = tmp_path / "split-math.docx"
    document.save(path)
    assert read_word_return(path)["conflicts"] == []


def test_changed_generated_title_is_not_promoted_to_a_section_paragraph(tmp_path):
    document = docx.Document()
    title = document.add_paragraph("A generated title", "Title")
    heading = document.add_heading("Methods", level=1)
    start_anchor(heading, "heading")
    end_anchor(heading, "heading")
    paragraph = document.add_paragraph("A methods paragraph.")
    start_anchor(paragraph, "p1")
    end_anchor(paragraph, "p1")
    state = {
        "schema_version": 1,
        "baseline_id": "b",
        "blocks": [
            block("heading", "Methods", kind="heading"),
            block("p1", "A methods paragraph."),
        ],
    }
    write_word_state(document, state)
    title.runs[0].text = "An authored title"
    path = tmp_path / "title.docx"
    document.save(path)
    result = read_word_return(path)
    assert [b["id"] for b in result["blocks"]] == ["heading", "p1"]
    assert result["conflicts"][0]["type"] == "generated_title_modified"


def test_numeric_recompilation_without_visual_projection_does_not_erase_author_format():
    original = protected_block()
    original["formatting"] = [{"start": 0, "end": 11, "bold": True}]
    original["formatting_source_text"] = original["rendered_text"]
    current = deepcopy(original)
    current.pop("formatting")
    current.pop("formatting_source_text")
    current["spans"][0]["text"] = "301"
    current["rendered_text"] = "Temperature is 301 K."
    result = merge_word_blocks([original], [deepcopy(original)], [current])
    assert result["conflicts"] == []
    assert result["blocks"][0]["formatting"] == original["formatting"]
    assert result["blocks"][0]["spans"][0]["text"] == "301"


def test_cross_section_move_is_not_silently_grouped_back_to_old_section(tmp_path):
    blocks = [
        block("h1", "Methods", kind="heading", section_id="methods"),
        block("p1", "A definition.", section_id="methods"),
        block("h2", "Results", kind="heading", section_id="results"),
        block("p2", "A result."),
    ]
    document, path = make_document(tmp_path, blocks)
    controls = document._element.xpath(".//w:body/w:sdt")
    controls[3].addnext(controls[1])
    document.save(path)
    result = read_word_return(path)
    assert any(
        c["type"] == "cross_section_move" and c["block_id"] == "p1" for c in result["conflicts"]
    )


@pytest.mark.parametrize("legacy_baseline", [False, True])
def test_word_localized_style_ids_are_not_author_style_changes(tmp_path, legacy_baseline):
    import json

    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from lxml import etree

    document = docx.Document()
    heading = document.add_heading("Methods", level=1)
    start_anchor(heading, "h1")
    end_anchor(heading, "h1")
    write_word_state(
        document,
        {
            "schema_version": 1,
            "baseline_id": "b",
            "blocks": [block("h1", "Methods", kind="heading")],
        },
    )
    if legacy_baseline:
        part = next(
            rel.target_part
            for rel in document.part.rels.values()
            if rel.reltype == RT.CUSTOM_XML
            and b"urn:cfdpaper:word-return:1" in rel.target_part.blob
        )
        root = etree.fromstring(part.blob)
        payload = root[0]
        state = json.loads(payload.text)
        state["_word_baseline"]["h1"]["paragraph_format"][0]["pStyle"] = {"val": "Heading1"}
        payload.text = json.dumps(state)
        part._blob = etree.tostring(root)
    document.styles["Heading 1"].element.set(qn("w:styleId"), "1")
    heading._p.pPr.pStyle.val = "1"
    path = tmp_path / "localized-style.docx"
    document.save(path)
    result = read_word_return(path)
    assert result["conflicts"] == []
    assert "paragraph_format" not in result["blocks"][0]
    assert result["state"]["_word_baseline"]["h1"]["paragraph_format"][0]["pStyle"] == {
        "name": "Heading 1"
    }
    document.styles["Heading 2"].element.set(qn("w:styleId"), "21")
    heading._p.pPr.pStyle.val = "21"
    document.save(path)
    returned = read_word_return(path)
    assert returned["blocks"][0]["paragraph_format"][0]["pStyle"] == {"name": "Heading 2"}
