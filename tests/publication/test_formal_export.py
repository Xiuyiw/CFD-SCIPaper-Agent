"""Formal views share ordered scientific content; generated files are generic fixtures."""

import json
import zipfile

import pytest
from PIL import Image

from cfdpaper.publication.formal_export import export_formal_manuscript, math_latex


def test_latex_caption_stays_with_its_float_and_tables_wrap(tmp_path):
    root = fixture(tmp_path)
    path = export_formal_manuscript(root, tmp_path / "conference.tex", profile="conference")
    text = path.read_text(encoding="utf-8")
    a = text.index(r"\begin{figure*}")
    b = text.index(r"\caption*{Figure", a)
    c = text.index(r"\end{figure*}", a)
    assert a < b < c
    assert r"\begin{tabularx}{\linewidth}" in text


def fixture(tmp_path):
    root = tmp_path / "source"
    root.mkdir()
    Image.new("RGB", (400, 200), "white").save(root / "field.png")
    data = {
        "title": "Conjugate heat transport",
        "section_id": "methods",
        "paragraphs": [
            {
                "paragraph_id": "opening",
                "text": "Interface temperature is 300 K.",
                "source_text": "Interface temperature is {{value:temperature}}.",
                "semantic_runs": [
                    {"text": "Interface temperature is "},
                    {"text": "300 K", "token": "{{value:temperature}}", "span_id": "temperature"},
                    {"text": "."},
                ],
                "figure_ids": [],
                "evidence_ids": [],
            },
            {
                "paragraph_id": "explanation",
                "text": "The ratio is dimensionless.",
                "figure_ids": ["1"],
                "evidence_ids": [],
            },
        ],
        "equations": [
            {
                "equation_id": "1",
                "expression": {
                    "kind": "fraction",
                    "children": [
                        {
                            "kind": "sub",
                            "children": [
                                {"kind": "symbol", "text": "q"},
                                {"kind": "text", "text": "out"},
                            ],
                        },
                        {"kind": "symbol", "text": "Q"},
                    ],
                },
                "evidence_ids": [],
            }
        ],
        "tables": [
            {
                "table_id": "1",
                "caption": "Controlled thermal input",
                "columns": ["Case", "Power (W)"],
                "rows": [["A", "25"]],
                "after_section_id": "methods",
                "numeric_columns": [1],
                "note": "Same prescribed input.",
            }
        ],
        "figures": [{"id": "1", "path": "field.png", "caption": "Interface temperature field."}],
        "block_order": [
            "paragraph:opening",
            "equation:1",
            "paragraph:explanation",
            "table:1",
            "figure:1",
        ],
        "references": [{"id": "source", "number": 1, "text": "Author. Thermal transport. 2020."}],
    }
    (root / "section.json").write_text(json.dumps(data), encoding="utf-8")
    return root


def test_markdown_tex_share_order_and_values_with_portable_assets(tmp_path):
    root = fixture(tmp_path)
    md = export_formal_manuscript(root, tmp_path / "export" / "paper.md")
    tex = export_formal_manuscript(root, tmp_path / "export" / "paper.tex")
    markdown, latex = md.read_text(encoding="utf-8"), tex.read_text(encoding="utf-8")
    for text in (markdown, latex):
        assert "300 K" in text and "25" in text
        assert text.index("temperature is") < text.index("frac")
        assert text.index("frac") < text.index("ratio is") < text.index("Controlled thermal")
        assert "Figure 1. Interface temperature field." in text
        assert "[1] Author. Thermal transport. 2020." in text
        assert str(tmp_path) not in text
    assert "\\frac" in latex and r"\text{out}" in latex
    assert (
        r"\end{tabularx}" + "\n" + r"\par\smallskip\noindent\small Same prescribed input.\par"
    ) in latex
    assert list((md.parent / "paper_md_assets").glob("*.png"))


def test_docx_has_native_math_table_order_and_editable_anchors(tmp_path):
    root = fixture(tmp_path)
    out = export_formal_manuscript(root, tmp_path / "paper.docx")
    with zipfile.ZipFile(out) as archive:
        xml = archive.read("word/document.xml").decode()
        styles = archive.read("word/styles.xml").decode()
        state = archive.read("customXml/cfdpaper-word-state.xml").decode()
    assert "<m:f>" in xml and "<w:tbl>" in xml
    assert "w:tblHeader" in xml and "<w:sdt>" in xml
    assert "w:lock" not in xml and "w:dataBinding" not in xml
    assert xml.index("temperature is") < xml.index("<m:f>") < xml.index("ratio is")
    assert 'firstLineChars="200"' in xml
    assert 'firstLineChars="200"' not in styles
    assert "{{value:temperature}}" in state
    from cfdpaper.publication.word_return import read_word_return

    returned = read_word_return(out)
    assert not returned["conflicts"]
    paragraph = next(b for b in returned["blocks"] if b["local_id"] == "opening")
    assert paragraph["text"] == "Interface temperature is {{value:temperature}}."


@pytest.mark.parametrize("profile,pt", [("research", 11), ("report", 12), ("conference", 10)])
def test_profiles_change_presentation_not_scientific_facts(tmp_path, profile, pt):
    from docx import Document

    root = fixture(tmp_path)
    out = export_formal_manuscript(root, tmp_path / f"{profile}.docx", profile=profile)
    doc = Document(out)
    assert doc.styles["Normal"].font.size.pt == pt
    text = "".join(t.text or "" for t in doc._element.xpath(".//w:t"))
    assert "300\u00a0K" in text and "25" in text
    assert "before" not in text


def test_formal_body_spacing_does_not_indent_figures_or_equations(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    root = fixture(tmp_path)
    path = root / "section.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    data["style"] = {"body_space_before_pt": 6, "body_space_after_pt": 12}
    path.write_text(json.dumps(data), encoding="utf-8")
    output = export_formal_manuscript(root, tmp_path / "spaced.docx")
    document = Document(output)
    paragraphs = [Paragraph(p, document._body) for p in document._element.body.xpath(".//w:p")]
    body = next(p for p in paragraphs if "temperature is" in p._p.xml)
    assert body.paragraph_format.space_before.pt == 6
    assert body.paragraph_format.space_after.pt == 12
    assert body.paragraph_format.first_line_indent.pt == 22
    assert document.styles["Normal"].paragraph_format.first_line_indent is None
    non_body = [p for p in paragraphs if p._p.xpath(".//wp:inline|.//m:oMath")]
    assert len(non_body) == 2
    for paragraph in non_body:
        indent = paragraph._p.pPr.find(qn("w:ind"))
        assert indent is None or indent.get(qn("w:firstLineChars")) in {None, "0"}
        assert paragraph.paragraph_format.space_before is None
        assert paragraph.paragraph_format.space_after is None


def test_formal_nonbody_indent_still_accepts_author_override(tmp_path):
    from docx import Document
    from docx.text.paragraph import Paragraph

    root = fixture(tmp_path)
    path = root / "section.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    data["equations"][0]["paragraph_format"] = {"first_line_indent_pt": 9}
    data["figures"][0]["paragraph_format"] = {"first_line_indent_pt": 7}
    path.write_text(json.dumps(data), encoding="utf-8")
    output = export_formal_manuscript(root, tmp_path / "author-indent.docx")
    document = Document(output)
    paragraphs = [Paragraph(p, document._body) for p in document._element.body.xpath(".//w:p")]
    equation = next(p for p in paragraphs if p._p.xpath(".//m:oMath"))
    figure = next(p for p in paragraphs if p._p.xpath(".//wp:inline"))
    assert equation.paragraph_format.first_line_indent.pt == 9
    assert figure.paragraph_format.first_line_indent.pt == 7


def test_explicit_word_template_preserves_page_and_normal_style(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    root = fixture(tmp_path)
    template = tmp_path / "template.docx"
    document = Document()
    document.styles["Normal"].font.size = Pt(13)
    document.sections[0].left_margin = Mm(30)
    document.sections[0]._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    document.sections[0].header.paragraphs[0].text = "Conference header"
    document.add_paragraph("Replace this placeholder")
    document.save(template)
    original = template.read_bytes()
    out = export_formal_manuscript(
        root, tmp_path / "paper.docx", profile="conference", template_path=template
    )
    actual = Document(out)
    assert actual.styles["Normal"].font.size.pt == 13
    assert actual.sections[0].left_margin.mm == pytest.approx(30, abs=0.05)
    assert actual.sections[0].header.paragraphs[0].text == "Conference header"
    report = json.loads(out.with_suffix(".layout.json").read_text(encoding="utf-8"))
    assert report["figures"][0]["width_mm"] <= report["column_width_mm"]
    assert report["column_width_mm"] < 90
    assert "Replace this placeholder" not in actual._element.xml
    assert template.read_bytes() == original


def test_word_template_uses_existing_named_roles_and_last_section_geometry(tmp_path):
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    root = fixture(tmp_path)
    document = Document()
    document.styles["Title"].name = "paper title"
    document.styles["paper title"].element.set(qn("w:styleId"), "papertitle")
    document.styles["paper title"].font.size = Pt(24)
    document.styles["Caption"].name = "figure caption"
    document.styles["figure caption"].element.set(qn("w:styleId"), "figurecaption")
    document.styles["figure caption"].font.size = Pt(8)
    numbers = OxmlElement("w:numPr")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), "9")
    numbers.append(number)
    document.styles["figure caption"].element.get_or_add_pPr().append(numbers)
    for option in document.settings.element.findall(".//" + qn("w:compatSetting")):
        if option.get(qn("w:name")) == "compatibilityMode":
            option.set(qn("w:val"), "11")
    document.styles["Body Text"].font.size = Pt(10)
    document.add_paragraph("Replace the actual body slot", style="Body Text")
    document.add_section(WD_SECTION_START.CONTINUOUS).left_margin = Mm(32)
    template = tmp_path / "role-template.docx"
    document.save(template)
    original = template.read_bytes()
    out = export_formal_manuscript(root, tmp_path / "report.docx", template_path=template)
    actual = Document(out)
    assert len(actual.sections) == 1
    assert actual.sections[0].left_margin.mm == pytest.approx(32, abs=0.05)
    assert actual.paragraphs[0].style.name == "paper title"
    with zipfile.ZipFile(out) as archive:
        xml = archive.read("word/document.xml").decode()
    assert 'w:val="BodyText"' in xml
    assert "<m:f>" in xml and "<w:tbl>" in xml and "w:tblHeader" in xml
    assert actual.styles["paper title"].font.size.pt == 24
    assert actual.styles["figure caption"].font.size.pt == 8
    for option in actual.settings.element.findall(".//" + qn("w:compatSetting")):
        if option.get(qn("w:name")) == "compatibilityMode":
            assert option.get(qn("w:val")) == "15"
    caption_ids = {actual.styles[name].style_id for name in ("Caption", "figure caption")}
    for paragraph in actual._element.body.iter(qn("w:p")):
        style = paragraph.find("./" + qn("w:pPr") + "/" + qn("w:pStyle"))
        if style is not None and style.get(qn("w:val")) in caption_ids:
            number = paragraph.find("./" + qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId"))
            assert number is not None and number.get(qn("w:val")) == "0"
    assert "Replace the actual body slot" not in xml
    report = json.loads(out.with_suffix(".layout.json").read_text(encoding="utf-8"))
    assert report["style"]["body_pt"] == 10
    assert report["style"]["title_pt"] == 24
    assert report["style"]["caption_pt"] == 8
    assert report["template_style_roles"]["Title"] == "paper title"
    assert template.read_bytes() == original


def test_existing_author_outputs_or_assets_are_not_overwritten(tmp_path):
    root = fixture(tmp_path)
    output = tmp_path / "paper.docx"
    output.write_bytes(b"author")
    with pytest.raises(FileExistsError):
        export_formal_manuscript(root, output)
    assert output.read_bytes() == b"author"
    (tmp_path / "paper_tex_assets").mkdir()
    with pytest.raises(FileExistsError):
        export_formal_manuscript(root, tmp_path / "paper.tex")


def test_tex_template_requires_explicit_body_marker_and_keeps_template(tmp_path):
    root = fixture(tmp_path)
    template = tmp_path / "template.tex"
    template.write_text(r"\documentclass{article}" + "\n%%CFDPAPER_BODY%%\n", encoding="utf-8")
    out = export_formal_manuscript(root, tmp_path / "paper.tex", template_path=template)
    assert "300 K" in out.read_text(encoding="utf-8")
    assert "%%CFDPAPER_BODY%%" in template.read_text(encoding="utf-8")
    assert math_latex({"kind": "sqrt", "children": [{"kind": "symbol", "text": "α"}]}) == (
        r"\sqrt{\alpha}"
    )


def test_latex_scientific_unicode_is_not_left_as_unsupported_inputenc_text(tmp_path):
    root = fixture(tmp_path)
    data = json.loads((root / "section.json").read_text(encoding="utf-8"))
    data["paragraphs"][1]["text"] = "Flux is −2 × 10 W m⁻² at 30 °C; CO₂ is retained."
    (root / "section.json").write_text(json.dumps(data), encoding="utf-8")
    out = export_formal_manuscript(root, tmp_path / "paper.tex")
    text = out.read_text(encoding="utf-8")
    assert r"m\textsuperscript{-2}" in text
    assert r"CO\textsubscript{2}" in text
    assert "−" not in text and "⁻" not in text


def test_author_formatting_comments_survive_changed_numeric_display(tmp_path):
    from cfdpaper.publication.document_blocks import build_document
    from cfdpaper.publication.word_return import read_word_return

    root = fixture(tmp_path)
    data = json.loads((root / "section.json").read_text(encoding="utf-8"))
    item = next(block for block in build_document(data)["blocks"] if block["local_id"] == "opening")
    old = "Interface temperature is 99\u00a0K."
    edits = {
        "blocks": {
            item["id"]: {
                "formatting": [
                    {"start": old.index("99"), "end": old.index("99") + 2, "bold": True}
                ],
                "formatting_source_text": old,
                "paragraph_format": [{"jc": {"val": "both"}, "spacing": {"after": "0"}}],
            }
        },
        "comments": [
            {
                "block_id": item["id"],
                "text": "Clarify interface sampling.",
                "author": "Research author",
                "start": 0,
                "end": 9,
                "anchor_text": "Interface",
            }
        ],
    }
    (root / "author-edits.json").write_text(json.dumps(edits), encoding="utf-8")
    out = export_formal_manuscript(root, tmp_path / "paper.docx")
    result = read_word_return(out)
    assert not result["conflicts"]
    comments = result["comments"]
    assert comments[0]["text"] == "Clarify interface sampling."
    assert comments[0]["anchor_text"] == "Interface"
    returned = next(b for b in result["blocks"] if b["id"] == item["id"])
    bold = next(p for p in returned["formatting"] if p.get("bold"))
    assert returned["rendered_text"][bold["start"] : bold["end"]] == "300"
    assert returned["paragraph_format"][0]["jc"]["val"] == "both"
    report = json.loads(out.with_suffix(".layout.json").read_text(encoding="utf-8"))
    assert not report["formatting_conflicts"]
    assert not report["comment_conflicts"]


def test_ambiguous_comments_and_character_ranges_remain_localized(tmp_path):
    from cfdpaper.publication.document_blocks import build_document

    root = fixture(tmp_path)
    data = json.loads((root / "section.json").read_text(encoding="utf-8"))
    block = next(b for b in build_document(data)["blocks"] if b["local_id"] == "opening")
    old = "Interface temperature is 9999 K."
    edits = {
        "blocks": {
            block["id"]: {
                "formatting_source_text": old,
                "formatting": [
                    {"start": old.index("9999") + 1, "end": old.index("9999") + 2, "italic": True}
                ],
            }
        },
        "comments": [
            {
                "block_id": block["id"],
                "text": "Retain this note.",
                "anchor_text": "prior missing phrase",
                "start": 0,
                "end": 3,
            }
        ],
    }
    (root / "author-edits.json").write_text(json.dumps(edits), encoding="utf-8")
    out = export_formal_manuscript(root, tmp_path / "paper.docx")
    report = json.loads(out.with_suffix(".layout.json").read_text(encoding="utf-8"))
    assert report["formatting_conflicts"][0]["block_id"] == block["id"]
    assert report["comment_conflicts"][0]["block_id"] == block["id"]
    with zipfile.ZipFile(out) as archive:
        assert "Retain this note." in archive.read("word/comments.xml").decode()


def test_local_figure_label_survives_object_renumber_without_fixed_scientific_text(tmp_path):
    from cfdpaper.publication.document_blocks import build_document
    from cfdpaper.publication.formal_export import _prefix, _runs

    root = fixture(tmp_path)
    data = json.loads((root / "section.json").read_text(encoding="utf-8"))
    paragraph = data["paragraphs"][1]
    paragraph.update(
        {
            "text": "See Figure 3.",
            "source_text": "See {{figure:1}}.",
            "semantic_runs": [
                {"text": "See "},
                {"text": "Figure 3", "token": "{{figure:1}}", "span_id": "figure-reference"},
                {"text": "."},
            ],
        }
    )
    data["figures"][0]["id"] = "3"
    paragraph["figure_ids"] = ["3"]
    data["object_origins"] = {"figure": {"3": {"section_id": "methods", "local_id": "1"}}}
    data.pop("block_order")
    first = build_document(data)
    edited = [
        b for b in first["blocks"] if b["kind"] == "caption" or b["local_id"] == "explanation"
    ]
    labels = {b["id"]: {"reference_labels": {"figure": "Fig."}} for b in edited}
    # Global numbering changes; local object identity is stable.
    data["object_origins"] = {"figure": {"4": {"section_id": "methods", "local_id": "1"}}}
    data["figures"][0]["id"] = "4"
    paragraph["text"] = "See Figure 4."
    paragraph["figure_ids"] = ["4"]
    paragraph["semantic_runs"][1]["text"] = "Figure 4"
    (root / "section.json").write_text(json.dumps(data), encoding="utf-8")
    (root / "author-edits.json").write_text(json.dumps({"blocks": labels}), encoding="utf-8")
    output = export_formal_manuscript(root, tmp_path / "renumbered.docx")
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "Fig. 4" in xml
    assert "Fig. 3" not in xml
    untouched = next(b for b in build_document(data)["blocks"] if b["kind"] == "caption")
    assert _prefix(untouched) == "Figure 4. "
    assert _runs(untouched)[0]["text"] == "Interface temperature field."
