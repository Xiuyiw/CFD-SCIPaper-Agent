import pytest

from cfdpaper.publication.elements import MathNode, SectionTable, math_text, math_xml


def test_numbered_markdown_table_matches_prose_and_docx_caption():
    from cfdpaper.publication.export import _markdown_table

    table = SectionTable(
        table_id="2",
        caption="Spatial diagnostics",
        after_section_id="results",
        columns=["Field", "Mean"],
        rows=[["A", "43.4 °C"]],
        evidence_ids=[],
    )
    assert _markdown_table(table, numbered=True).startswith("**Table 2. Spatial diagnostics**")
    assert _markdown_table(table).startswith("**Spatial diagnostics**")


def symbol(text):
    return {"kind": "symbol", "text": text}


@pytest.mark.parametrize(
    "kind,count,tag",
    [
        ("sub", 2, "sSub"),
        ("sup", 2, "sSup"),
        ("subsup", 3, "sSubSup"),
        ("fraction", 2, "f"),
        ("sqrt", 1, "rad"),
        ("dot", 1, "acc"),
        ("overbar", 1, "acc"),
    ],
)
def test_math_nodes_are_editable_structures(kind, count, tag):
    pytest.importorskip("docx")
    node = MathNode(kind=kind, children=[symbol("x")] * count)
    xml = math_xml(node)[0]
    assert xml.tag.endswith("}" + tag)
    assert math_text(node)
    with pytest.raises(ValueError):
        MathNode(kind=kind, children=[symbol("x")] * (count + 1))


@pytest.mark.parametrize(
    "node",
    [
        {"kind": "latex", "text": "x"},
        {"kind": "symbol"},
        {"kind": "row", "text": "discarded", "children": [symbol("x")]},
    ],
)
def test_unsupported_or_ambiguous_math_is_not_silently_rendered(node):
    with pytest.raises(ValueError):
        MathNode.model_validate(node)


def test_table_retains_rectangular_semantics_and_widths():
    with pytest.raises(ValueError, match="column count"):
        SectionTable(
            table_id="1", caption="Data", columns=["a", "b"], rows=[["1"]], after_section_id="s1"
        )
    with pytest.raises(ValueError, match="widths"):
        SectionTable(
            table_id="1",
            caption="Data",
            columns=["a"],
            rows=[["1"]],
            after_section_id="s1",
            column_widths_mm=[float("nan")],
        )


def test_nested_flow_accents_and_table_note_pagination():
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn

    from cfdpaper.publication.elements import add_table
    from cfdpaper.publication.style import PublicationStyle

    node = MathNode(kind="overbar", children=[MathNode(kind="dot", children=[symbol("m")])])
    xml = math_xml(node)[0]
    assert [e.get(qn("m:val")) for e in xml.iter(qn("m:chr"))] == ["\u0305", "\u0307"]
    assert [e.text for e in xml.iter(qn("m:t"))] == ["m"]
    document = docx.Document()
    add_table(
        document,
        SectionTable(
            table_id="1",
            caption="Data",
            columns=["x"],
            rows=[["1"], ["2"]],
            after_section_id="s",
            note="Definition.",
        ),
        PublicationStyle(),
    )
    table = document.tables[0]
    assert table.rows[-1]._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None
    assert table.rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next is True
    assert table.rows[-2].cells[0].paragraphs[0].paragraph_format.keep_with_next is True
    assert document.paragraphs[-1].text == "Definition."


@pytest.mark.parametrize("value", ["Re", "Im", "Nu"])
def test_multi_letter_symbols_are_not_importer_keywords(value):
    pytest.importorskip("docx")
    from docx.oxml.ns import qn

    node = MathNode(kind="sub", children=[symbol(value), symbol("D")])
    xml = math_xml(node)[0]
    base = xml.find(qn("m:e"))
    assert [item.text for item in base.iter(qn("m:t"))] == list(value)
    assert not list(base.iter(qn("m:nor")))
    assert math_text(node) == f"{value}_(D)"


@pytest.mark.parametrize("value", ["(", ")", "mean (local)"])
def test_text_parentheses_do_not_receive_normal_text_escaping(value):
    pytest.importorskip("docx")
    from docx.oxml.ns import qn

    node = MathNode(kind="text", text=value)
    runs = math_xml(node)
    assert "".join(run.find(qn("m:t")).text for run in runs) == value
    for run in runs:
        text = run.find(qn("m:t")).text
        assert "\\" not in text
        assert bool(list(run.iter(qn("m:nor")))) is (text not in {"(", ")"})
    assert math_text(node) == value


@pytest.mark.parametrize("rows,compact", [(6, True), (60, False)])
@pytest.mark.parametrize("note", ["", "Definition."])
def test_compact_tables_keep_all_rows_without_chaining_following_text(rows, compact, note):
    docx = pytest.importorskip("docx")

    from cfdpaper.publication.elements import add_table
    from cfdpaper.publication.style import PublicationStyle

    document = docx.Document()
    add_table(
        document,
        SectionTable(
            table_id="1",
            caption="Data",
            columns=["Case", "Value"],
            rows=[[f"Case {i}", str(i)] for i in range(rows)],
            after_section_id="s",
            note=note,
        ),
        PublicationStyle(),
    )
    table = document.tables[0]
    assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    for i, row in enumerate(table.rows):
        assert row._tr.xpath("./w:trPr/w:cantSplit")
        expected = bool(note) if i == rows else compact or i == 0
        for cell in row.cells:
            assert cell.paragraphs[0].paragraph_format.keep_with_next is expected
    tail = document.add_paragraph("Following text.")
    assert tail.paragraph_format.keep_with_next is None
    assert [cell.text for cell in table.rows[-1].cells] == [f"Case {rows - 1}", str(rows - 1)]


def test_few_rows_with_long_wrapped_text_are_not_forced_onto_one_page():
    docx = pytest.importorskip("docx")
    from cfdpaper.publication.elements import add_table
    from cfdpaper.publication.style import PublicationStyle

    document = docx.Document()
    add_table(
        document,
        SectionTable(
            table_id="1",
            caption="Data",
            columns=["Description"],
            rows=[["Long text " * 200], ["Another row"], ["Last row"]],
            after_section_id="s",
            column_widths_mm=[30],
        ),
        PublicationStyle(),
    )
    assert (
        document.tables[0].rows[1].cells[0].paragraphs[0].paragraph_format.keep_with_next is False
    )
