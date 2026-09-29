"""Finite pagination controls do not change scientific text or physical sizing."""

import json

import pytest
from pydantic import ValidationError

from cfdpaper.publication.elements import SectionTable, add_table, apply_figure_pagination
from cfdpaper.publication.style import PublicationStyle


@pytest.mark.parametrize("mode", ["keep", "allow-split"])
@pytest.mark.parametrize("page_break", [False, True])
def test_figure_caption_pagination_ends_keep_chain(mode, page_break):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.styles["Caption"].paragraph_format.keep_with_next = True
    image = document.add_paragraph()
    caption = document.add_paragraph("Figure 1. Response at 300 K.", style="Caption")
    tail = document.add_paragraph("Independent following paragraph.")
    apply_figure_pagination(
        image,
        caption,
        PublicationStyle(figure_caption_pagination=mode),
        page_break_before=page_break,
    )
    assert image.paragraph_format.keep_with_next is True
    assert image.paragraph_format.page_break_before is page_break
    assert caption.paragraph_format.keep_together is (mode == "keep")
    assert caption.paragraph_format.keep_with_next is False
    assert caption.paragraph_format.widow_control is True
    assert caption.text == "Figure 1. Response at 300 K."
    assert tail.paragraph_format.keep_with_next is None


@pytest.mark.parametrize("rows", [3, 65])
@pytest.mark.parametrize("mode", ["auto", "allow-split"])
@pytest.mark.parametrize("note", ["", "Values retain the original units."])
def test_table_pagination_preserves_headers_rows_and_note(rows, mode, note):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    # Explicit terminal values must win over a keep-next caption template.
    document.styles["Caption"].paragraph_format.keep_with_next = True
    source = SectionTable(
        table_id="1",
        caption="Measured response",
        after_section_id="results",
        columns=["Case", "Temperature (K)"],
        rows=[[f"Case {index}", f"{300 + index:.2f}"] for index in range(rows)],
        numeric_columns=[1],
        note=note,
    )
    original = source.model_dump()
    add_table(document, source, PublicationStyle(table_pagination=mode))
    table = document.tables[0]
    assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert document.paragraphs[0].paragraph_format.keep_with_next is True
    assert document.paragraphs[0].paragraph_format.keep_together is True
    for index, row in enumerate(table.rows):
        assert row._tr.xpath("./w:trPr/w:cantSplit")
        if index == rows:
            expected = bool(note)
        else:
            expected = index == 0 or (rows == 3 and mode == "auto")
        for cell in row.cells:
            assert cell.paragraphs[0].paragraph_format.keep_with_next is expected
    if note:
        assert document.paragraphs[-1].paragraph_format.keep_with_next is False
        assert document.paragraphs[-1].paragraph_format.keep_together is True
    assert source.model_dump() == original


@pytest.mark.parametrize("field", ["figure_caption_pagination", "table_pagination"])
def test_pagination_options_are_bounded(field):
    with pytest.raises(ValidationError):
        PublicationStyle.model_validate({field: "shrink-to-fit"})


@pytest.mark.parametrize("mode", ["keep", "allow-split"])
def test_export_applies_caption_option_without_changing_body_or_image(tmp_path, mode):
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn
    from PIL import Image

    from cfdpaper.publication.section import export_section_docx

    Image.new("RGB", (800, 500), "white").save(tmp_path / "plot.png")
    source = {
        "title": "Thermal response",
        "style": {"figure_caption_pagination": mode},
        "paragraphs": [{"text": "The change is -2.20 K.", "figure_ids": ["1"]}],
        "figures": [{"id": "1", "path": "plot.png", "caption": "Local temperature."}],
        "references": [],
    }
    path = tmp_path / "section.json"
    path.write_text(json.dumps(source), encoding="utf-8")
    before = path.read_bytes()
    output = export_section_docx(tmp_path, tmp_path / "result.docx", layout="near-reference")
    document = docx.Document(output)
    body, image, caption = document.paragraphs[1:4]
    assert body.paragraph_format.space_before.pt == 0
    assert body.paragraph_format.space_after.pt == 0
    assert body._p.get_or_add_pPr().get_or_add_ind().get(qn("w:firstLineChars")) == "200"
    assert body.text == "The change is −\u20602.20\u00a0K."
    assert image.paragraph_format.keep_with_next is True
    assert caption.paragraph_format.keep_with_next is False
    assert caption.paragraph_format.keep_together is (mode == "keep")
    assert document.inline_shapes[0].width.mm == pytest.approx(160)
    assert document.inline_shapes[0].height.mm == pytest.approx(100)
    assert path.read_bytes() == before
